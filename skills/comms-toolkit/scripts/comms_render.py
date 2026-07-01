#!/usr/bin/env python3
"""
comms-toolkit — render a change-communications plan as an xlsx master grid + a
self-contained interactive HTML dashboard, for ANY project, from a structured
comms plan (built by Claude with the user in plain language).

Method baked in (portable change-comms):
  - Audience segmentation (Champions / Managers / Org-wide / BU-specific / SteerCo).
  - Channel mix + cadence; weekly sequencing + monthly refresh.
  - Comms-by-impact coverage; uncovered impacts flagged as GAPs.
  - Approval / QA gate (draft -> route -> reviewer SLA -> send).
  - Crisis comms scenarios (T+24h holding / T+72h substantive / T+14d recovery).

Usage:
  python3 comms_render.py --config project.json --plan comms_plan.json --outdir OUT

Only stdlib + openpyxl. Cross-platform (pathlib). Output HTML has no external
dependencies (single self-contained file).
"""
import argparse
import html
import json
from pathlib import Path

DEFAULT_BRAND = {
    "navy": "#162B75",
    "accent": "#EE2C81",
    "font": "Calibri, system-ui, sans-serif",
    "footer": "Confidential",
}

# ----------------------------------------------------------------------------- helpers
def g(rec, *keys, default=""):
    for k in keys:
        if k in rec and rec[k] not in (None, ""):
            return rec[k]
    return default


def load_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def audience_name(plan, ref):
    """Resolve an audience id or name to a display name."""
    for a in plan.get("audiences", []):
        if str(ref) in (str(a.get("id")), str(a.get("name"))):
            return a.get("name", str(ref))
    return str(ref)


# ----------------------------------------------------------------------------- normalize
def normalize_activities(plan):
    out = []
    for i, a in enumerate(plan.get("activities", []), 1):
        out.append({
            "id": i,
            "week": g(a, "week"),
            "date": g(a, "date"),
            "audience": audience_name(plan, g(a, "audience")),
            "channel": g(a, "channel"),
            "title": g(a, "title"),
            "owner": g(a, "owner"),
            "comm_ref": g(a, "comm_ref"),
            "review_gate": g(a, "review_gate"),
            "status": g(a, "status", default="Planned"),
        })
    return out


def normalize_coverage(plan):
    out = []
    for c in plan.get("coverage", []):
        status = str(g(c, "status", default="covered")).lower()
        out.append({
            "impact": g(c, "impact"),
            "impact_id": g(c, "impact_id"),
            "vehicle": g(c, "vehicle"),
            "audience": audience_name(plan, g(c, "audience")),
            "status": "gap" if status == "gap" else "covered",
        })
    return out


# ----------------------------------------------------------------------------- xlsx
def write_xlsx(path, sheet_name, columns, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    thin = Side(style="thin", color="D9D9D9")
    border = Border(thin, thin, thin, thin)
    ws.append([c[1] for c in columns])
    for ci, _ in enumerate(columns, 1):
        cell = ws.cell(1, ci)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor="162B75")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 26
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
    for r in rows:
        ws.append([r.get(c[0], "") for c in columns])
    for ci, col in enumerate(columns, 1):
        L = get_column_letter(ci)
        ws.column_dimensions[L].width = col[2] if len(col) > 2 else 22
        for cell in ws[L][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            cell.font = Font(size=10)
    wb.save(path)


# ----------------------------------------------------------------------------- HTML engine
TEMPLATE = r"""<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0"><title>%%TITLE%%</title>
<style>
:root{--navy:%%NAVY%%;--accent:%%ACCENT%%;--ink:#1d2433;--mut:#6b7280;--line:#e5e7eb;--bg:#f5f6f8}
*{box-sizing:border-box}body{margin:0;font-family:%%FONT%%;color:var(--ink);background:var(--bg);font-size:14px}
header{background:linear-gradient(135deg,var(--navy),#0E2841);color:#fff;padding:18px 28px}
header h1{margin:0;font-size:1.15rem;font-weight:700}header .sub{opacity:.85;font-size:.85rem;margin-top:3px}
.wrap{max-width:1500px;margin:0 auto;padding:18px 28px 80px}
.stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:14px;margin:16px 0}
.card{background:#fff;border:1px solid var(--line);border-radius:10px;padding:14px 16px}
.card .v{font-size:1.6rem;font-weight:800;color:var(--navy)}
.card .l{font-size:.72rem;color:var(--mut);text-transform:uppercase;letter-spacing:.04em;margin-top:2px}
.card.gap .v{color:#b40020}
.bar{display:flex;flex-wrap:wrap;gap:8px;align-items:center;margin:10px 0}
.bar select,.bar input{padding:7px 10px;border:1px solid var(--line);border-radius:7px;font:inherit;background:#fff}
.bar input{flex:1;min-width:180px}
.btn{background:var(--navy);color:#fff;border:none;border-radius:7px;padding:8px 14px;cursor:pointer;font:inherit}
.tabs{display:flex;gap:4px;border-bottom:2px solid var(--line);margin-top:8px;flex-wrap:wrap}
.tab{padding:8px 14px;cursor:pointer;border-radius:7px 7px 0 0;font-weight:600;color:var(--mut)}
.tab.active{color:var(--navy);background:#fff;border:1px solid var(--line);border-bottom:2px solid #fff;margin-bottom:-2px}
.view{display:none}.view.active{display:block}
.pv-head{margin:14px 0 6px}.pv-head h3{color:var(--navy);margin:0}
.pv-note{color:var(--mut);font-size:.78rem;margin-top:3px}
table{width:100%;border-collapse:collapse;background:#fff;border:1px solid var(--line);border-radius:10px;overflow:hidden}
thead th{position:sticky;top:0;background:var(--navy);color:#fff;text-align:left;padding:9px 11px;font-size:.74rem;cursor:pointer;white-space:nowrap}
tbody td{padding:9px 11px;border-top:1px solid var(--line);font-size:.82rem;vertical-align:top}
tbody tr:hover{background:#f0f3fb}
tr.gap-row{background:#fdecea}tr.gap-row:hover{background:#fbdcd8}
.chip{display:inline-block;padding:2px 9px;border-radius:11px;color:#fff;font-size:.68rem;font-weight:600;white-space:nowrap;background:var(--navy)}
.chip.acc{background:var(--accent)}
.flag{display:inline-block;padding:2px 8px;border-radius:6px;background:#b40020;color:#fff;font-size:.66rem;font-weight:800;text-transform:uppercase}
.ok{display:inline-block;padding:2px 8px;border-radius:6px;background:#04A577;color:#fff;font-size:.66rem;font-weight:700}
.st{font-weight:700}.st.Sent{color:#04A577}.st.Approved{color:#1d6fb8}.st.Draft{color:#d35400}.st.Planned{color:#6b7280}
.matrix{overflow:auto}.matrix table{min-width:520px}
.matrix td.n{text-align:center;font-weight:700;color:var(--navy)}
.matrix td.z{text-align:center;color:#c7ccd6}
.crisis{display:grid;grid-template-columns:1fr;gap:14px}
.cbox{background:#fff;border:1px solid var(--line);border-radius:12px;padding:14px 16px;border-top:4px solid var(--accent)}
.cbox h4{margin:0 0 8px;color:var(--navy)}
.cstep{margin:8px 0}.cstep .lbl{font-size:.68rem;text-transform:uppercase;letter-spacing:.04em;color:var(--mut);font-weight:700;margin-bottom:2px}
.gate{background:#fff;border:1px solid var(--line);border-radius:12px;padding:14px 16px;margin:12px 0}
.gate .row{display:flex;flex-wrap:wrap;gap:22px;margin:8px 0}
.gate .m .v{font-size:1.2rem;font-weight:800;color:var(--navy)}.gate .m .l{font-size:.66rem;color:var(--mut);text-transform:uppercase}
footer{text-align:center;color:var(--mut);font-size:.72rem;padding:18px}
</style></head><body>
<header><h1>%%TITLE%%</h1><div class="sub">%%SUBTITLE%%</div></header>
<div class="wrap">
  <div class="tabs" id="tabs"></div>
  <div class="view active" id="v-cal">
    <div class="stats" id="stats"></div>
    <div class="bar" id="bar"></div>
    <div style="overflow:auto"><table id="tbl"><thead><tr id="head"></tr></thead><tbody id="body"></tbody></table></div>
  </div>
  <div class="view" id="v-cov">
    <div class="pv-head"><h3>Comms-by-Impact Coverage</h3>
      <div class="pv-note">Every change impact should map to a comm vehicle. Rows highlighted red are GAPs — the escalation trigger.</div></div>
    <div style="overflow:auto"><table id="covtbl"><thead><tr><th>Impact</th><th>Audience</th><th>Vehicle</th><th>Status</th></tr></thead><tbody id="covbody"></tbody></table></div>
  </div>
  <div class="view" id="v-mix">
    <div class="pv-head"><h3>Audience x Channel Mix</h3>
      <div class="pv-note">Count of planned activities per audience and channel — spot thin coverage.</div></div>
    <div class="matrix" id="mix"></div>
  </div>
  <div class="view" id="v-gate">
    <div class="pv-head"><h3>Approval / QA Gate</h3>
      <div class="pv-note">Draft -> route -> reviewer SLA -> send. Factual, tone, scope-leak, sensitivity checks.</div></div>
    <div id="gate"></div>
  </div>
  <div class="view" id="v-crisis">
    <div class="pv-head"><h3>Crisis Comms</h3>
      <div class="pv-note">Activation scenarios, each with a T+24h holding statement, T+72h substantive, T+14d recovery.</div></div>
    <div class="crisis" id="crisis"></div>
  </div>
</div>
<footer>%%FOOTER%%</footer>
<script>
const DATA=%%DATA%%, COV=%%COV%%, PLAN=%%PLAN%%;
const $=s=>document.querySelector(s), esc=s=>String(s==null?'':s).replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
const COLS=[{k:'week',l:'Week'},{k:'date',l:'Date'},{k:'audience',l:'Audience'},{k:'channel',l:'Channel'},
  {k:'title',l:'Title'},{k:'owner',l:'Owner'},{k:'comm_ref',l:'Comm Ref'},{k:'review_gate',l:'Review Gate'},{k:'status',l:'Status'}];
const FILTERS=[{k:'audience',l:'All Audiences'},{k:'channel',l:'All Channels'},{k:'status',l:'All Status'}];
let sortKey=null, sortDir=1;
function distinct(k){return [...new Set(DATA.map(d=>d[k]).filter(v=>v!==''&&v!=null))].sort();}
function buildBar(){
  const bar=$('#bar');
  FILTERS.forEach(f=>{const s=document.createElement('select');s.id='f-'+f.k;
    s.innerHTML='<option value="">'+esc(f.l)+'</option>'+distinct(f.k).map(v=>'<option>'+esc(v)+'</option>').join('');
    s.onchange=render;bar.appendChild(s);});
  const q=document.createElement('input');q.id='q';q.placeholder='Search...';q.oninput=render;bar.appendChild(q);
  const b=document.createElement('button');b.className='btn';b.textContent='Reset';
  b.onclick=()=>{FILTERS.forEach(f=>$('#f-'+f.k).value='');$('#q').value='';render();};bar.appendChild(b);
}
function filtered(){
  let d=DATA.slice();
  FILTERS.forEach(f=>{const v=$('#f-'+f.k).value;if(v)d=d.filter(x=>String(x[f.k])===v);});
  const q=($('#q').value||'').toLowerCase();
  if(q)d=d.filter(x=>JSON.stringify(x).toLowerCase().includes(q));
  if(sortKey)d.sort((a,b)=>String(a[sortKey]).localeCompare(String(b[sortKey]),undefined,{numeric:true})*sortDir);
  return d;
}
function cell(d,c){
  const v=d[c.k];
  if(c.k==='status')return '<span class="st '+esc(v)+'">'+esc(v)+'</span>';
  if(c.k==='audience')return '<span class="chip">'+esc(v)+'</span>';
  if(c.k==='channel')return '<span class="chip acc">'+esc(v)+'</span>';
  return esc(v);
}
function render(){
  const rows=filtered();
  $('#head').innerHTML=COLS.map(c=>'<th data-k="'+c.k+'">'+esc(c.l)+(sortKey===c.k?(sortDir>0?' ▴':' ▾'):'')+'</th>').join('');
  $('#head').querySelectorAll('th').forEach(th=>th.onclick=()=>{const k=th.dataset.k;sortDir=(sortKey===k?-sortDir:1);sortKey=k;render();});
  $('#body').innerHTML=rows.map(d=>'<tr>'+COLS.map(c=>'<td>'+cell(d,c)+'</td>').join('')+'</tr>').join('');
  buildStats();
}
function buildStats(){
  const gaps=COV.filter(c=>c.status==='gap').length;
  const cards=[['Activities',DATA.length,''],['Audiences',distinct('audience').length,''],
    ['Channels',distinct('channel').length,''],['Coverage Gaps',gaps,gaps>0?'gap':'']];
  $('#stats').innerHTML=cards.map(c=>'<div class="card '+c[2]+'"><div class="v">'+c[1]+'</div><div class="l">'+esc(c[0])+'</div></div>').join('');
}
function buildCoverage(){
  $('#covbody').innerHTML=COV.map(c=>{
    const gap=c.status==='gap';
    return '<tr class="'+(gap?'gap-row':'')+'"><td>'+esc(c.impact)+'</td><td>'+esc(c.audience)+'</td><td>'+esc(c.vehicle)+
      '</td><td>'+(gap?'<span class="flag">GAP</span>':'<span class="ok">covered</span>')+'</td></tr>';
  }).join('');
}
function buildMix(){
  const auds=[...new Set(DATA.map(d=>d.audience).filter(Boolean))].sort();
  const chans=[...new Set(DATA.map(d=>d.channel).filter(Boolean))].sort();
  if(!auds.length||!chans.length){$('#mix').innerHTML='<p style="color:var(--mut)">No activities yet.</p>';return;}
  let h='<table><thead><tr><th>Audience \\ Channel</th>'+chans.map(c=>'<th>'+esc(c)+'</th>').join('')+'<th>Total</th></tr></thead><tbody>';
  auds.forEach(a=>{
    let rowTot=0;
    h+='<tr><td><b>'+esc(a)+'</b></td>'+chans.map(c=>{
      const n=DATA.filter(d=>d.audience===a&&d.channel===c).length;rowTot+=n;
      return n?'<td class="n">'+n+'</td>':'<td class="z">.</td>';}).join('')+'<td class="n">'+rowTot+'</td></tr>';
  });
  h+='</tbody></table>';
  $('#mix').innerHTML=h;
}
function buildGate(){
  const gate=PLAN.approval_gate||{};
  if(!Object.keys(gate).length){$('#gate').innerHTML='<p style="color:var(--mut)">No approval gate defined.</p>';return;}
  const revs=(gate.reviewers||[]).map(r=>'<span class="chip">'+esc(r)+'</span>').join(' ');
  $('#gate').innerHTML='<div class="gate"><div class="row">'+
    '<div class="m"><div class="v">'+esc(gate.draft_by||'-')+'</div><div class="l">Draft by</div></div>'+
    '<div class="m"><div class="v">'+esc(gate.sla_hours||'-')+'h</div><div class="l">Reviewer SLA</div></div>'+
    '<div class="m"><div class="v">'+esc(gate.send||'-')+'</div><div class="l">Send</div></div></div>'+
    '<div class="cstep"><div class="lbl">Route</div>'+esc(gate.route||'-')+'</div>'+
    '<div class="cstep"><div class="lbl">Reviewers</div>'+(revs||'-')+'</div></div>';
}
function buildCrisis(){
  const cr=PLAN.crisis||[];
  if(!cr.length){$('#crisis').innerHTML='<p style="color:var(--mut)">No crisis scenarios defined.</p>';return;}
  $('#crisis').innerHTML=cr.map(c=>'<div class="cbox"><h4>'+esc(c.scenario)+'</h4>'+
    '<div class="cstep"><div class="lbl">T+24h holding statement</div>'+esc(c.holding_24h||'-')+'</div>'+
    '<div class="cstep"><div class="lbl">T+72h substantive</div>'+esc(c.substantive_72h||'-')+'</div>'+
    '<div class="cstep"><div class="lbl">T+14d recovery</div>'+esc(c.recovery_14d||'-')+'</div></div>').join('');
}
function buildTabs(){
  const tabs=[{l:'Comms Calendar',v:'v-cal'},{l:'Coverage',v:'v-cov'},{l:'Audience x Channel',v:'v-mix'},
    {l:'Approval Gate',v:'v-gate'},{l:'Crisis Comms',v:'v-crisis'}];
  const t=$('#tabs');
  tabs.forEach((tb,i)=>{const e=document.createElement('div');e.className='tab'+(i===0?' active':'');e.textContent=tb.l;
    e.onclick=()=>{document.querySelectorAll('.tab').forEach(x=>x.classList.remove('active'));
      document.querySelectorAll('.view').forEach(x=>x.classList.remove('active'));
      e.classList.add('active');$('#'+tb.v).classList.add('active');};t.appendChild(e);});
}
buildTabs();buildBar();render();buildCoverage();buildMix();buildGate();buildCrisis();
</script></body></html>"""


def render_html(out_path, *, title, subtitle, data, coverage, plan, brand):
    out = (TEMPLATE
           .replace("%%TITLE%%", html.escape(title))
           .replace("%%SUBTITLE%%", html.escape(subtitle))
           .replace("%%NAVY%%", brand["navy"])
           .replace("%%ACCENT%%", brand["accent"])
           .replace("%%FONT%%", brand["font"])
           .replace("%%FOOTER%%", html.escape(brand["footer"]))
           .replace("%%DATA%%", json.dumps(data, ensure_ascii=False))
           .replace("%%COV%%", json.dumps(coverage, ensure_ascii=False))
           .replace("%%PLAN%%", json.dumps(plan, ensure_ascii=False)))
    Path(out_path).write_text(out, encoding="utf-8")


# ----------------------------------------------------------------------------- build
def build(project, plan, outdir):
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    brand = {**DEFAULT_BRAND, **project.get("brand", {})}
    pname = project.get("project_name", "Project")

    activities = normalize_activities(plan)
    coverage = normalize_coverage(plan)
    gaps = sum(1 for c in coverage if c["status"] == "gap")

    # ---- xlsx master grid ----
    cols = [("week", "Week", 8), ("date", "Date", 14), ("audience", "Audience", 22),
            ("channel", "Channel", 22), ("title", "Title", 40), ("owner", "Owner", 18),
            ("comm_ref", "Comm Ref", 16), ("review_gate", "Review Gate", 16),
            ("status", "Status", 14)]
    xlsx_path = outdir / f"{pname} Comms Master Grid.xlsx"
    write_xlsx(xlsx_path, "Comms Master Grid", cols, activities)

    # ---- dashboard ----
    html_path = outdir / f"{pname} Comms Dashboard.html"
    subtitle = (f"{len(activities)} activities · {len({a['audience'] for a in activities if a['audience']})} audiences "
                f"· {len({a['channel'] for a in activities if a['channel']})} channels · {gaps} coverage gap(s)")
    render_html(html_path, title=f"{pname} — Change Communications",
                subtitle=subtitle, data=activities, coverage=coverage, plan=plan, brand=brand)

    print(f"Comms Master Grid -> {xlsx_path}")
    print(f"Comms Dashboard   -> {html_path}")
    print(f"{len(activities)} activities, {len(coverage)} coverage rows, {gaps} gap(s)")


def _guard(outdir, names, force):
    """Refuse to overwrite existing deliverables unless --force is set."""
    import os, sys
    existing = [n for n in names if os.path.exists(os.path.join(outdir, n))]
    if existing and not force:
        sys.exit("Refusing to overwrite existing output(s) in %r:\n  %s\n"
                 "Re-run with --force to overwrite, or choose a different --outdir."
                 % (outdir, "\n  ".join(existing)))


def main():
    ap = argparse.ArgumentParser(description="Render a change-comms plan to xlsx + HTML dashboard.")
    ap.add_argument("--config", required=True, help="project.json")
    ap.add_argument("--plan", required=True, help="comms_plan.json")
    ap.add_argument("--outdir", default=".")
    ap.add_argument("--force", action="store_true", help="overwrite existing outputs in --outdir")
    a = ap.parse_args()
    project = load_json(a.config)
    plan = load_json(a.plan)
    _pn = project.get("project_name", "Project")
    _guard(a.outdir, [f"{_pn} Comms Master Grid.xlsx", f"{_pn} Comms Dashboard.html"], a.force)
    build(project, plan, a.outdir)


if __name__ == "__main__":
    main()
