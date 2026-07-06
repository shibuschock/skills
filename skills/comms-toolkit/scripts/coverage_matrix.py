#!/usr/bin/env python3
"""
coverage_matrix.py — cross-reference every CIA change impact against a comms
plan's coverage entries, and emit a coverage matrix. Uncovered impacts are
flagged GAP (the non-negotiable escalation trigger). Stdlib + openpyxl.

An impact is "covered" when the plan has at least one coverage row that maps to
it (by impact_id index OR by a near-identical title) with status != "gap".
Impacts with no such row — or only status="gap" rows — are flagged GAP.

Usage:
  python3 coverage_matrix.py --cia cia_records.json --plan comms_plan.json --outdir OUT
                             [--config project.json] [--threshold 0.80]

CIA records follow the cia-builder schema (impacts have `title`,
`roles_impacted`, etc.). Ids are assigned 1-based by position, matching the
cia-builder renderer.
"""
import argparse
import difflib
import json
import re
from pathlib import Path


def norm(s):
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", "", str(s or "").lower())).strip()


def load_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def audience_name(plan, ref):
    for a in plan.get("audiences", []):
        if str(ref) in (str(a.get("id")), str(a.get("name"))):
            return a.get("name", str(ref))
    return str(ref)


def coverage_for(impact_id, title, plan, threshold):
    """Return (covered_hits, gap_hits) — coverage rows mapping to this impact,
    split by status. gap_hits are tracked gaps: an intended vehicle is named
    but not yet built."""
    hits, gap_hits = [], []
    ntitle = norm(title)
    for c in plan.get("coverage", []):
        by_id = c.get("impact_id") not in (None, "") and str(c.get("impact_id")) == str(impact_id)
        by_title = ntitle and difflib.SequenceMatcher(None, ntitle, norm(c.get("impact"))).ratio() >= threshold
        if by_id or by_title:
            (gap_hits if str(c.get("status", "covered")).lower() == "gap" else hits).append(c)
    return hits, gap_hits


def write_xlsx(path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    cols = [("impact", "Impact", 48), ("audience", "Audience", 26),
            ("vehicle", "Vehicle", 26), ("status", "Status", 12)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Coverage Matrix"
    thin = Side(style="thin", color="D9D9D9")
    border = Border(thin, thin, thin, thin)
    ws.append([c[1] for c in cols])
    for ci, _ in enumerate(cols, 1):
        cell = ws.cell(1, ci)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor="162B75")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 26
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"
    gap_fill = PatternFill("solid", fgColor="FDECEA")
    for r in rows:
        ws.append([r.get(c[0], "") for c in cols])
        if r.get("status") == "GAP":
            for ci in range(1, len(cols) + 1):
                ws.cell(ws.max_row, ci).fill = gap_fill
    for ci, col in enumerate(cols, 1):
        L = get_column_letter(ci)
        ws.column_dimensions[L].width = col[2]
        for cell in ws[L][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            cell.font = Font(size=10)
    wb.save(path)


def build(cia, plan, outdir, project, threshold):
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    pname = (project or {}).get("project_name", "Project")

    rows = []
    tracked_gaps = unmapped_gaps = 0
    for i, impact in enumerate(cia, 1):
        title = impact.get("title", f"Impact {i}")
        hits, gap_hits = coverage_for(i, title, plan, threshold)
        if hits:
            for c in hits:
                rows.append({"impact": title,
                             "audience": audience_name(plan, c.get("audience", "")),
                             "vehicle": c.get("vehicle", ""),
                             "status": "Covered"})
        elif gap_hits:
            tracked_gaps += 1
            for c in gap_hits:
                rows.append({"impact": title,
                             "audience": audience_name(plan, c.get("audience", "")),
                             "vehicle": c.get("vehicle", ""),
                             "status": "GAP"})
        else:
            unmapped_gaps += 1
            rows.append({"impact": title, "audience": "", "vehicle": "", "status": "GAP"})

    gaps = tracked_gaps + unmapped_gaps
    path = outdir / f"{pname} Coverage Matrix.xlsx"
    write_xlsx(path, rows)
    print(f"Coverage Matrix -> {path}")
    print(f"{len(cia)} impacts checked; {gaps} GAP(s): "
          f"{tracked_gaps} tracked gap(s) (intended vehicle named, not yet built) + "
          f"{unmapped_gaps} unmapped impact(s) (no coverage row at all).")
    if gaps:
        print("ESCALATE: gap impacts are the non-negotiable escalation trigger "
              "(tracked gaps need the vehicle built; unmapped impacts need a coverage decision).")


def _guard(outdir, names, force):
    """Refuse to overwrite existing deliverables unless --force is set."""
    import os, sys
    existing = [n for n in names if os.path.exists(os.path.join(outdir, n))]
    if existing and not force:
        sys.exit("Refusing to overwrite existing output(s) in %r:\n  %s\n"
                 "Re-run with --force to overwrite, or choose a different --outdir."
                 % (outdir, "\n  ".join(existing)))


def main():
    ap = argparse.ArgumentParser(description="Cross-reference CIA impacts against a comms plan's coverage.")
    ap.add_argument("--cia", required=True, help="cia_records.json (cia-builder schema)")
    ap.add_argument("--plan", required=True, help="comms_plan.json")
    ap.add_argument("--outdir", default=".")
    ap.add_argument("--force", action="store_true", help="overwrite existing outputs in --outdir")
    ap.add_argument("--config", help="project.json (optional, for project name)")
    ap.add_argument("--threshold", type=float, default=0.80, help="title-match ratio for coverage (default 0.80)")
    a = ap.parse_args()
    cia = load_json(a.cia)
    plan = load_json(a.plan)
    project = load_json(a.config) if a.config else None
    _pn = (project or {}).get("project_name", "Project")
    _guard(a.outdir, [f"{_pn} Coverage Matrix.xlsx"], a.force)
    build(cia, plan, a.outdir, project, a.threshold)


if __name__ == "__main__":
    main()
