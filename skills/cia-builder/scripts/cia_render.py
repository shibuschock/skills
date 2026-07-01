#!/usr/bin/env python3
"""
cia-builder — render a Change Impact Assessment (CIA) as an xlsx workbook + a
self-contained interactive HTML dashboard, for ANY project, from structured
records (extracted by Claude from raw current/future-state transcripts).

Methodology baked in (portable OCM):
  - 6 MECE Change Dimensions + a single Primary; severity_label + impact_score
    (= complexity x severity) derived; scope disposition field.
  - Optional analytical layers: TOM Shifts, Value Realization (benefits/value
    levers, non-monetary), Training Needs, Future-State / Workshop status.

Usage:
  python3 cia_render.py cia    --records cia_records.json --config project.json --outdir OUT
  python3 cia_render.py readme --config project.json --outdir OUT

Records = a JSON array of objects (see references/CIA_SCHEMA.md).
Only stdlib + openpyxl. Cross-platform. Output HTML has no external dependencies.
"""
import argparse, datetime, html, json, os, re, sys

# ----------------------------------------------------------------------------- helpers
SEV_LABEL = {5: "Critical", 4: "High", 3: "Medium", 2: "Low", 1: "Minimal"}
DIMENSIONS = ["Role/Accountability", "Process/Policy", "Ways of Working",
              "Data/Decision Inputs", "Skills/Capability", "Mindset/Culture"]
DIM_COLORS = {"Role/Accountability": "#162B75", "Ways of Working": "#04A577",
              "Skills/Capability": "#F08301", "Mindset/Culture": "#EE2C81",
              "Process/Policy": "#FF533C", "Data/Decision Inputs": "#7A4FBE"}
# Future-state status -> ribbon color (green confirmed / amber pending / grey assumed)
STATUS_COLORS = {"Discussed-Confirmed": "#04A577", "Discussed-Pending": "#F08301", "Assumed": "#9aa3b2"}
DEFAULT_BRAND = {"navy": "#162B75", "magenta": "#EE2C81", "orange": "#F08301",
                 "teal": "#04A577", "coral": "#FF533C", "font": "Calibri, system-ui, sans-serif",
                 "footer": "Confidential"}
DEFAULT_STALE_DAYS = 30

def g(rec, *keys, default=""):
    for k in keys:
        if k in rec and rec[k] not in (None, ""):
            return rec[k]
    return default

def as_int(v):
    try: return int(float(v))
    except Exception: return None

def as_list(v):
    if v is None or v == "": return []
    if isinstance(v, (list, tuple)): return [str(x).strip() for x in v if str(x).strip()]
    return [s.strip() for s in re.split(r"[;,]", str(v)) if s.strip()]

def days_since(date_str):
    if not date_str: return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d"):
        try:
            d = datetime.datetime.strptime(str(date_str)[:10], fmt).date()
            return (datetime.date.today() - d).days
        except Exception:
            continue
    return None

# ----------------------------------------------------------------------------- normalize
def normalize_cia(records, stale_days=DEFAULT_STALE_DAYS):
    out = []
    for i, r in enumerate(records, 1):
        sev = as_int(g(r, "severity")); cx = as_int(g(r, "complexity"))
        dims = [d for d in (r.get("dimensions") or []) if d in DIMENSIONS]
        primary = g(r, "primary_dimension") or (dims[0] if dims else "")
        if primary and primary not in dims and primary in DIMENSIONS:
            dims = [primary] + dims
        # Future-state status: explicit value wins; else infer Assumed when no evidence/workshop ref.
        status = g(r, "future_state_status")
        if status not in STATUS_COLORS:
            has_ev = bool(g(r, "workshop_reference") or g(r, "evidence_source"))
            status = "Discussed-Pending" if has_ev else "Assumed"
        last_rev = g(r, "last_reviewed")
        ds = days_since(last_rev)
        stale = bool(ds is not None and ds > stale_days)
        out.append({
            "id": r.get("id", i),
            "title": g(r, "title"),
            "functional_area": g(r, "functional_area", "area"),
            "business_unit": g(r, "business_unit", "bu"),
            "l1_process": g(r, "l1_process"), "l2_process": g(r, "l2_process"),
            "l3_process": g(r, "l3_process"),
            "current_state": g(r, "current_state"), "future_state": g(r, "future_state"),
            "process_change": g(r, "process_change"),
            "roles_impacted": g(r, "roles_impacted"),
            "role_changes": r.get("role_changes") or {},
            "severity": sev or "", "complexity": cx or "",
            "severity_label": SEV_LABEL.get(sev, ""),
            "impact_score": (cx * sev) if (cx and sev) else "",
            "priority": g(r, "priority"),
            "sentiment": g(r, "sentiment"),
            "change_considerations": g(r, "change_considerations"),
            "mitigation": g(r, "mitigation"),
            "dim": {d: True for d in dims},
            "primary_dimension": primary,
            "scope": g(r, "scope", default="In Scope"),
            # optional analytical layers
            "future_state_status": status,
            "tom_shift": g(r, "tom_shift"),
            "value_levers": as_list(r.get("value_levers")),
            "benefit_at_risk": g(r, "benefit_at_risk"),
            "training_modality": as_list(r.get("training_modality")),
            "estimated_training_hours": (as_int(g(r, "estimated_training_hours")) or ""),
            "access_constraints": g(r, "access_constraints"),
            "evidence_source": g(r, "evidence_source"),
            "last_reviewed": last_rev,
            "workshop_reference": g(r, "workshop_reference"),
            "stale": stale,
        })
    return out

# ----------------------------------------------------------------------------- xlsx
def write_xlsx(path, sheet_name, columns, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = sheet_name[:31]
    thin = Side(style="thin", color="D9D9D9"); border = Border(thin, thin, thin, thin)
    ws.append([c[1] for c in columns])
    for ci, _ in enumerate(columns, 1):
        c = ws.cell(1, ci); c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor="162B75"); c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = border
    ws.freeze_panes = "B2"; ws.row_dimensions[1].height = 26
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
    for r in rows:
        ws.append([fmt_cell(r.get(c[0], "")) for c in columns])
    for ci, col in enumerate(columns, 1):
        L = get_column_letter(ci); ws.column_dimensions[L].width = col[2] if len(col) > 2 else 22
        for cell in ws[L][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top"); cell.border = border
            cell.font = Font(size=10)
    wb.save(path)

def fmt_cell(v):
    if isinstance(v, bool):
        return "Yes" if v else ""
    if isinstance(v, dict):
        return ", ".join(k for k, val in v.items() if val)
    if isinstance(v, (list, tuple)):
        return ", ".join(str(x) for x in v)
    return v

# ----------------------------------------------------------------------------- HTML engine
TEMPLATE = r"""<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0"><title>%%TITLE%%</title>
<style>
:root{--navy:%%NAVY%%;--mag:%%MAG%%;--ink:#1d2433;--mut:#6b7280;--line:#e5e7eb;--bg:#f5f6f8}
*{box-sizing:border-box}body{margin:0;font-family:%%FONT%%;color:var(--ink);background:var(--bg);font-size:14px}
header{background:linear-gradient(135deg,var(--navy),#0E2841);color:#fff;padding:18px 28px}
header h1{margin:0;font-size:1.15rem;font-weight:700}header .sub{opacity:.85;font-size:.85rem;margin-top:3px}
.wrap{max-width:1500px;margin:0 auto;padding:18px 28px 80px}
.stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:14px;margin:16px 0}
.card{background:#fff;border:1px solid var(--line);border-radius:10px;padding:14px 16px}
.card .v{font-size:1.6rem;font-weight:800;color:var(--navy)}.card .l{font-size:.72rem;color:var(--mut);text-transform:uppercase;letter-spacing:.04em;margin-top:2px}
.bar{display:flex;flex-wrap:wrap;gap:8px;align-items:center;margin:10px 0}
.bar select,.bar input{padding:7px 10px;border:1px solid var(--line);border-radius:7px;font:inherit;background:#fff}
.bar input{flex:1;min-width:180px}.btn{background:var(--navy);color:#fff;border:none;border-radius:7px;padding:8px 14px;cursor:pointer;font:inherit}
.tabs{display:flex;gap:4px;border-bottom:2px solid var(--line);margin-top:8px;flex-wrap:wrap}
.tab{padding:8px 14px;cursor:pointer;border-radius:7px 7px 0 0;font-weight:600;color:var(--mut)}
.tab.active{color:var(--navy);background:#fff;border:1px solid var(--line);border-bottom:2px solid #fff;margin-bottom:-2px}
.view{display:none}.view.active{display:block}
.pv-head{display:flex;align-items:baseline;justify-content:space-between;margin:14px 0 6px}
.pv-head h3{color:var(--navy);margin:0}.pv-note{color:var(--mut);font-size:.78rem}
.tiles{display:grid;grid-template-columns:repeat(auto-fill,minmax(260px,1fr));gap:14px}
.ptile{background:#fff;border:1px solid var(--line);border-radius:12px;padding:14px 16px;cursor:pointer;transition:box-shadow .12s,transform .12s;border-top:4px solid var(--navy)}
.ptile:hover{box-shadow:0 6px 18px rgba(20,40,90,.10);transform:translateY(-1px)}
.ptile.pv-zero{opacity:.7;border-top-color:#c7ccd6}
.pt-h{font-weight:700;color:var(--navy);font-size:.98rem;display:flex;align-items:center;gap:8px;justify-content:space-between}
.pt-meta{color:var(--mut);font-size:.78rem;margin:5px 0 2px;line-height:1.4}
.pm-row{display:flex;flex-wrap:wrap;gap:14px;margin-top:10px}
.pm{min-width:54px}.pm-v{font-size:1.25rem;font-weight:800;color:var(--ink)}.pm-l{font-size:.64rem;color:var(--mut);text-transform:uppercase;letter-spacing:.03em}
.zero-flag{font-size:.6rem;background:#fdecea;color:#b40020;padding:2px 7px;border-radius:9px;font-weight:700;text-transform:uppercase}
.pt-print{background:none;border:1px solid var(--line);border-radius:6px;cursor:pointer;color:var(--mut);font-size:.8rem;padding:2px 7px}
.pt-print:hover{color:var(--navy);border-color:var(--navy)}
table{width:100%;border-collapse:collapse;background:#fff;border:1px solid var(--line);border-radius:10px;overflow:hidden}
thead th{position:sticky;top:0;background:var(--navy);color:#fff;text-align:left;padding:9px 11px;font-size:.74rem;cursor:pointer;white-space:nowrap}
tbody td{padding:9px 11px;border-top:1px solid var(--line);font-size:.82rem;vertical-align:top}
tbody tr.row,tbody tr.prow{cursor:pointer}tbody tr.row:hover,tbody tr.prow:hover{background:#f0f3fb}
.chip{display:inline-block;padding:2px 9px;border-radius:11px;color:#fff;font-size:.68rem;font-weight:600;margin:1px 2px 1px 0;white-space:nowrap}
.chip.prim{border:2px solid #111;font-weight:800}.chip.sec{opacity:.72}
.chip.lever{background:#7A4FBE}.chip.mod{background:#5a6b86}
.scope{display:inline-block;padding:2px 8px;border-radius:6px;background:#6b7280;color:#fff;font-size:.66rem;font-weight:700}
.statc{display:inline-block;padding:2px 8px;border-radius:6px;color:#fff;font-size:.66rem;font-weight:700;white-space:nowrap}
.stale-flag{display:inline-block;padding:1px 7px;border-radius:9px;background:#fdecea;color:#b40020;font-size:.62rem;font-weight:700}
.pri{font-weight:700}.pri.Critical{color:#b40020}.pri.High{color:#d35400}.pri.Medium{color:#1d6fb8}.pri.Low{color:#6b7280}
.detail{background:#fbfcfe}.detail .sec{margin:9px 0}.detail .lbl{font-size:.68rem;text-transform:uppercase;letter-spacing:.04em;color:var(--mut);font-weight:700;margin-bottom:3px}
.cf{display:grid;grid-template-columns:1fr 1fr;gap:12px}.cf .b{background:#fff;border:1px solid var(--line);border-radius:8px;padding:9px}
.cf .b.cur{border-left:3px solid #c0392b}.cf .b.fut{border-left:3px solid var(--teal,#04A577)}
.cf .b.fut.assumed{opacity:.6;font-style:italic}
ul.roles{margin:4px 0;padding-left:18px}ul.roles li{margin:3px 0}
.grid2{position:relative;width:100%;max-width:620px;height:460px;border:1px solid var(--line);background:#fff;border-radius:10px;margin:14px 0}
.grid2 .q{position:absolute;width:50%;height:50%;display:flex;align-items:flex-start;justify-content:flex-start;padding:8px;font-size:.7rem;font-weight:700;color:var(--mut)}
.grid2 .dot{position:absolute;width:11px;height:11px;border-radius:50%;background:var(--navy);transform:translate(-50%,50%);cursor:pointer;border:1px solid #fff}
.grid2 .axis{position:absolute;font-size:.68rem;color:var(--mut)}
footer{text-align:center;color:var(--mut);font-size:.72rem;padding:18px}
#onepager{display:none}
@media print{
  body{background:#fff}
  header,.tabs,.bar,.stats,footer,.view,.pt-print{display:none !important}
  body.printing #onepager{display:block !important}
  #onepager h2{color:var(--navy)}
  #onepager .op-cf{display:grid;grid-template-columns:1fr 1fr;gap:14px;margin:12px 0}
  #onepager .op-b{border:1px solid #ccc;border-radius:8px;padding:10px}
  #onepager .op-m{display:inline-block;margin:0 18px 8px 0}
  #onepager .op-m b{font-size:1.3rem;color:var(--navy)}
  #onepager li{margin:3px 0}
  .op-foot{margin-top:18px;color:#888;font-size:.7rem;border-top:1px solid #ccc;padding-top:6px}
}
</style></head><body>
<header><h1>%%TITLE%%</h1><div class="sub">%%SUBTITLE%%</div></header>
<div class="wrap">
  <div class="tabs" id="tabs"></div>
  <div class="view active" id="v-table">
    <div class="stats" id="stats"></div>
    <div class="bar" id="bar"></div>
    <div style="overflow:auto"><table id="tbl"><thead><tr id="head"></tr></thead><tbody id="body"></tbody></table></div>
  </div>
  %%EXTRA_VIEWS%%
</div>
<div id="onepager"></div>
<footer>%%FOOTER%%</footer>
<script>
const DATA=%%DATA%%, CONFIG=%%CONFIG%%, STATUS_COLORS=%%STATUSCOLORS%%, BRAND=%%BRAND%%;
const $=s=>document.querySelector(s), esc=s=>String(s==null?'':s).replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
let sortKey=null, sortDir=1, expanded=null;
function vals(d,k){const v=d[k];return Array.isArray(v)?v:(v==null||v===''?[]:[v]);}
function distinct(k){const s=new Set();DATA.forEach(d=>vals(d,k).forEach(v=>s.add(v)));return [...s].sort();}
function buildBar(){
  const bar=$('#bar');
  CONFIG.filters.forEach(f=>{const s=document.createElement('select');s.id='f-'+f.key;
    s.innerHTML='<option value="">'+esc(f.label)+'</option>'+distinct(f.key).map(v=>'<option>'+esc(v)+'</option>').join('');
    s.onchange=render;bar.appendChild(s);});
  const q=document.createElement('input');q.id='q';q.placeholder='Search…';q.oninput=render;bar.appendChild(q);
  const b=document.createElement('button');b.className='btn';b.textContent='Reset';b.onclick=()=>{CONFIG.filters.forEach(f=>$('#f-'+f.key).value='');$('#q').value='';render();};bar.appendChild(b);
}
function filtered(){
  let d=DATA.slice();
  CONFIG.filters.forEach(f=>{const v=$('#f-'+f.key).value;if(v)d=d.filter(x=>{const xv=x[f.key];return Array.isArray(xv)?xv.map(String).includes(v):String(xv)===v;});});
  const q=($('#q').value||'').toLowerCase();
  if(q)d=d.filter(x=>JSON.stringify(x).toLowerCase().includes(q));
  if(sortKey)d.sort((a,b)=>{let x=a[sortKey],y=b[sortKey];if(typeof x==='number'&&typeof y==='number')return (x-y)*sortDir;return String(x).localeCompare(String(y))*sortDir;});
  return d;
}
function dimChips(d){
  if(!d.dim)return '';
  const prim=d.primary_dimension||'';
  const order=Object.keys(d.dim).filter(k=>d.dim[k]).sort((a,b)=>(b===prim)-(a===prim));
  return order.map(n=>'<span class="chip '+(n===prim?'prim':'sec')+'" style="background:'+(CONFIG.dimColors[n]||'#888')+'">'+esc(n)+(n===prim?' ★':'')+'</span>').join('')
    +(d.scope&&d.scope!=='In Scope'?' <span class="scope">'+esc(d.scope)+'</span>':'');
}
function statusChip(v){if(!v)return '';return '<span class="statc" style="background:'+(STATUS_COLORS[v]||'#888')+'">'+esc(v)+'</span>';}
function listChips(v,cls){const a=Array.isArray(v)?v:(v?[v]:[]);return a.length?a.map(x=>'<span class="chip '+(cls||'')+'">'+esc(x)+'</span>').join(''):'<span style="color:#bbb">—</span>';}
function cell(d,c){
  const v=d[c.key];
  if(c.kind==='dims')return dimChips(d);
  if(c.kind==='pri')return '<span class="pri '+esc(v)+'">'+esc(v)+'</span>';
  if(c.kind==='status')return statusChip(v);
  if(c.kind==='levers')return listChips(v,'lever');
  if(c.kind==='mods')return listChips(v,'mod');
  if(c.kind==='stale')return v?'<span class="stale-flag">stale</span>':'';
  if(c.kind==='chip'&&v)return '<span class="chip" style="background:'+esc(c.color||'#888')+'">'+esc(v)+'</span>';
  return esc(v).slice(0,c.clip||9999);
}
function render(){
  const rows=filtered();
  $('#head').innerHTML=CONFIG.columns.map(c=>'<th data-k="'+c.key+'">'+esc(c.label)+(sortKey===c.key?(sortDir>0?' ▴':' ▾'):'')+'</th>').join('');
  $('#head').querySelectorAll('th').forEach(th=>th.onclick=()=>{const k=th.dataset.k;sortDir=(sortKey===k?-sortDir:1);sortKey=k;render();});
  $('#body').innerHTML=rows.map(d=>{
    const tr='<tr class="row" data-id="'+d[CONFIG.idKey]+'">'+CONFIG.columns.map(c=>'<td>'+cell(d,c)+'</td>').join('')+'</tr>';
    const det=(expanded===d[CONFIG.idKey])?'<tr class="detail"><td colspan="'+CONFIG.columns.length+'">'+detail(d)+'</td></tr>':'';
    return tr+det;}).join('');
  $('#body').querySelectorAll('tr.row').forEach(tr=>tr.onclick=()=>{const id=tr.dataset.id;expanded=(String(expanded)===id)?null:(isNaN(id)?id:+id);render();});
  buildStats(rows);
}
function detail(d){
  return CONFIG.detail.map(s=>{
    if(s.type==='currentfuture'){const assumed=(d.future_state_status==='Assumed')?' assumed':'';
      return '<div class="sec"><div class="cf"><div class="b cur"><div class="lbl">Current state</div>'+esc(d.current_state)+'</div><div class="b fut'+assumed+'"><div class="lbl">Future state'+(d.future_state_status?(' — '+esc(d.future_state_status)):'')+'</div>'+esc(d.future_state)+'</div></div></div>';}
    if(s.type==='dims')return '<div class="sec"><div class="lbl">'+esc(s.label)+'</div>'+dimChips(d)+'</div>';
    if(s.type==='roles'){const rc=d.role_changes||{};const names=String(d.roles_impacted||'').split(/[;,]/).map(x=>x.trim()).filter(Boolean);
      if(!names.length)return '';return '<div class="sec"><div class="lbl">'+esc(s.label)+'</div><ul class="roles">'+names.map(n=>'<li><b>'+esc(n)+'</b>'+(rc[n]?(' — '+esc(typeof rc[n]==='string'?rc[n]:(rc[n].desc||''))):'')+'</li>').join('')+'</ul></div>';}
    if(s.type==='levers'){if(!vals(d,'value_levers').length)return '';return '<div class="sec"><div class="lbl">'+esc(s.label)+'</div>'+listChips(d.value_levers,'lever')+'</div>';}
    if(s.type==='mods'){if(!vals(d,'training_modality').length)return '';return '<div class="sec"><div class="lbl">'+esc(s.label)+'</div>'+listChips(d.training_modality,'mod')+'</div>';}
    const v=d[s.key];if(v==null||v==='')return '';
    return '<div class="sec"><div class="lbl">'+esc(s.label)+'</div>'+esc(v)+'</div>';
  }).join('');
}
function buildStats(rows){
  $('#stats').innerHTML=CONFIG.stats.map(s=>{
    let val;
    if(s.type==='count')val=rows.length;
    else if(s.type==='distinct')val=new Set(rows.flatMap(r=>vals(r,s.key))).size;
    else if(s.type==='countif')val=rows.filter(r=>s.test(r)).length;
    return '<div class="card"><div class="v">'+val+'</div><div class="l">'+esc(s.label)+'</div></div>';
  }).join('');
}
CONFIG.stats.forEach(s=>{if(s.countifKey){s.type='countif';s.test=(r=>String(r[s.countifKey])===s.countifVal);}});
// ---- pivot / tiles engine (TOM Shifts, Value Realization, Training, Workshop status)
function testRec(r,t){
  if(t.startsWith!==undefined)return String(r[t.key]||'').startsWith(t.startsWith);
  if(t.eq!==undefined)return String(r[t.key])===String(t.eq);
  if(t.inSet!==undefined)return t.inSet.map(String).includes(String(r[t.key]));
  if(t.gte!==undefined)return (parseFloat(r[t.key])||0)>=t.gte;
  if(t.truthy!==undefined)return !!r[t.key]===t.truthy;
  if(t.all!==undefined)return t.all.every(tt=>testRec(r,tt));
  return false;
}
function pivotGroups(spec){
  const groups={};
  DATA.forEach(d=>{
    let keys;
    if(spec.explode)keys=vals(d,spec.explode);
    else if(spec.roles)keys=String(d.roles_impacted||'').split(/[;,]/).map(s=>s.trim()).filter(Boolean);
    else keys=vals(d,spec.groupBy);
    keys.forEach(k=>{(groups[k]=groups[k]||[]).push(d);});
  });
  (spec.vocab||[]).forEach(k=>{if(!groups[k])groups[k]=[];});
  return groups;
}
function metricVal(rows,m){
  if(m.kind==='count')return rows.length;
  if(m.kind==='sum'){const t=rows.reduce((s,r)=>s+(parseFloat(r[m.key])||0),0);return Math.round(t*10)/10;}
  if(m.kind==='distinct')return new Set(rows.flatMap(r=>vals(r,m.key))).size;
  if(m.kind==='countif')return rows.filter(r=>testRec(r,m.test)).length;
  return '';
}
function renderPivots(){
  (CONFIG.pivots||[]).forEach(spec=>{
    const host=document.getElementById(spec.view);if(!host)return;
    const body=host.querySelector('.pv-body');
    const groups=pivotGroups(spec);
    const names=Object.keys(groups).sort((a,b)=>groups[b].length-groups[a].length||String(a).localeCompare(b));
    if(!names.length){body.innerHTML='<p style="color:var(--mut)">No data for this view yet — populate the relevant fields in your records.</p>';return;}
    if(spec.layout==='tiles'){
      body.className='pv-body tiles';
      body.innerHTML=names.map(n=>{
        const rows=groups[n], meta=(spec.meta&&spec.meta[n])||null, zero=rows.length===0;
        const color=(spec.colorBy&&spec.colorBy[n])||BRAND.navy;
        const mets=spec.metrics.map(m=>'<div class="pm"><div class="pm-v">'+metricVal(rows,m)+'</div><div class="pm-l">'+esc(m.label)+'</div></div>').join('');
        const metaLine=meta?('<div class="pt-meta">'+esc(meta.description||'')+(meta.target_metric?(' <b>Target:</b> '+esc(meta.target_metric)):'')+'</div>'):'';
        const pr=spec.print?'<button class="pt-print" data-print="'+esc(n)+'" title="Print one-pager">⎙</button>':'';
        return '<div class="ptile'+(zero?' pv-zero':'')+'" style="border-top-color:'+esc(color)+'" data-drill="'+esc(spec.drill||'')+'" data-search="'+esc(spec.drillSearch||'')+'" data-val="'+esc(n)+'"><div class="pt-h"><span>'+esc(n)+(zero?' <span class="zero-flag">no impacts</span>':'')+'</span>'+pr+'</div>'+metaLine+'<div class="pm-row">'+mets+'</div></div>';
      }).join('');
    } else {
      body.className='pv-body';
      body.innerHTML='<div style="overflow:auto"><table><thead><tr><th>'+esc(spec.groupLabel||'Group')+'</th>'+spec.metrics.map(m=>'<th>'+esc(m.label)+'</th>').join('')+'</tr></thead><tbody>'+
        names.map(n=>{const rows=groups[n];return '<tr class="prow" data-drill="'+esc(spec.drill||'')+'" data-search="'+esc(spec.drillSearch||'')+'" data-val="'+esc(n)+'"><td>'+esc(n)+(rows.length===0?' <span class="zero-flag">no impacts</span>':'')+'</td>'+spec.metrics.map(m=>'<td>'+metricVal(rows,m)+'</td>').join('')+'</tr>';}).join('')+'</tbody></table></div>';
    }
    body.querySelectorAll('[data-val]').forEach(node=>node.onclick=ev=>{
      if(ev.target.classList.contains('pt-print')){printOnePager(spec,node.dataset.val,groups[node.dataset.val]);return;}
      const k=node.dataset.drill,sk=node.dataset.search,val=node.dataset.val;
      if(k)setRegisterFilter(k,val);else if(sk)setRegisterSearch(val);
    });
  });
}
function gotoRegister(){
  const t=[...document.querySelectorAll('.tab')].find(x=>x.dataset.view==='v-table');
  document.querySelectorAll('.tab').forEach(x=>x.classList.remove('active'));document.querySelectorAll('.view').forEach(x=>x.classList.remove('active'));
  if(t)t.classList.add('active');$('#v-table').classList.add('active');
}
function setRegisterFilter(key,val){
  gotoRegister();const sel=$('#f-'+key);
  if(sel){if(![...sel.options].some(o=>o.value===val)){const o=document.createElement('option');o.value=val;o.textContent=val;sel.appendChild(o);}sel.value=val;}
  render();
}
function setRegisterSearch(val){gotoRegister();const q=$('#q');if(q)q.value=val;render();}
function printOnePager(spec,name,rows){
  const meta=(spec.meta&&spec.meta[name])||{};
  const mets=spec.metrics.map(m=>'<span class="op-m"><b>'+metricVal(rows,m)+'</b> '+esc(m.label)+'</span>').join('');
  const cf=(meta.today||meta.future)?('<div class="op-cf"><div class="op-b"><b>Today</b><div>'+esc(meta.today||'—')+'</div></div><div class="op-b"><b>Future</b><div>'+esc(meta.future||'—')+'</div></div></div>'):'';
  const items=rows.slice().sort((a,b)=>(b.impact_score||0)-(a.impact_score||0)).map(r=>'<li><b>'+esc(r.title)+'</b> — '+esc(r.future_state_status)+(r.priority?(' · '+esc(r.priority)):'')+'</li>').join('');
  $('#onepager').innerHTML='<h2>'+esc(name)+'</h2>'+(meta.description?'<p>'+esc(meta.description)+'</p>':'')+cf+'<div>'+mets+'</div>'+(items?'<h3>Mapped impacts</h3><ul>'+items+'</ul>':'')+'<div class="op-foot">'+esc(CONFIG.title||'')+' — '+esc(BRAND.footer||'')+'</div>';
  document.body.classList.add('printing');window.print();
}
window.addEventListener('afterprint',()=>document.body.classList.remove('printing'));
function buildTabs(){
  const t=$('#tabs');CONFIG.tabs.forEach((tb,i)=>{const e=document.createElement('div');e.className='tab'+(i===0?' active':'');e.textContent=tb.label;e.dataset.view=tb.view;
    if(i===0)$('#v-table').classList.add('active');
    e.onclick=()=>{document.querySelectorAll('.tab').forEach(x=>x.classList.remove('active'));document.querySelectorAll('.view').forEach(x=>x.classList.remove('active'));
      e.classList.add('active');$('#'+tb.view).classList.add('active');if(tb.view==='v-grid')drawGrid();};t.appendChild(e);});
  // first tab may not be v-table; activate the configured first view
  document.querySelectorAll('.view').forEach(x=>x.classList.remove('active'));
  $('#'+CONFIG.tabs[0].view).classList.add('active');
}
function drawGrid(){
  if(!CONFIG.grid)return;const el=$('#grid');if(!el||el.dataset.done)return;el.dataset.done=1;
  el.insertAdjacentHTML('beforeend','<div class="q" style="left:0;top:0">'+CONFIG.grid.q.tl+'</div><div class="q" style="right:0;top:0;text-align:right;justify-content:flex-end">'+CONFIG.grid.q.tr+'</div><div class="q" style="left:0;bottom:0;align-items:flex-end">'+CONFIG.grid.q.bl+'</div><div class="q" style="right:0;bottom:0;align-items:flex-end;justify-content:flex-end">'+CONFIG.grid.q.br+'</div>');
  el.insertAdjacentHTML('beforeend','<div class="axis" style="left:50%;bottom:-18px;transform:translateX(-50%)">Interest →</div><div class="axis" style="left:-6px;top:50%;transform:rotate(-90deg);transform-origin:left">Influence →</div>');
  DATA.forEach(d=>{const x=d[CONFIG.grid.x],y=d[CONFIG.grid.y];if(!x||!y)return;
    const dot=document.createElement('div');dot.className='dot';dot.style.left=((x-1)/4*100)+'%';dot.style.bottom=((y-1)/4*100)+'%';
    dot.title=d[CONFIG.grid.label]+' (Infl '+y+', Int '+x+')';el.appendChild(dot);});
}
buildBar();buildTabs();renderPivots();render();
</script></body></html>"""

def render_html(out_path, *, title, subtitle, data, config, brand, extra_views=""):
    config = {**config, "title": title}
    htmlout = (TEMPLATE
        .replace("%%TITLE%%", html.escape(title))
        .replace("%%SUBTITLE%%", html.escape(subtitle))
        .replace("%%NAVY%%", brand["navy"]).replace("%%MAG%%", brand["magenta"])
        .replace("%%FONT%%", brand["font"]).replace("%%FOOTER%%", html.escape(brand["footer"]))
        .replace("%%EXTRA_VIEWS%%", extra_views)
        .replace("%%DATA%%", json.dumps(data, ensure_ascii=False))
        .replace("%%CONFIG%%", json.dumps(config, ensure_ascii=False))
        .replace("%%STATUSCOLORS%%", json.dumps(STATUS_COLORS, ensure_ascii=False))
        .replace("%%BRAND%%", json.dumps(brand, ensure_ascii=False)))
    open(out_path, "w", encoding="utf-8").write(htmlout)

def pivot_view_html(view_id, heading, note=""):
    n = f'<div class="pv-note">{html.escape(note)}</div>' if note else ""
    return (f'<div class="view" id="{view_id}"><div class="pv-head"><h3>{html.escape(heading)}</h3>{n}</div>'
            f'<div class="pv-body"></div></div>')

# ----------------------------------------------------------------------------- builders
def build_cia(records, project, outdir):
    stale_days = int((project.get("training") or {}).get("stale_days", DEFAULT_STALE_DAYS))
    data = normalize_cia(records, stale_days=stale_days)
    brand = {**DEFAULT_BRAND, **project.get("brand", {})}
    pname = project.get("project_name", "Project")
    areas = len({d["functional_area"] for d in data if d["functional_area"]})

    # what optional layers does the data / config support?
    tom_shifts = project.get("tom_shifts") or []
    benefits = project.get("benefits") or []
    has_tom = any(d["tom_shift"] for d in data) or bool(tom_shifts)
    has_value = any(d["value_levers"] for d in data) or bool(benefits)
    has_training = any(d["training_modality"] or d["estimated_training_hours"] for d in data)
    has_status = any(d["last_reviewed"] or d["workshop_reference"] for d in data)

    # ---- xlsx (one row per impact; arrays + status flattened) ----
    cols = [("id", "ID", 6), ("title", "Title", 40), ("functional_area", "Functional Area", 24),
            ("business_unit", "Business Unit", 16), ("l1_process", "L1 Process", 20),
            ("l2_process", "L2 Process", 22), ("l3_process", "L3 Process", 22),
            ("current_state", "Current State", 46), ("future_state", "Future State", 46),
            ("future_state_status", "Future State Status", 18),
            ("process_change", "Process Change", 40), ("roles_impacted", "Roles Impacted", 30),
            ("severity", "Severity", 10), ("complexity", "Complexity", 10),
            ("severity_label", "Severity Label", 14), ("impact_score", "Impact Score", 12),
            ("priority", "Priority", 12), ("sentiment", "Sentiment", 14),
            ("change_considerations", "Change Considerations", 46), ("mitigation", "Mitigation", 36),
            ("dim", "Change Dimensions", 34), ("primary_dimension", "Primary Change Dimension", 22),
            ("tom_shift", "TOM Shift", 26), ("value_levers", "Value Levers", 30),
            ("benefit_at_risk", "Benefit at Risk", 28),
            ("training_modality", "Training Modality", 24), ("estimated_training_hours", "Est. Training Hours", 14),
            ("access_constraints", "Access Constraints", 26),
            ("evidence_source", "Evidence Source", 26), ("workshop_reference", "Workshop Reference", 24),
            ("last_reviewed", "Last Reviewed", 14), ("scope", "Scope", 22)]
    write_xlsx(os.path.join(outdir, f"{pname} CIA.xlsx"), "CIA Register", cols, data)

    # ---- register columns / filters (add status, tom, levers only when data exists) ----
    columns = [{"key": "id", "label": "ID"}, {"key": "functional_area", "label": "Area"},
               {"key": "title", "label": "Title", "clip": 64},
               {"key": "severity", "label": "Sev"}, {"key": "priority", "label": "Priority", "kind": "pri"}]
    if has_status or has_tom or has_value:
        columns.append({"key": "future_state_status", "label": "Status", "kind": "status"})
    columns.append({"key": "dim", "label": "Change Dimensions", "kind": "dims"})

    filters = [{"key": "functional_area", "label": "All Areas"}, {"key": "business_unit", "label": "All BUs"},
               {"key": "priority", "label": "All Priorities"}, {"key": "primary_dimension", "label": "All Primary Dims"},
               {"key": "sentiment", "label": "All Sentiment"}]
    if has_status or has_tom or has_value:
        filters.append({"key": "future_state_status", "label": "All Status"})
    if has_tom: filters.append({"key": "tom_shift", "label": "All TOM Shifts"})
    if has_value: filters.append({"key": "value_levers", "label": "All Value Levers"})
    filters.append({"key": "scope", "label": "All Scope"})

    detail = [{"type": "currentfuture"}, {"label": "Process Change", "key": "process_change"},
              {"type": "dims", "label": "Change Dimensions"}, {"type": "roles", "label": "Roles Impacted"},
              {"label": "Change Considerations", "key": "change_considerations"},
              {"label": "Mitigation", "key": "mitigation"}, {"label": "TOM Shift", "key": "tom_shift"},
              {"type": "levers", "label": "Value Levers"}, {"label": "Benefit at Risk", "key": "benefit_at_risk"},
              {"type": "mods", "label": "Training Modality"}, {"label": "Est. Training Hours", "key": "estimated_training_hours"},
              {"label": "Access Constraints", "key": "access_constraints"},
              {"label": "Evidence Source", "key": "evidence_source"}, {"label": "Workshop Reference", "key": "workshop_reference"},
              {"label": "Last Reviewed", "key": "last_reviewed"}, {"label": "Scope", "key": "scope"}]

    tabs = [{"label": "Impact Register", "view": "v-table"}]
    pivots = []
    extra = []

    if has_tom:
        tabs.append({"label": "TOM Shifts", "view": "v-tom"})
        extra.append(pivot_view_html("v-tom", "Target Operating Model Shifts",
                                     "Each tile is an operating-model transition. Click to drill into its impacts; ⎙ prints a one-pager."))
        tom_meta = {}
        for t in tom_shifts:
            if isinstance(t, dict):
                tom_meta[t.get("name", "")] = {"description": t.get("description", ""), "today": t.get("today", ""),
                                               "future": t.get("future", ""), "target_metric": ""}
        pivots.append({"view": "v-tom", "groupBy": "tom_shift", "layout": "tiles", "drill": "tom_shift",
                       "print": True, "meta": tom_meta,
                       "metrics": [{"label": "Impacts", "kind": "count"},
                                   {"label": "Discussed", "kind": "countif", "test": {"key": "future_state_status", "startsWith": "Discussed"}},
                                   {"label": "Assumed", "kind": "countif", "test": {"key": "future_state_status", "eq": "Assumed"}},
                                   {"label": "Value Levers", "kind": "distinct", "key": "value_levers"},
                                   {"label": "Roles", "kind": "distinct", "key": "roles_impacted"}]})

    if has_value:
        tabs.append({"label": "Value Realization", "view": "v-value"})
        extra.append(pivot_view_html("v-value", "Value Realization",
                                     "Mapped impacts per value lever, by future-state status. Non-monetary. Click a lever to drill in."))
        ben_meta = {}; ben_vocab = []
        for b in benefits:
            if isinstance(b, dict):
                nm = b.get("benefit", "")
                ben_meta[nm] = {"description": b.get("description", ""), "target_metric": b.get("target_metric", "")}
                ben_vocab.append(nm)
        pivots.append({"view": "v-value", "explode": "value_levers", "layout": "tiles", "drill": "value_levers",
                       "meta": ben_meta, "vocab": ben_vocab,
                       "metrics": [{"label": "Impacts", "kind": "count"},
                                   {"label": "Confirmed", "kind": "countif", "test": {"key": "future_state_status", "eq": "Discussed-Confirmed"}},
                                   {"label": "Pending", "kind": "countif", "test": {"key": "future_state_status", "eq": "Discussed-Pending"}},
                                   {"label": "Assumed", "kind": "countif", "test": {"key": "future_state_status", "eq": "Assumed"}},
                                   {"label": "At Risk", "kind": "countif", "test": {"all": [{"key": "sentiment", "inSet": ["Mixed", "Negative"]}, {"key": "complexity", "gte": 4}]}}]})

    if has_training:
        tabs.append({"label": "Training Needs", "view": "v-training"})
        extra.append(pivot_view_html("v-training", "Training Needs",
                                     "Per audience (role): impacts, estimated hours, and modality mix. Click a row to search the register."))
        pivots.append({"view": "v-training", "roles": True, "layout": "rows", "drillSearch": True, "groupLabel": "Audience / Role",
                       "metrics": [{"label": "Impacts", "kind": "count"},
                                   {"label": "Est. Hours", "kind": "sum", "key": "estimated_training_hours"},
                                   {"label": "Modalities", "kind": "distinct", "key": "training_modality"},
                                   {"label": "High Complexity", "kind": "countif", "test": {"key": "complexity", "gte": 4}}]})

    if has_status:
        columns_extra_added = False
        # surface review columns on the register so the Workshop Status drill shows them
        for col in [{"key": "workshop_reference", "label": "Workshop"}, {"key": "last_reviewed", "label": "Last Reviewed"},
                    {"key": "stale", "label": "Stale", "kind": "stale"}]:
            columns.append(col); columns_extra_added = True
        filters.append({"key": "future_state_status", "label": "All Status"}) if not any(f["key"] == "future_state_status" for f in filters) else None
        tabs.append({"label": "Workshop Status", "view": "v-status"})
        extra.append(pivot_view_html("v-status", "Workshop / Future-State Status",
                                     f"Impacts by future-state status, with staleness (> {stale_days} days since last review)."))
        pivots.append({"view": "v-status", "groupBy": "future_state_status", "layout": "tiles", "drill": "future_state_status",
                       "colorBy": STATUS_COLORS,
                       "metrics": [{"label": "Impacts", "kind": "count"},
                                   {"label": "Stale", "kind": "countif", "test": {"key": "stale", "truthy": True}}]})

    config = {"idKey": "id", "dimColors": DIM_COLORS, "tabs": tabs, "columns": columns,
              "filters": filters,
              "stats": [{"label": "Impacts", "type": "count"}, {"label": "Functional Areas", "type": "distinct", "key": "functional_area"},
                        {"label": "Critical", "countifKey": "priority", "countifVal": "Critical"},
                        {"label": "High", "countifKey": "priority", "countifVal": "High"}],
              "detail": detail, "pivots": pivots}

    render_html(os.path.join(outdir, f"{pname} CIA Dashboard.html"),
                title=f"{pname} — Change Impact Assessment",
                subtitle=f"{len(data)} impacts across {areas} functional areas",
                data=data, config=config, brand=brand, extra_views="\n".join(extra))
    return len(data), areas

def write_readme(outdir, project):
    pname = project.get("project_name", "Project")
    files = sorted(f for f in os.listdir(outdir)
                   if f.startswith(pname) and "CIA" in f and (f.endswith(".xlsx") or f.endswith(".html")))
    lines = [f"# {pname} \u2014 Change Impact Assessment", "",
             "Generated by the `cia-builder` skill from current/future-state transcripts.", "",
             "## Deliverables", ""]
    for f in files:
        what = ("interactive dashboard \u2014 open in a browser; filter, sort, click a row for full detail"
                if f.endswith(".html") else "data workbook \u2014 one row per impact, all fields")
        lines.append(f"- **{f}** \u2014 {what}")
    lines += ["", "## CIA \u2014 Change Impact Assessment",
              "Each impact is a current->future change to how people work, scored on severity/complexity, "
              "classified across the 6 MECE Change Dimensions (with one Primary), and given a disposition "
              "(In Scope / Out of Scope / Gap). Where the records carry them, the dashboard also surfaces "
              "**TOM Shifts**, **Value Realization** (non-monetary), **Training Needs**, and **Workshop / "
              "Future-State status** as extra tabs.", "",
              "## Methodology", "See the `cia-builder` skill references (CIA_SCHEMA, METHODOLOGY, UPDATE_LOOP)."]
    open(os.path.join(outdir, "README.md"), "w", encoding="utf-8").write("\n".join(lines) + "\n")

# ----------------------------------------------------------------------------- cli
def _guard(outdir, names, force):
    """Refuse to overwrite existing deliverables unless --force is set."""
    import os, sys
    existing = [n for n in names if os.path.exists(os.path.join(outdir, n))]
    if existing and not force:
        sys.exit("Refusing to overwrite existing output(s) in %r:\n  %s\n"
                 "Re-run with --force to overwrite, or choose a different --outdir."
                 % (outdir, "\n  ".join(existing)))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("kind", choices=["cia", "readme"])
    ap.add_argument("--records"); ap.add_argument("--config", required=True)
    ap.add_argument("--outdir", default=".")
    ap.add_argument("--force", action="store_true", help="overwrite existing outputs in --outdir")
    a = ap.parse_args()
    os.makedirs(a.outdir, exist_ok=True)
    project = json.load(open(a.config, encoding="utf-8"))
    if a.kind == "cia":
        _pn = project.get("project_name", "Project")
        _guard(a.outdir, [f"{_pn} CIA.xlsx", f"{_pn} CIA Dashboard.html"], a.force)
    if a.kind == "readme":
        write_readme(a.outdir, project); print(f"README.md -> {a.outdir}"); return
    if not a.records:
        sys.exit("--records is required for cia")
    records = json.load(open(a.records, encoding="utf-8"))
    n, areas = build_cia(records, project, a.outdir); print(f"CIA: {n} impacts / {areas} areas -> {a.outdir}")

if __name__ == "__main__":
    main()
