#!/usr/bin/env python3
"""
training-rollout-builder — render a training rollout & measurement plan as an
xlsx workbook + a self-contained HTML dashboard, for ANY project, from
rollout_plan.json (see references/ROLLOUT_SCHEMA.md).

Timeline anchors BACKWARD from go-live using week offsets — no fabricated dates.

Usage:
  python3 rollout_render.py --records rollout_plan.json --config project.json --outdir OUT [--force]

Only stdlib + openpyxl. Output HTML has no external dependencies.
"""
import argparse, html, json, os, sys

DEFAULT_BRAND = {"navy": "#162B75", "magenta": "#EE2C81",
                 "font": "Calibri, system-ui, sans-serif", "footer": "Confidential"}
WAVE_COLORS = ["#EE2C81", "#162B75", "#04A577", "#F08301", "#7A4FBE", "#FF533C", "#1d6fb8"]
KIRK_COLORS = {"L1": "#6b7280", "L2": "#1d6fb8", "L3": "#F08301", "L4": "#04A577"}

def fmt(v):
    if isinstance(v, bool): return "Yes" if v else ""
    if isinstance(v, (list, tuple)): return ", ".join(str(x) for x in v)
    return v

def window_label(w):
    s, e = w.get("start_weeks_before_golive"), w.get("end_weeks_before_golive")
    if s is None: return ""
    return f"GL−{s} to " + ("GL" if not e else f"GL−{e}") + " wks"

def style_sheet(ws, columns, navy):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    thin = Side(style="thin", color="D9D9D9"); border = Border(thin, thin, thin, thin)
    for ci, _ in enumerate(columns, 1):
        c = ws.cell(1, ci); c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=navy.lstrip("#"))
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = border
    ws.freeze_panes = "B2"; ws.row_dimensions[1].height = 26
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
    for ci, col in enumerate(columns, 1):
        L = get_column_letter(ci); ws.column_dimensions[L].width = col[2] if len(col) > 2 else 22
        for cell in ws[L][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top"); cell.border = border
            cell.font = Font(size=10)

def write_xlsx(path, plan, navy):
    from openpyxl import Workbook
    wb = Workbook()
    wcols = [("id", "Wave", 8), ("name", "Name", 34), ("go_live_ref", "Go-Live", 10),
             ("audiences", "Audiences", 32), ("business_units", "Business Units", 20),
             ("locations", "Locations", 24), ("modules", "Modules", 28),
             ("delivery_mode", "Delivery Mode", 22), ("headcount", "Headcount", 24),
             ("window", "Window (backward from GL)", 20), ("sessions", "Sessions", 50), ("notes", "Notes", 40)]
    ws = wb.active; ws.title = "Waves & Sessions"
    ws.append([c[1] for c in wcols])
    for w in plan.get("waves", []):
        sess = "\n".join(f"{s.get('module','')} — {s.get('format','')} — {s.get('audience','')}"
                         + (f" ({s.get('notes')})" if s.get("notes") else "")
                         for s in w.get("sessions", []))
        ws.append([w.get("id"), w.get("name"), w.get("go_live_ref"), fmt(w.get("audiences")),
                   fmt(w.get("business_units")), fmt(w.get("locations")), fmt(w.get("modules")),
                   w.get("delivery_mode"), w.get("headcount"), window_label(w), sess, w.get("notes", "")])
    style_sheet(ws, wcols, navy)

    rcols = [("audience", "Audience", 34), ("criteria", "Readiness Criteria", 60), ("gate", "Gate (go/no-go)", 50)]
    ws2 = wb.create_sheet("Readiness")
    ws2.append([c[1] for c in rcols])
    for r in plan.get("readiness_criteria", []):
        ws2.append([r.get("audience"), "\n".join(r.get("criteria", [])), r.get("gate", "")])
    style_sheet(ws2, rcols, navy)

    mcols = [("level", "Level", 8), ("label", "Label", 12), ("metric", "Metric", 44),
             ("method", "Method", 44), ("owner", "Owner", 18), ("timing", "Timing", 20),
             ("value_lever", "Value Lever (CIA)", 30)]
    ws3 = wb.create_sheet("Measurement")
    ws3.append([c[1] for c in mcols])
    for m in plan.get("measurement", []):
        ws3.append([m.get(c[0], "") for c in mcols])
    # reinforcement schedule appended below measurement for one-workbook completeness
    ws3.append([]); ws3.append(["", "Reinforcement / sustainment schedule"])
    for r in plan.get("reinforcement", []):
        ws3.append(["", r.get("timing", ""), r.get("activity", ""), "", r.get("owner", ""), "", ""])
    style_sheet(ws3, mcols, navy)
    wb.save(path)

# ----------------------------------------------------------------------------- html
CSS = r"""
:root{--navy:%%NAVY%%;--mag:%%MAG%%;--ink:#1d2433;--mut:#6b7280;--line:#e5e7eb;--bg:#f5f6f8}
*{box-sizing:border-box}body{margin:0;font-family:%%FONT%%;color:var(--ink);background:var(--bg);font-size:14px}
header{background:linear-gradient(135deg,var(--navy),#0E2841);color:#fff;padding:18px 28px}
header h1{margin:0;font-size:1.15rem;font-weight:700}header .sub{opacity:.85;font-size:.85rem;margin-top:3px}
.wrap{max-width:1300px;margin:0 auto;padding:18px 28px 80px}
.stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:14px;margin:16px 0}
.card{background:#fff;border:1px solid var(--line);border-radius:10px;padding:14px 16px}
.card .v{font-size:1.6rem;font-weight:800;color:var(--navy)}.card .l{font-size:.72rem;color:var(--mut);text-transform:uppercase;letter-spacing:.04em;margin-top:2px}
.panel{background:#fff;border:1px solid var(--line);border-radius:10px;padding:16px 20px;margin:18px 0}
.panel h2{color:var(--navy);font-size:1.02rem;margin:0 0 10px}
.tl{position:relative;margin:14px 0 6px;overflow-x:auto}
.tl-grid{position:relative;min-width:640px}
.tl-axis{display:flex;border-bottom:2px solid var(--navy);margin-left:230px;position:relative}
.tl-axis span{flex:1;font-size:.68rem;color:var(--mut);text-align:right;padding:2px 4px;border-right:1px dashed #d8dce4}
.tl-row{display:flex;align-items:center;margin:8px 0}
.tl-name{width:230px;flex:none;font-size:.78rem;font-weight:700;color:var(--navy);padding-right:10px}
.tl-track{position:relative;flex:1;height:26px;background:#eef0f4;border-radius:6px}
.tl-bar{position:absolute;top:0;height:26px;border-radius:6px;color:#fff;font-size:.68rem;font-weight:700;display:flex;align-items:center;padding:0 8px;white-space:nowrap;overflow:hidden}
.gl-line{position:absolute;top:0;bottom:0;right:0;width:3px;background:var(--mag)}
.gl-note{font-size:.72rem;color:var(--mag);font-weight:800;text-align:right;margin-left:230px}
.tl-h{color:var(--navy);font-size:.9rem;margin:16px 0 2px}
.defer{border:1px solid #eeddc2;border-left:5px solid #B45309;background:#fffaf1;border-radius:8px;padding:10px 14px;margin:8px 0;font-size:.82rem}
.defer b{color:#B45309}
.wavecard{border:1px solid var(--line);border-left:5px solid var(--navy);border-radius:8px;padding:10px 14px;margin:8px 0;font-size:.82rem}
.wavecard b.wid{color:var(--navy)}
.chip{display:inline-block;padding:2px 9px;border-radius:11px;background:#5a6b86;color:#fff;font-size:.66rem;font-weight:600;margin:1px 3px 1px 0;white-space:nowrap}
.chip.mod{background:#0E2841}.chip.lever{background:#7A4FBE}
table{width:100%;border-collapse:collapse}
th{background:var(--navy);color:#fff;text-align:left;padding:8px 10px;font-size:.72rem;white-space:nowrap}
td{padding:8px 10px;border-top:1px solid var(--line);font-size:.8rem;vertical-align:top}
.klvl{display:inline-block;width:34px;text-align:center;padding:3px 0;border-radius:6px;color:#fff;font-weight:800;font-size:.72rem}
ul.crit{margin:4px 0;padding-left:18px}ul.crit li{margin:3px 0}
.gate{background:#f6f8fc;border:1px solid var(--line);border-radius:7px;padding:6px 10px;font-size:.78rem;margin-top:4px}
.risk{font-size:.82rem;margin:5px 0}
footer{text-align:center;color:var(--mut);font-size:.72rem;padding:18px}
"""

def esc(s): return html.escape(str(s if s is not None else ""))

def chips(items, cls=""):
    return "".join(f'<span class="chip {cls}">{esc(x)}</span>' for x in (items or []))

def _gl_note(gl):
    gl_label = gl.get("label", "Go-live")
    gl_date = f' — {esc(gl["date"])}' if gl.get("date") else " — date TBC (offsets only, no fabricated dates)"
    return f'<div class="gl-note">▲ {esc(gl_label)}{gl_date}</div>'

def _single_gl_section(gl, waves, color_offset=0, heading=""):
    """One backward-from-GL timeline for the waves anchored to a single go-live."""
    span = max([w.get("start_weeks_before_golive", 0) or 0 for w in waves] + [1])
    axis = "".join(f"<span>{('GL−' + str(span - i - 1)) if span - i - 1 else 'GL'}</span>" for i in range(span))
    rows = []
    for i, w in enumerate(waves):
        s = w.get("start_weeks_before_golive", 0) or 0
        e = w.get("end_weeks_before_golive", 0) or 0
        left = (span - s) / span * 100
        width = max((s - e) / span * 100, 4)
        color = WAVE_COLORS[(color_offset + i) % len(WAVE_COLORS)]
        rows.append(f'<div class="tl-row"><div class="tl-name">{esc(w.get("id"))} · {esc(w.get("name"))}</div>'
                    f'<div class="tl-track"><div class="tl-bar" style="left:{left:.1f}%;width:{width:.1f}%;background:{color}" '
                    f'title="{esc(window_label(w))}">{esc(window_label(w))}</div><div class="gl-line"></div></div></div>')
    head = f'<h3 class="tl-h">{esc(heading)}</h3>' if heading else ""
    return f'{head}<div class="tl"><div class="tl-grid"><div class="tl-axis">{axis}</div>{"".join(rows)}{_gl_note(gl)}</div></div>'

def _absolute_timeline(gls, waves):
    """All go-lives carry offset_weeks: plot every wave on one absolute week axis
    (week 0 = first go-live) with a marker per go-live."""
    off = {g.get("id"): float(g.get("offset_weeks") or 0) for g in gls}
    starts, ends = [], []
    for w in waves:
        o = off.get(w.get("go_live_ref"), 0)
        starts.append(o - (w.get("start_weeks_before_golive", 0) or 0))
        ends.append(o - (w.get("end_weeks_before_golive", 0) or 0))
    lo = min(starts + [0]); hi = max(ends + list(off.values()) + [1])
    total = hi - lo or 1
    axis = "".join(f"<span>{int(lo + i + 1):+d}</span>" for i in range(int(total)))
    def pos(x): return (x - lo) / total * 100
    marks = "".join(f'<div class="gl-line" style="left:calc({pos(off.get(g.get("id"), 0)):.1f}% - 1px);right:auto"></div>'
                    for g in gls)
    rows = []
    for i, w in enumerate(waves):
        o = off.get(w.get("go_live_ref"), 0)
        s = o - (w.get("start_weeks_before_golive", 0) or 0)
        e = o - (w.get("end_weeks_before_golive", 0) or 0)
        color = WAVE_COLORS[i % len(WAVE_COLORS)]
        lbl = f'{esc(w.get("go_live_ref"))} {esc(window_label(w))}'
        rows.append(f'<div class="tl-row"><div class="tl-name">{esc(w.get("id"))} · {esc(w.get("name"))}</div>'
                    f'<div class="tl-track"><div class="tl-bar" style="left:{pos(s):.1f}%;width:{max((e - s) / total * 100, 4):.1f}%;background:{color}" '
                    f'title="{lbl}">{lbl}</div>{marks}</div></div>')
    notes = "".join(
        f'<div class="gl-note">▲ {esc(g.get("label", g.get("id")))} — week {int(off.get(g.get("id"), 0)):+d}'
        + (f' — {esc(g["date"])}' if g.get("date") else " — date TBC (offsets only, no fabricated dates)") + "</div>"
        for g in gls)
    return (f'<div class="tl"><div class="tl-grid"><div class="tl-axis">{axis}</div>{"".join(rows)}{notes}'
            f'<div class="gl-note" style="color:var(--mut);font-weight:400">Axis: weeks relative to the first go-live (week 0)</div></div></div>')

def timeline(plan):
    waves = plan.get("waves", [])
    gls = plan.get("go_lives", [{}])
    refs = {w.get("go_live_ref") for w in waves}
    used = [g for g in gls if g.get("id") in refs] or gls[:1]
    if len(used) <= 1:
        return _single_gl_section(used[0] if used else {}, waves)
    if all(isinstance(g.get("offset_weeks"), (int, float)) for g in used):
        return _absolute_timeline(used, waves)
    # No offsets: never collapse onto go_lives[0] — one section per go-live.
    out, coff = [], 0
    for g in used:
        gw = [w for w in waves if w.get("go_live_ref") == g.get("id")]
        out.append(_single_gl_section(g, gw, color_offset=coff, heading=f'{g.get("id")} — {g.get("label", "")}'))
        coff += len(gw)
    return "".join(out)

def write_html(path, plan, project, brand):
    pname = project.get("project_name", "Project")
    waves = plan.get("waves", [])
    body = []
    stats = [("Waves", len(waves)),
             ("Modules Deployed", len({m for w in waves for m in w.get("modules", [])})),
             ("Audiences", len({a for w in waves for a in w.get("audiences", [])})),
             ("Readiness Gates", len(plan.get("readiness_criteria", []))),
             ("Measurement Levels", len({m.get("level") for m in plan.get("measurement", [])}))]
    body.append('<div class="stats">' + "".join(
        f'<div class="card"><div class="v">{v}</div><div class="l">{esc(l)}</div></div>' for l, v in stats) + "</div>")

    body.append(f'<div class="panel"><h2>Wave timeline — anchored backward from go-live (weeks)</h2>{timeline(plan)}'
                + "".join(
        f'<div class="wavecard" style="border-left-color:{WAVE_COLORS[i % len(WAVE_COLORS)]}">'
        f'<b class="wid">{esc(w.get("id"))} — {esc(w.get("name"))}</b> · {esc(window_label(w))} · {esc(w.get("delivery_mode",""))}'
        f' · headcount: {esc(w.get("headcount","—"))}<br>Audiences: {esc(", ".join(w.get("audiences", [])))}'
        f' · BUs: {esc(", ".join(w.get("business_units", [])))}<br>Modules: {chips(w.get("modules"), "mod")}'
        + (f'<br><i>{esc(w.get("notes"))}</i>' if w.get("notes") else "") + "</div>"
        for i, w in enumerate(waves)) + "</div>")

    su = plan.get("super_user_plan") or {}
    if su:
        body.append('<div class="panel"><h2>Super-user / champion enablement</h2>' + "".join(
            f'<p style="font-size:.84rem;margin:5px 0"><b>{esc(k.replace("_", " ").title())}:</b> {esc(v)}</p>'
            for k, v in su.items()) + "</div>")

    rc = plan.get("readiness_criteria", [])
    if rc:
        body.append('<div class="panel"><h2>Readiness gates (go/no-go per audience)</h2>' + "".join(
            f'<div class="wavecard"><b class="wid">{esc(r.get("audience"))}</b>'
            f'<ul class="crit">{"".join(f"<li>{esc(c)}</li>" for c in r.get("criteria", []))}</ul>'
            f'<div class="gate"><b>Gate:</b> {esc(r.get("gate",""))}</div></div>' for r in rc) + "</div>")

    pg, sg = plan.get("program_gates", []), plan.get("site_gates", [])
    if pg or sg:
        parts = []
        if pg:
            parts.append('<h3 class="tl-h">Program-level gates</h3>' + "".join(
                f'<div class="wavecard" style="border-left-color:#B45309"><b class="wid">{esc(g.get("gate",""))}</b>'
                + (f'<br>Source: {esc(g.get("source_need"))}' if g.get("source_need") else "")
                + (f' · Blocks: {esc(", ".join(g.get("blocks", [])))}' if g.get("blocks") else "") + "</div>"
                for g in pg))
        if sg:
            parts.append('<h3 class="tl-h">Site-level gates</h3>' + "".join(
                f'<div class="wavecard" style="border-left-color:#B45309"><b class="wid">{esc(s.get("site",""))}</b>'
                f'<ul class="crit">{"".join(f"<li>{esc(c)}</li>" for c in s.get("gates", []))}</ul>'
                + (f'<div class="gate">Source: {esc(s.get("source_need"))}</div>' if s.get("source_need") else "") + "</div>"
                for s in sg))
        body.append('<div class="panel"><h2>Program & site gates (non-training preconditions)</h2>' + "".join(parts) + "</div>")

    dm = plan.get("deferred_modules", [])
    if dm:
        body.append('<div class="panel"><h2>Not scheduled — gated</h2>' + "".join(
            f'<div class="defer"><b>{esc(d.get("module",""))}</b> — {esc(d.get("reason",""))}</div>' for d in dm)
            + '<p style="font-size:.78rem;color:#6b7280;margin:6px 0 0">These modules are deliberately excluded from all waves until their gates clear — not omissions.</p></div>')

    ms = plan.get("measurement", [])
    if ms:
        rows = "".join(
            f'<tr><td><span class="klvl" style="background:{KIRK_COLORS.get(m.get("level"), "#888")}">{esc(m.get("level"))}</span> '
            f'{esc(m.get("label",""))}</td><td>{esc(m.get("metric",""))}</td><td>{esc(m.get("method",""))}</td>'
            f'<td>{esc(m.get("owner",""))}</td><td>{esc(m.get("timing",""))}</td>'
            f'<td>{chips([x.strip() for x in str(m.get("value_lever","")).split(";") if x.strip()], "lever") or "—"}</td></tr>'
            for m in ms)
        body.append('<div class="panel"><h2>Measurement — Kirkpatrick L1–L4, tied to CIA value levers</h2>'
                    '<div style="overflow-x:auto"><table><thead><tr><th>Level</th><th>Metric</th><th>Method</th>'
                    f'<th>Owner</th><th>Timing</th><th>Value Lever</th></tr></thead><tbody>{rows}</tbody></table></div></div>')

    rf = plan.get("reinforcement", [])
    if rf:
        rows = "".join(f'<tr><td style="white-space:nowrap;font-weight:700;color:var(--navy)">{esc(r.get("timing"))}</td>'
                       f'<td>{esc(r.get("activity",""))}</td><td>{esc(r.get("audience",""))}</td><td>{esc(r.get("owner",""))}</td></tr>'
                       for r in rf)
        body.append('<div class="panel"><h2>Reinforcement / sustainment schedule</h2><div style="overflow-x:auto">'
                    f'<table><thead><tr><th>Timing</th><th>Activity</th><th>Audience</th><th>Owner</th></tr></thead><tbody>{rows}</tbody></table></div></div>')

    ct = plan.get("comms_touchpoints", [])
    if ct:
        rows = "".join(f'<tr><td style="white-space:nowrap;font-weight:700;color:var(--navy)">{esc(c.get("timing"))}</td>'
                       f'<td>{esc(c.get("audience",""))}</td><td>{esc(c.get("message",""))}</td><td>{esc(c.get("channel",""))}</td></tr>'
                       for c in ct)
        body.append('<div class="panel"><h2>Learner comms touchpoints (training-specific)</h2><div style="overflow-x:auto">'
                    f'<table><thead><tr><th>Timing</th><th>Audience</th><th>Message</th><th>Channel</th></tr></thead><tbody>{rows}</tbody></table></div></div>')

    ra = plan.get("risks_assumptions", [])
    if ra:
        body.append('<div class="panel"><h2>Risks & assumptions</h2>' +
                    "".join(f'<div class="risk">• {esc(r)}</div>' for r in ra) + "</div>")

    doc = ('<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">'
           '<meta name="viewport" content="width=device-width, initial-scale=1.0">'
           f'<title>{esc(pname)} — Training Rollout</title><style>'
           + CSS.replace("%%NAVY%%", brand["navy"]).replace("%%MAG%%", brand["magenta"]).replace("%%FONT%%", brand["font"])
           + f'</style></head><body><header><h1>{esc(pname)} — Training Rollout & Measurement</h1>'
           f'<div class="sub">{len(waves)} waves scheduled backward from go-live · Kirkpatrick L1–L4 · super-users first</div></header>'
           f'<div class="wrap">{"".join(body)}</div><footer>{esc(brand["footer"])}</footer></body></html>')
    open(path, "w", encoding="utf-8").write(doc)

# ----------------------------------------------------------------------------- cli
def _guard(outdir, names, force):
    existing = [n for n in names if os.path.exists(os.path.join(outdir, n))]
    if existing and not force:
        sys.exit("Refusing to overwrite existing output(s) in %r:\n  %s\n"
                 "Re-run with --force to overwrite, or choose a different --outdir."
                 % (outdir, "\n  ".join(existing)))

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--records", required=True); ap.add_argument("--config", required=True)
    ap.add_argument("--outdir", default=".")
    ap.add_argument("--force", action="store_true", help="overwrite existing outputs in --outdir")
    a = ap.parse_args()
    os.makedirs(a.outdir, exist_ok=True)
    project = json.load(open(a.config, encoding="utf-8"))
    brand = {**DEFAULT_BRAND, **project.get("brand", {})}
    pname = project.get("project_name", "Project")
    names = [f"{pname} Training Rollout.xlsx", f"{pname} Rollout Dashboard.html"]
    _guard(a.outdir, names, a.force)
    plan = json.load(open(a.records, encoding="utf-8"))
    write_xlsx(os.path.join(a.outdir, names[0]), plan, brand["navy"])
    write_html(os.path.join(a.outdir, names[1]), plan, project, brand)
    print(f"Rollout: {len(plan.get('waves', []))} waves -> {a.outdir}")

if __name__ == "__main__":
    main()
