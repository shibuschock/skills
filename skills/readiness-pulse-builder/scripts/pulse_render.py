#!/usr/bin/env python3
"""
readiness-pulse-builder — render (a) a Pulse Survey Guide xlsx from a survey
design, and (b) a self-contained Readiness Readout HTML dashboard from fielded
results, for ANY project. Claude writes the JSON inputs; this script renders.

Usage:
  python pulse_render.py design  --design pulse_design.json --config project.json --outdir OUT
  python pulse_render.py results --results pulse_results.json --design pulse_design.json --config project.json --outdir OUT

Rules baked in:
  - Reverse items flipped (scale_max + 1 - score) before averaging.
  - Segments with n < anonymity_threshold are suppressed and flagged, never scored.
  - Dimension score = mean of that dimension's likert question means; overall
    tile = n-weighted mean across reportable segments.
  - No fabrication: missing scores stay blank.

Only stdlib + openpyxl (xlsx only). HTML output has no external dependencies.
Render scaffolding conventions per cia-builder (canonical source of shared style).
"""
import argparse, html, json, os, sys

DEFAULT_BRAND = {"navy": "#162B75", "magenta": "#EE2C81",
                 "font": "Calibri, system-ui, sans-serif", "footer": "Confidential"}
RAG = {"green": "#04A577", "amber": "#F08301", "red": "#FF533C", "grey": "#9aa3b2"}
DEFAULT_THRESHOLD = 5


def load_json(path, what):
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        sys.exit(f"ERROR reading {what} ({path}): {e}")


def brand_of(cfg):
    b = dict(DEFAULT_BRAND)
    b.update(cfg.get("brand") or {})
    return b


def _guard(path, force):
    if os.path.exists(path) and not force:
        sys.exit(f"REFUSING to overwrite existing output: {path}\nPass --force to regenerate.")


def esc(s):
    return html.escape(str(s if s is not None else ""))


# --------------------------------------------------------------------- scoring
def mean_of(value):
    """Score value is a mean (number) or a distribution (list of counts, idx 0 = point 1)."""
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, (list, tuple)) and value and sum(value) > 0:
        total = sum(value)
        return sum((i + 1) * c for i, c in enumerate(value)) / total
    return None


def question_index(design):
    """id -> question dict, likert-scored only."""
    return {q["id"]: q for q in design.get("questions", [])
            if q.get("scale", "likert5") == "likert5"}


def adj_score(q, val, scale_max=5):
    m = mean_of(val)
    if m is None:
        return None
    return (scale_max + 1 - m) if q.get("reverse") else m


def dim_scores(segment, qidx, dims, scale_max=5):
    """Per-dimension mean for one segment. Missing -> None."""
    acc = {d: [] for d in dims}
    for qid, val in (segment.get("scores") or {}).items():
        q = qidx.get(qid)
        if not q:
            continue
        s = adj_score(q, val, scale_max)
        if s is not None and q.get("dimension") in acc:
            acc[q["dimension"]].append(s)
    return {d: (sum(v) / len(v) if v else None) for d, v in acc.items()}


def rag_color(score):
    if score is None:
        return RAG["grey"]
    if score >= 3.75:
        return RAG["green"]
    if score >= 3.0:
        return RAG["amber"]
    return RAG["red"]


def compute_wave(wave, qidx, dims, threshold, scale_max=5):
    """Returns dict with reportable segment rows, suppressed names, dim overall means."""
    rows, suppressed = [], []
    for seg in wave.get("segments") or []:
        n = seg.get("n") or 0
        if n < threshold:
            suppressed.append({"name": seg.get("name", "?"), "n": n})
            continue
        rows.append({"name": seg.get("name", "?"), "n": n,
                     "dims": dim_scores(seg, qidx, dims, scale_max)})
    overall = {}
    for d in dims:
        pairs = [(r["dims"][d], r["n"]) for r in rows if r["dims"][d] is not None]
        overall[d] = (sum(s * n for s, n in pairs) / sum(n for _, n in pairs)) if pairs else None
    total_n = sum(r["n"] for r in rows) + sum(s["n"] for s in suppressed)
    return {"label": wave.get("label", ""), "date": wave.get("date", ""),
            "invited": wave.get("invited"), "responses": total_n,
            "rows": rows, "suppressed": suppressed, "overall": overall,
            "themes": wave.get("open_text_themes") or []}


# --------------------------------------------------------------------- design xlsx
def render_design(design, cfg, outdir, force):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("openpyxl required for design mode: pip install openpyxl")

    project = cfg.get("project_name", "Project")
    navy = brand_of(cfg)["navy"].lstrip("#")
    path = os.path.join(outdir, f"{project} Pulse Survey Guide.xlsx")
    _guard(path, force)

    wb = Workbook()
    thin = Side(style="thin", color="D9D9D9")
    border = Border(thin, thin, thin, thin)

    def style_sheet(ws, widths):
        for ci, w in enumerate(widths, 1):
            L = get_column_letter(ci)
            ws.column_dimensions[L].width = w
            for cell in ws[L]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = border
                if cell.font.bold is not True:
                    cell.font = Font(size=10)

    def header_row(ws, labels):
        ws.append(labels)
        for ci in range(1, len(labels) + 1):
            c = ws.cell(ws.max_row, ci)
            c.font = Font(bold=True, color="FFFFFF", size=11)
            c.fill = PatternFill("solid", fgColor=navy)
            c.alignment = Alignment(wrap_text=True, vertical="center")

    # --- Questions sheet
    ws = wb.active
    ws.title = "Questions"
    header_row(ws, ["ID", "Dimension", "Question", "Scale", "Type", "Reverse-scored", "Use when"])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G1"
    for q in design.get("questions", []):
        ws.append([q.get("id", ""), q.get("dimension", ""), q.get("text", ""),
                   q.get("scale", "likert5"),
                   "Anchor" if q.get("anchor") else "Rotating",
                   "Yes" if q.get("reverse") else "",
                   q.get("use_when", "")])
    style_sheet(ws, [8, 20, 60, 10, 11, 14, 26])
    ws.row_dimensions[1].height = 24

    # --- Admin Plan sheet
    ws2 = wb.create_sheet("Admin Plan")
    sc = design.get("scale") or {}
    labels = sc.get("labels") or []
    def section(title):
        ws2.append([])
        ws2.append([title])
        c = ws2.cell(ws2.max_row, 1)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=navy)

    ws2.append([design.get("instrument_name", f"{project} Readiness Pulse")])
    ws2.cell(1, 1).font = Font(bold=True, size=13)
    section("Response scale")
    ws2.append([f"{sc.get('points', 5)}-point Likert",
                " | ".join(f"{i+1} = {l}" for i, l in enumerate(labels))])
    section("Audience segments (reporting level)")
    for s in design.get("segments", []):
        ws2.append([s])
    section("Cadence")
    ws2.append(["Phase", "Frequency", "Focus"])
    for c_ in design.get("cadence", []):
        ws2.append([c_.get("phase", ""), c_.get("frequency", ""), c_.get("focus", "")])
    section("Anonymity")
    thr = design.get("anonymity_threshold", DEFAULT_THRESHOLD)
    ws2.append([f"Minimum reporting n: {thr}. Segments with fewer than {thr} responses are never reported."])
    section("Comms plan")
    for k, label in [("announced_by", "Announced by"), ("distribution", "Distribution"),
                     ("reminders", "Reminders"), ("anonymity_statement", "Anonymity statement"),
                     ("results_feedback", "Results feedback (close the loop)")]:
        v = (design.get("comms_plan") or {}).get(k, "")
        if v:
            ws2.append([label, v])
    style_sheet(ws2, [30, 60, 40])

    wb.save(path)
    print(f"Wrote {path}")


# --------------------------------------------------------------------- results html
def render_results(results, design, cfg, outdir, force):
    project = cfg.get("project_name", "Project")
    b = brand_of(cfg)
    path = os.path.join(outdir, f"{project} Readiness Readout.html")
    _guard(path, force)

    dims = design.get("dimensions") or []
    qidx = question_index(design)
    scale_max = (design.get("scale") or {}).get("points", 5)
    threshold = results.get("anonymity_threshold",
                            design.get("anonymity_threshold", DEFAULT_THRESHOLD))
    waves = [compute_wave(w, qidx, dims, threshold, scale_max)
             for w in results.get("waves") or []]
    if not waves:
        sys.exit("No waves in results file.")
    latest, prev = waves[-1], (waves[-2] if len(waves) > 1 else None)

    def fmt(s):
        return f"{s:.2f}" if s is not None else "&ndash;"

    def delta_html(cur, old):
        if cur is None or old is None:
            return ""
        d = cur - old
        if abs(d) < 0.05:
            return '<span class="delta flat">&#8596; 0.00</span>'
        arrow, cls = ("&#9650;", "up") if d > 0 else ("&#9660;", "down")
        return f'<span class="delta {cls}">{arrow} {d:+.2f}</span>'

    # tiles: latest overall per dimension (+ delta vs prev)
    tiles = []
    for d in dims:
        cur = latest["overall"].get(d)
        old = prev["overall"].get(d) if prev else None
        tiles.append(
            f'<div class="card" style="border-top:4px solid {rag_color(cur)}">'
            f'<div class="v">{fmt(cur)}</div><div class="l">{esc(d)}</div>'
            f'{delta_html(cur, old) if prev else ""}</div>')

    # response-rate strip per wave
    rr = []
    for w in waves:
        rate = (f"{100.0 * w['responses'] / w['invited']:.0f}%"
                if w.get("invited") else "&ndash;")
        rr.append(f'<div class="card"><div class="v">{rate}</div>'
                  f'<div class="l">{esc(w["label"])} &middot; {esc(w["date"])}<br>'
                  f'{w["responses"]} responses'
                  f'{" of " + str(w["invited"]) + " invited" if w.get("invited") else ""}</div></div>')

    # heat table: dimension x segment (latest wave)
    head = "".join(f"<th>{esc(r['name'])}<br><span class='n'>n={r['n']}</span></th>"
                   for r in latest["rows"])
    body = []
    for d in dims:
        cells = []
        for r in latest["rows"]:
            s = r["dims"].get(d)
            txt = fmt(s)
            cells.append(f'<td style="background:{rag_color(s)}{"" if s is None else ""};'
                         f'color:#fff;font-weight:700;text-align:center">{txt}</td>')
        ov = latest["overall"].get(d)
        body.append(f"<tr><td class='dimname'>{esc(d)}</td>"
                    f"<td style='background:{rag_color(ov)};color:#fff;font-weight:800;"
                    f"text-align:center'>{fmt(ov)}</td>{''.join(cells)}</tr>")
    suppress_note = ""
    if latest["suppressed"]:
        names = ", ".join(f"{esc(s['name'])} (n={s['n']})" for s in latest["suppressed"])
        suppress_note = (f'<div class="suppress">&#9888; Suppressed below anonymity threshold '
                         f'(n &lt; {threshold}): {names} &mdash; scores withheld to protect anonymity.</div>')

    # wave-over-wave table (all waves, overall per dimension)
    trend = ""
    if len(waves) > 1:
        thead = "".join(f"<th>{esc(w['label'])}</th>" for w in waves)
        trows = []
        for d in dims:
            tds = []
            for i, w in enumerate(waves):
                cur = w["overall"].get(d)
                old = waves[i - 1]["overall"].get(d) if i else None
                tds.append(f"<td style='text-align:center'>{fmt(cur)} "
                           f"{delta_html(cur, old) if i else ''}</td>")
            trows.append(f"<tr><td class='dimname'>{esc(d)}</td>{''.join(tds)}</tr>")
        trend = (f"<h2>Wave-over-wave</h2><div class='tblwrap'><table><thead><tr><th>Dimension</th>"
                 f"{thead}</tr></thead><tbody>{''.join(trows)}</tbody></table></div>")

    # open-text themes (latest wave)
    themes = ""
    if latest["themes"]:
        items = "".join(
            f"<li><b>{esc(t.get('theme',''))}</b>"
            f"{' &mdash; ' + str(t['mentions']) + ' mentions' if t.get('mentions') else ''}"
            f"{('<br><em>&ldquo;' + esc(t['sample_quote']) + '&rdquo;</em>') if t.get('sample_quote') else ''}</li>"
            for t in latest["themes"])
        themes = f"<h2>Open-text themes &mdash; {esc(latest['label'])}</h2><ul class='themes'>{items}</ul>"

    page = f"""<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{esc(project)} Readiness Readout</title>
<style>
:root{{--navy:{b['navy']};--mag:{b['magenta']};--ink:#1d2433;--mut:#6b7280;--line:#e5e7eb;--bg:#f5f6f8}}
*{{box-sizing:border-box}}body{{margin:0;font-family:{b['font']};color:var(--ink);background:var(--bg);font-size:14px}}
header{{background:linear-gradient(135deg,var(--navy),#0E2841);color:#fff;padding:18px 28px}}
header h1{{margin:0;font-size:1.15rem;font-weight:700}}header .sub{{opacity:.85;font-size:.85rem;margin-top:3px}}
.wrap{{max-width:1300px;margin:0 auto;padding:18px 28px 60px}}
h2{{color:var(--navy);font-size:1.02rem;margin:26px 0 10px}}
.stats{{display:grid;grid-template-columns:repeat(auto-fit,minmax(160px,1fr));gap:14px;margin:14px 0}}
.card{{background:#fff;border:1px solid var(--line);border-radius:10px;padding:14px 16px}}
.card .v{{font-size:1.6rem;font-weight:800;color:var(--navy)}}
.card .l{{font-size:.72rem;color:var(--mut);text-transform:uppercase;letter-spacing:.04em;margin-top:2px;line-height:1.5}}
.delta{{display:inline-block;margin-top:6px;font-size:.74rem;font-weight:700;padding:2px 8px;border-radius:9px}}
.delta.up{{background:#e6f6f0;color:#036c4e}}.delta.down{{background:#fdecea;color:#b40020}}
.delta.flat{{background:#eef0f3;color:var(--mut)}}
.tblwrap{{overflow-x:auto}}
table{{width:100%;border-collapse:collapse;background:#fff;border:1px solid var(--line);border-radius:10px;overflow:hidden}}
thead th{{background:var(--navy);color:#fff;text-align:center;padding:9px 11px;font-size:.74rem;white-space:nowrap}}
thead th:first-child{{text-align:left}}thead .n{{font-weight:400;opacity:.8}}
tbody td{{padding:9px 11px;border-top:1px solid var(--line);font-size:.84rem}}
td.dimname{{font-weight:700;color:var(--navy);white-space:nowrap}}
.legend{{display:flex;gap:16px;margin:8px 0;font-size:.74rem;color:var(--mut);flex-wrap:wrap}}
.legend span{{display:inline-flex;align-items:center;gap:6px}}
.legend i{{display:inline-block;width:12px;height:12px;border-radius:3px}}
.suppress{{margin:10px 0;padding:10px 14px;background:#fff7ed;border:1px solid #F08301;border-radius:8px;font-size:.82rem}}
ul.themes{{background:#fff;border:1px solid var(--line);border-radius:10px;padding:14px 14px 14px 32px;margin:0}}
ul.themes li{{margin:8px 0;line-height:1.45}}
footer{{text-align:center;color:var(--mut);font-size:.72rem;padding:18px}}
</style></head><body>
<header><h1>{esc(project)} &mdash; Readiness Readout</h1>
<div class="sub">{esc(latest['label'])} &middot; {esc(latest['date'])} &middot; anonymity threshold n &ge; {threshold}</div></header>
<div class="wrap">
<h2>Response rate</h2><div class="stats">{''.join(rr)}</div>
<h2>Readiness by dimension &mdash; {esc(latest['label'])}</h2>
<div class="stats">{''.join(tiles)}</div>
<h2>Dimension &times; segment</h2>
<div class="legend"><span><i style="background:{RAG['green']}"></i>&ge; 3.75</span>
<span><i style="background:{RAG['amber']}"></i>3.00&ndash;3.74</span>
<span><i style="background:{RAG['red']}"></i>&lt; 3.00</span>
<span><i style="background:{RAG['grey']}"></i>no data</span></div>
<div class="tblwrap"><table><thead><tr><th>Dimension</th><th>Overall</th>{head}</tr></thead>
<tbody>{''.join(body)}</tbody></table></div>
{suppress_note}
{trend}
{themes}
</div><footer>{esc(b['footer'])} &middot; Generated by readiness-pulse-builder</footer>
</body></html>"""

    os.makedirs(outdir, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(page)
    print(f"Wrote {path}")
    if latest["suppressed"]:
        print(f"Suppressed segments (n < {threshold}): "
              + ", ".join(s["name"] for s in latest["suppressed"]))


# --------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Render pulse survey guide (xlsx) or readiness readout (html).")
    sub = ap.add_subparsers(dest="mode", required=True)
    d = sub.add_parser("design", help="Render the Pulse Survey Guide xlsx")
    d.add_argument("--design", required=True)
    d.add_argument("--config", required=True)
    d.add_argument("--outdir", default=".")
    d.add_argument("--force", action="store_true")
    r = sub.add_parser("results", help="Render the Readiness Readout html")
    r.add_argument("--results", required=True)
    r.add_argument("--design", required=True)
    r.add_argument("--config", required=True)
    r.add_argument("--outdir", default=".")
    r.add_argument("--force", action="store_true")
    a = ap.parse_args()

    cfg = load_json(a.config, "project config")
    design = load_json(a.design, "pulse design")
    os.makedirs(a.outdir, exist_ok=True)
    if a.mode == "design":
        render_design(design, cfg, a.outdir, a.force)
    else:
        results = load_json(a.results, "pulse results")
        render_results(results, design, cfg, a.outdir, a.force)


if __name__ == "__main__":
    main()
