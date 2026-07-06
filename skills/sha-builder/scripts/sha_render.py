#!/usr/bin/env python3
"""
sha-builder — render a Stakeholder Assessment (SHA) as an xlsx workbook + a
self-contained interactive HTML dashboard, for ANY project, from structured
records (extracted by Claude from raw transcripts/interviews).

Methodology baked in (portable OCM):
  - Influence x Interest (Mendelow) grid -> Engagement Strategy quadrant derived
    (Manage Closely / Keep Satisfied / Keep Informed / Monitor).

Usage:
  python3 sha_render.py sha    --records sha_records.json --config project.json --outdir OUT
  python3 sha_render.py readme --config project.json --outdir OUT

Records = a JSON array of objects (see references/SHA_SCHEMA.md).
Only stdlib + openpyxl. Cross-platform. Output HTML has no external dependencies.
"""
import argparse, datetime, html, json, os, re, sys

# ----------------------------------------------------------------------------- helpers
DIM_COLORS = {"Role/Accountability": "#162B75", "Ways of Working": "#04A577",
              "Skills/Capability": "#F08301", "Mindset/Culture": "#EE2C81",
              "Process/Policy": "#FF533C", "Data/Decision Inputs": "#7A4FBE"}
QUADS = ["Manage Closely", "Keep Satisfied", "Keep Informed", "Monitor"]
# Future-state status -> ribbon color (green confirmed / amber pending / grey assumed)
STATUS_COLORS = {"Discussed-Confirmed": "#04A577", "Discussed-Pending": "#F08301", "Assumed": "#9aa3b2"}
DEFAULT_BRAND = {"navy": "#162B75", "magenta": "#EE2C81", "orange": "#F08301",
                 "teal": "#04A577", "coral": "#FF533C", "font": "Calibri, system-ui, sans-serif",
                 "footer": "Confidential"}
DEFAULT_STALE_DAYS = 30

def g(rec, *keys, default=""):
    for k in keys:
        if k in rec and rec[k] not in (None, ""):
            return rec[k]
    return default

def as_int(v):
    try: return int(float(v))
    except Exception: return None

def as_list(v):
    if v is None or v == "": return []
    if isinstance(v, (list, tuple)): return [str(x).strip() for x in v if str(x).strip()]
    return [s.strip() for s in re.split(r"[;,]", str(v)) if s.strip()]

def engagement_quadrant(infl, inte):
    hi = lambda n: isinstance(n, (int, float)) and n >= 4
    if hi(infl) and hi(inte): return "Manage Closely"
    if hi(infl) and not hi(inte): return "Keep Satisfied"
    if not hi(infl) and hi(inte): return "Keep Informed"
    return "Monitor"

def days_since(date_str):
    if not date_str: return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d"):
        try:
            d = datetime.datetime.strptime(str(date_str)[:10], fmt).date()
            return (datetime.date.today() - d).days
        except Exception:
            continue
    return None

# ----------------------------------------------------------------------------- normalize
def normalize_sha(records):
    out = []
    for i, r in enumerate(records, 1):
        infl = as_int(g(r, "influence")); inte = as_int(g(r, "interest"))
        eng = g(r, "engagement_strategy")
        quad = engagement_quadrant(infl, inte) if (infl is not None and inte is not None) else ""
        if eng:  # ensure lead quadrant matches the grid
            tail = eng
            for q in QUADS:
                m = re.match(re.escape(q) + r"\s*[—\-:]*\s*", eng, re.I)
                if m: tail = eng[m.end():]; break
            eng = (quad + " — " + tail).strip() if (quad and tail.strip()) else (eng or quad)
        else:
            eng = quad
        out.append({
            "id": r.get("id", i),
            "group": g(r, "stakeholder_group", "group", "role"),
            "business_unit": g(r, "business_unit", "bu"),
            "location": g(r, "location"),
            "category": g(r, "category"),
            "description": g(r, "group_description", "description"),
            "influence": infl if infl is not None else "", "interest": inte if inte is not None else "",
            "quadrant": quad,
            "assessment_status": g(r, "assessment_status"),
            "influence_rationale": g(r, "influence_rationale"),
            "interest_rationale": g(r, "interest_rationale"),
            "impact_from_change": g(r, "impact_from_change"),
            "impact_rationale": g(r, "impact_rationale"),
            "decision_authority": g(r, "decision_authority"),
            "sentiment": g(r, "current_sentiment", "sentiment"),
            "pain_points": g(r, "key_pain_points", "pain_points"),
            "comms": g(r, "communication_preferences"),
            "engagement": eng,
            "training": g(r, "training_preferences"),
            "poc": g(r, "main_point_of_contact", "poc"),
            "champion": g(r, "change_champion"),
            "notes": g(r, "additional_notes"),
            "follow_up": g(r, "follow_up_needed"),
        })
    return out

# ----------------------------------------------------------------------------- xlsx
def write_xlsx(path, sheet_name, columns, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = sheet_name[:31]
    thin = Side(style="thin", color="D9D9D9"); border = Border(thin, thin, thin, thin)
    ws.append([c[1] for c in columns])
    for ci, _ in enumerate(columns, 1):
        c = ws.cell(1, ci); c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor="162B75"); c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = border
    ws.freeze_panes = "B2"; ws.row_dimensions[1].height = 26
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
    for r in rows:
        ws.append([fmt_cell(r.get(c[0], "")) for c in columns])
    for ci, col in enumerate(columns, 1):
        L = get_column_letter(ci); ws.column_dimensions[L].width = col[2] if len(col) > 2 else 22
        for cell in ws[L][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top"); cell.border = border
            cell.font = Font(size=10)
    wb.save(path)

def fmt_cell(v):
    if isinstance(v, bool):
        return "Yes" if v else ""
    if isinstance(v, dict):
        return ", ".join(k for k, val in v.items() if val)
    if isinstance(v, (list, tuple)):
        return ", ".join(str(x) for x in v)
    return v

# ----------------------------------------------------------------------------- HTML engine
TEMPLATE = r"""<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0"><title>%%TITLE%%</title>
<style>
:root{--navy:%%NAVY%%;--mag:%%MAG%%;--ink:#1d2433;--mut:#6b7280;--line:#e5e7eb;--bg:#f5f6f8}
*{box-sizing:border-box}body{margin:0;font-family:%%FONT%%;color:var(--ink);background:var(--bg);font-size:14px}
header{background:linear-gradient(135deg,var(--navy),#0E2841);color:#fff;padding:18px 28px}
header h1{margin:0;font-size:1.15rem;font-weight:700}header .sub{opacity:.85;font-size:.85rem;margin-top:3px}
.wrap{max-width:1500px;margin:0 auto;padding:18px 28px 80px}
.stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:14px;margin:16px 0}
.card{background:#fff;border:1px solid var(--line);border-radius:10px;padding:14px 16px}
.card .v{font-size:1.6rem;font-weight:800;color:var(--navy)}.card .l{font-size:.72rem;color:var(--mut);text-transform:uppercase;letter-spacing:.04em;margin-top:2px}
.bar{display:flex;flex-wrap:wrap;gap:8px;align-items:center;margin:10px 0}
.bar select,.bar input{padding:7px 10px;border:1px solid var(--line);border-radius:7px;font:inherit;background:#fff}
.bar input{flex:1;min-width:180px}.btn{background:var(--navy);color:#fff;border:none;border-radius:7px;padding:8px 14px;cursor:pointer;font:inherit}
.tabs{display:flex;gap:4px;border-bottom:2px solid var(--line);margin-top:8px;flex-wrap:wrap}
.tab{padding:8px 14px;cursor:pointer;border-radius:7px 7px 0 0;font-weight:600;color:var(--mut)}
.tab.active{color:var(--navy);background:#fff;border:1px solid var(--line);border-bottom:2px solid #fff;margin-bottom:-2px}
.view{display:none}.view.active{display:block}
.pv-head{display:flex;align-items:baseline;justify-content:space-between;margin:14px 0 6px}
.pv-head h3{color:var(--navy);margin:0}.pv-note{color:var(--mut);font-size:.78rem}
.tiles{display:grid;grid-template-columns:repeat(auto-fill,minmax(260px,1fr));gap:14px}
.ptile{background:#fff;border:1px solid var(--line);border-radius:12px;padding:14px 16px;cursor:pointer;transition:box-shadow .12s,transform .12s;border-top:4px solid var(--navy)}
.ptile:hover{box-shadow:0 6px 18px rgba(20,40,90,.10);transform:translateY(-1px)}
.ptile.pv-zero{opacity:.7;border-top-color:#c7ccd6}
.pt-h{font-weight:700;color:var(--navy);font-size:.98rem;display:flex;align-items:center;gap:8px;justify-content:space-between}
.pt-meta{color:var(--mut);font-size:.78rem;margin:5px 0 2px;line-height:1.4}
.pm-row{display:flex;flex-wrap:wrap;gap:14px;margin-top:10px}
.pm{min-width:54px}.pm-v{font-size:1.25rem;font-weight:800;color:var(--ink)}.pm-l{font-size:.64rem;color:var(--mut);text-transform:uppercase;letter-spacing:.03em}
.zero-flag{font-size:.6rem;background:#fdecea;color:#b40020;padding:2px 7px;border-radius:9px;font-weight:700;text-transform:uppercase}
.pt-print{background:none;border:1px solid var(--line);border-radius:6px;cursor:pointer;color:var(--mut);font-size:.8rem;padding:2px 7px}
.pt-print:hover{color:var(--navy);border-color:var(--navy)}
table{width:100%;border-collapse:collapse;background:#fff;border:1px solid var(--line);border-radius:10px;overflow:hidden}
thead th{position:sticky;top:0;background:var(--navy);color:#fff;text-align:left;padding:9px 11px;font-size:.74rem;cursor:pointer;white-space:nowrap}
tbody td{padding:9px 11px;border-top:1px solid var(--line);font-size:.82rem;vertical-align:top}
tbody tr.row,tbody tr.prow{cursor:pointer}tbody tr.row:hover,tbody tr.prow:hover{background:#f0f3fb}
.chip{display:inline-block;padding:2px 9px;border-radius:11px;color:#fff;font-size:.68rem;font-weight:600;margin:1px 2px 1px 0;white-space:nowrap}
.chip.prim{border:2px solid #111;font-weight:800}.chip.sec{opacity:.72}
.chip.lever{background:#7A4FBE}.chip.mod{background:#5a6b86}
.scope{display:inline-block;padding:2px 8px;border-radius:6px;background:#6b7280;color:#fff;font-size:.66rem;font-weight:700}
.statc{display:inline-block;padding:2px 8px;border-radius:6px;color:#fff;font-size:.66rem;font-weight:700;white-space:nowrap}
.stale-flag{display:inline-block;padding:1px 7px;border-radius:9px;background:#fdecea;color:#b40020;font-size:.62rem;font-weight:700}
.pri{font-weight:700}.pri.Critical{color:#b40020}.pri.High{color:#d35400}.pri.Medium{color:#1d6fb8}.pri.Low{color:#6b7280}
.detail{background:#fbfcfe}.detail .sec{margin:9px 0}.detail .lbl{font-size:.68rem;text-transform:uppercase;letter-spacing:.04em;color:var(--mut);font-weight:700;margin-bottom:3px}
.cf{display:grid;grid-template-columns:1fr 1fr;gap:12px}.cf .b{background:#fff;border:1px solid var(--line);border-radius:8px;padding:9px}
.cf .b.cur{border-left:3px solid #c0392b}.cf .b.fut{border-left:3px solid var(--teal,#04A577)}
.cf .b.fut.assumed{opacity:.6;font-style:italic}
ul.roles{margin:4px 0;padding-left:18px}ul.roles li{margin:3px 0}
.grid2{position:relative;width:100%;max-width:620px;height:460px;border:1px solid var(--line);background:#fff;border-radius:10px;margin:14px 0}
.grid2 .q{position:absolute;width:50%;height:50%;display:flex;align-items:flex-start;justify-content:flex-start;padding:8px;font-size:.7rem;font-weight:700;color:var(--mut)}
.grid2 .dot{position:absolute;width:11px;height:11px;border-radius:50%;background:var(--navy);transform:translate(-50%,50%);cursor:pointer;border:1px solid #fff}
.grid2 .axis{position:absolute;font-size:.68rem;color:var(--mut)}
footer{text-align:center;color:var(--mut);font-size:.72rem;padding:18px}
#onepager{display:none}
@media print{
  body{background:#fff}
  header,.tabs,.bar,.stats,footer,.view,.pt-print{display:none !important}
  body.printing #onepager{display:block !important}
  #onepager h2{color:var(--navy)}
  #onepager .op-cf{display:grid;grid-template-columns:1fr 1fr;gap:14px;margin:12px 0}
  #onepager .op-b{border:1px solid #ccc;border-radius:8px;padding:10px}
  #onepager .op-m{display:inline-block;margin:0 18px 8px 0}
  #onepager .op-m b{font-size:1.3rem;color:var(--navy)}
  #onepager li{margin:3px 0}
  .op-foot{margin-top:18px;color:#888;font-size:.7rem;border-top:1px solid #ccc;padding-top:6px}
}
</style></head><body>
<header><h1>%%TITLE%%</h1><div class="sub">%%SUBTITLE%%</div></header>
<div class="wrap">
  <div class="tabs" id="tabs"></div>
  <div class="view active" id="v-table">
    <div class="stats" id="stats"></div>
    <div class="bar" id="bar"></div>
    <div style="overflow:auto"><table id="tbl"><thead><tr id="head"></tr></thead><tbody id="body"></tbody></table></div>
  </div>
  %%EXTRA_VIEWS%%
</div>
<div id="onepager"></div>
<footer>%%FOOTER%%</footer>
<script>
const DATA=%%DATA%%, CONFIG=%%CONFIG%%, STATUS_COLORS=%%STATUSCOLORS%%, BRAND=%%BRAND%%;
const $=s=>document.querySelector(s), esc=s=>String(s==null?'':s).replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
let sortKey=null, sortDir=1, expanded=null;
function vals(d,k){const v=d[k];return Array.isArray(v)?v:(v==null||v===''?[]:[v]);}
function distinct(k){const s=new Set();DATA.forEach(d=>vals(d,k).forEach(v=>s.add(v)));return [...s].sort();}
function buildBar(){
  const bar=$('#bar');
  CONFIG.filters.forEach(f=>{const s=document.createElement('select');s.id='f-'+f.key;
    s.innerHTML='<option value="">'+esc(f.label)+'</option>'+distinct(f.key).map(v=>'<option>'+esc(v)+'</option>').join('');
    s.onchange=render;bar.appendChild(s);});
  const q=document.createElement('input');q.id='q';q.placeholder='Search…';q.oninput=render;bar.appendChild(q);
  const b=document.createElement('button');b.className='btn';b.textContent='Reset';b.onclick=()=>{CONFIG.filters.forEach(f=>$('#f-'+f.key).value='');$('#q').value='';render();};bar.appendChild(b);
}
function filtered(){
  let d=DATA.slice();
  CONFIG.filters.forEach(f=>{const v=$('#f-'+f.key).value;if(v)d=d.filter(x=>{const xv=x[f.key];return Array.isArray(xv)?xv.map(String).includes(v):String(xv)===v;});});
  const q=($('#q').value||'').toLowerCase();
  if(q)d=d.filter(x=>JSON.stringify(x).toLowerCase().includes(q));
  if(sortKey)d.sort((a,b)=>{let x=a[sortKey],y=b[sortKey];if(typeof x==='number'&&typeof y==='number')return (x-y)*sortDir;return String(x).localeCompare(String(y))*sortDir;});
  return d;
}
function dimChips(d){
  if(!d.dim)return '';
  const prim=d.primary_dimension||'';
  const order=Object.keys(d.dim).filter(k=>d.dim[k]).sort((a,b)=>(b===prim)-(a===prim));
  return order.map(n=>'<span class="chip '+(n===prim?'prim':'sec')+'" style="background:'+(CONFIG.dimColors[n]||'#888')+'">'+esc(n)+(n===prim?' ★':'')+'</span>').join('')
    +(d.scope&&d.scope!=='In Scope'?' <span class="scope">'+esc(d.scope)+'</span>':'');
}
function statusChip(v){if(!v)return '';return '<span class="statc" style="background:'+(STATUS_COLORS[v]||'#888')+'">'+esc(v)+'</span>';}
function listChips(v,cls){const a=Array.isArray(v)?v:(v?[v]:[]);return a.length?a.map(x=>'<span class="chip '+(cls||'')+'">'+esc(x)+'</span>').join(''):'<span style="color:#bbb">—</span>';}
function cell(d,c){
  const v=d[c.key];
  if(c.kind==='dims')return dimChips(d);
  if(c.kind==='pri')return '<span class="pri '+esc(v)+'">'+esc(v)+'</span>';
  if(c.kind==='status')return statusChip(v);
  if(c.kind==='levers')return listChips(v,'lever');
  if(c.kind==='mods')return listChips(v,'mod');
  if(c.kind==='stale')return v?'<span class="stale-flag">stale</span>':'';
  if(c.kind==='chip'&&v)return '<span class="chip" style="background:'+esc(c.color||'#888')+'">'+esc(v)+'</span>';
  return clipTxt(v,c.clip||9999);
}
function clipTxt(v,n){
  const s=String(v==null?'':v);
  if(s.length<=n)return esc(s);
  let t=s.slice(0,n);
  const i=t.lastIndexOf(' ');
  if(i>n*0.5)t=t.slice(0,i);
  return esc(t.replace(/[\s,;:\-—]+$/,''))+'…';
}
function render(){
  const rows=filtered();
  $('#head').innerHTML=CONFIG.columns.map(c=>'<th data-k="'+c.key+'">'+esc(c.label)+(sortKey===c.key?(sortDir>0?' ▴':' ▾'):'')+'</th>').join('');
  $('#head').querySelectorAll('th').forEach(th=>th.onclick=()=>{const k=th.dataset.k;sortDir=(sortKey===k?-sortDir:1);sortKey=k;render();});
  $('#body').innerHTML=rows.map(d=>{
    const tr='<tr class="row" data-id="'+d[CONFIG.idKey]+'">'+CONFIG.columns.map(c=>'<td>'+cell(d,c)+'</td>').join('')+'</tr>';
    const det=(expanded===d[CONFIG.idKey])?'<tr class="detail"><td colspan="'+CONFIG.columns.length+'">'+detail(d)+'</td></tr>':'';
    return tr+det;}).join('');
  $('#body').querySelectorAll('tr.row').forEach(tr=>tr.onclick=()=>{const id=tr.dataset.id;expanded=(String(expanded)===id)?null:(isNaN(id)?id:+id);render();});
  buildStats(rows);
}
function detail(d){
  return CONFIG.detail.map(s=>{
    if(s.type==='currentfuture'){const assumed=(d.future_state_status==='Assumed')?' assumed':'';
      return '<div class="sec"><div class="cf"><div class="b cur"><div class="lbl">Current state</div>'+esc(d.current_state)+'</div><div class="b fut'+assumed+'"><div class="lbl">Future state'+(d.future_state_status?(' — '+esc(d.future_state_status)):'')+'</div>'+esc(d.future_state)+'</div></div></div>';}
    if(s.type==='dims')return '<div class="sec"><div class="lbl">'+esc(s.label)+'</div>'+dimChips(d)+'</div>';
    if(s.type==='roles'){const rc=d.role_changes||{};const names=String(d.roles_impacted||'').split(/[;,]/).map(x=>x.trim()).filter(Boolean);
      if(!names.length)return '';return '<div class="sec"><div class="lbl">'+esc(s.label)+'</div><ul class="roles">'+names.map(n=>'<li><b>'+esc(n)+'</b>'+(rc[n]?(' — '+esc(typeof rc[n]==='string'?rc[n]:(rc[n].desc||''))):'')+'</li>').join('')+'</ul></div>';}
    if(s.type==='levers'){if(!vals(d,'value_levers').length)return '';return '<div class="sec"><div class="lbl">'+esc(s.label)+'</div>'+listChips(d.value_levers,'lever')+'</div>';}
    if(s.type==='mods'){if(!vals(d,'training_modality').length)return '';return '<div class="sec"><div class="lbl">'+esc(s.label)+'</div>'+listChips(d.training_modality,'mod')+'</div>';}
    const v=d[s.key];if(v==null||v==='')return '';
    return '<div class="sec"><div class="lbl">'+esc(s.label)+'</div>'+esc(v)+'</div>';
  }).join('');
}
function buildStats(rows){
  $('#stats').innerHTML=CONFIG.stats.map(s=>{
    let val;
    if(s.type==='count')val=rows.length;
    else if(s.type==='distinct')val=new Set(rows.flatMap(r=>vals(r,s.key))).size;
    else if(s.type==='countif')val=rows.filter(r=>s.test(r)).length;
    return '<div class="card"><div class="v">'+val+'</div><div class="l">'+esc(s.label)+'</div></div>';
  }).join('');
}
CONFIG.stats.forEach(s=>{if(s.countifKey){s.type='countif';s.test=(r=>String(r[s.countifKey])===s.countifVal);}});
// ---- pivot / tiles engine (TOM Shifts, Value Realization, Training, Workshop status)
function testRec(r,t){
  if(t.startsWith!==undefined)return String(r[t.key]||'').startsWith(t.startsWith);
  if(t.eq!==undefined)return String(r[t.key])===String(t.eq);
  if(t.inSet!==undefined)return t.inSet.map(String).includes(String(r[t.key]));
  if(t.gte!==undefined)return (parseFloat(r[t.key])||0)>=t.gte;
  if(t.truthy!==undefined)return !!r[t.key]===t.truthy;
  if(t.all!==undefined)return t.all.every(tt=>testRec(r,tt));
  return false;
}
function pivotGroups(spec){
  const groups={};
  DATA.forEach(d=>{
    let keys;
    if(spec.explode)keys=vals(d,spec.explode);
    else if(spec.roles)keys=String(d.roles_impacted||'').split(/[;,]/).map(s=>s.trim()).filter(Boolean);
    else keys=vals(d,spec.groupBy);
    keys.forEach(k=>{(groups[k]=groups[k]||[]).push(d);});
  });
  (spec.vocab||[]).forEach(k=>{if(!groups[k])groups[k]=[];});
  return groups;
}
function metricVal(rows,m){
  if(m.kind==='count')return rows.length;
  if(m.kind==='sum'){const t=rows.reduce((s,r)=>s+(parseFloat(r[m.key])||0),0);return Math.round(t*10)/10;}
  if(m.kind==='distinct')return new Set(rows.flatMap(r=>vals(r,m.key))).size;
  if(m.kind==='countif')return rows.filter(r=>testRec(r,m.test)).length;
  return '';
}
function renderPivots(){
  (CONFIG.pivots||[]).forEach(spec=>{
    const host=document.getElementById(spec.view);if(!host)return;
    const body=host.querySelector('.pv-body');
    const groups=pivotGroups(spec);
    const names=Object.keys(groups).sort((a,b)=>groups[b].length-groups[a].length||String(a).localeCompare(b));
    if(!names.length){body.innerHTML='<p style="color:var(--mut)">No data for this view yet — populate the relevant fields in your records.</p>';return;}
    if(spec.layout==='tiles'){
      body.className='pv-body tiles';
      body.innerHTML=names.map(n=>{
        const rows=groups[n], meta=(spec.meta&&spec.meta[n])||null, zero=rows.length===0;
        const color=(spec.colorBy&&spec.colorBy[n])||BRAND.navy;
        const mets=spec.metrics.map(m=>'<div class="pm"><div class="pm-v">'+metricVal(rows,m)+'</div><div class="pm-l">'+esc(m.label)+'</div></div>').join('');
        const metaLine=meta?('<div class="pt-meta">'+esc(meta.description||'')+(meta.target_metric?(' <b>Target:</b> '+esc(meta.target_metric)):'')+'</div>'):'';
        const pr=spec.print?'<button class="pt-print" data-print="'+esc(n)+'" title="Print one-pager">⎙</button>':'';
        return '<div class="ptile'+(zero?' pv-zero':'')+'" style="border-top-color:'+esc(color)+'" data-drill="'+esc(spec.drill||'')+'" data-search="'+esc(spec.drillSearch||'')+'" data-val="'+esc(n)+'"><div class="pt-h"><span>'+esc(n)+(zero?' <span class="zero-flag">no impacts</span>':'')+'</span>'+pr+'</div>'+metaLine+'<div class="pm-row">'+mets+'</div></div>';
      }).join('');
    } else {
      body.className='pv-body';
      body.innerHTML='<div style="overflow:auto"><table><thead><tr><th>'+esc(spec.groupLabel||'Group')+'</th>'+spec.metrics.map(m=>'<th>'+esc(m.label)+'</th>').join('')+'</tr></thead><tbody>'+
        names.map(n=>{const rows=groups[n];return '<tr class="prow" data-drill="'+esc(spec.drill||'')+'" data-search="'+esc(spec.drillSearch||'')+'" data-val="'+esc(n)+'"><td>'+esc(n)+(rows.length===0?' <span class="zero-flag">no impacts</span>':'')+'</td>'+spec.metrics.map(m=>'<td>'+metricVal(rows,m)+'</td>').join('')+'</tr>';}).join('')+'</tbody></table></div>';
    }
    body.querySelectorAll('[data-val]').forEach(node=>node.onclick=ev=>{
      if(ev.target.classList.contains('pt-print')){printOnePager(spec,node.dataset.val,groups[node.dataset.val]);return;}
      const k=node.dataset.drill,sk=node.dataset.search,val=node.dataset.val;
      if(k)setRegisterFilter(k,val);else if(sk)setRegisterSearch(val);
    });
  });
}
function gotoRegister(){
  const t=[...document.querySelectorAll('.tab')].find(x=>x.dataset.view==='v-table');
  document.querySelectorAll('.tab').forEach(x=>x.classList.remove('active'));document.querySelectorAll('.view').forEach(x=>x.classList.remove('active'));
  if(t)t.classList.add('active');$('#v-table').classList.add('active');
}
function setRegisterFilter(key,val){
  gotoRegister();const sel=$('#f-'+key);
  if(sel){if(![...sel.options].some(o=>o.value===val)){const o=document.createElement('option');o.value=val;o.textContent=val;sel.appendChild(o);}sel.value=val;}
  render();
}
function setRegisterSearch(val){gotoRegister();const q=$('#q');if(q)q.value=val;render();}
function printOnePager(spec,name,rows){
  const meta=(spec.meta&&spec.meta[name])||{};
  const mets=spec.metrics.map(m=>'<span class="op-m"><b>'+metricVal(rows,m)+'</b> '+esc(m.label)+'</span>').join('');
  const cf=(meta.today||meta.future)?('<div class="op-cf"><div class="op-b"><b>Today</b><div>'+esc(meta.today||'—')+'</div></div><div class="op-b"><b>Future</b><div>'+esc(meta.future||'—')+'</div></div></div>'):'';
  const items=rows.slice().sort((a,b)=>(b.impact_score||0)-(a.impact_score||0)).map(r=>'<li><b>'+esc(r.title)+'</b> — '+esc(r.future_state_status)+(r.priority?(' · '+esc(r.priority)):'')+'</li>').join('');
  $('#onepager').innerHTML='<h2>'+esc(name)+'</h2>'+(meta.description?'<p>'+esc(meta.description)+'</p>':'')+cf+'<div>'+mets+'</div>'+(items?'<h3>Mapped impacts</h3><ul>'+items+'</ul>':'')+'<div class="op-foot">'+esc(CONFIG.title||'')+' — '+esc(BRAND.footer||'')+'</div>';
  document.body.classList.add('printing');window.print();
}
window.addEventListener('afterprint',()=>document.body.classList.remove('printing'));
function buildTabs(){
  const t=$('#tabs');CONFIG.tabs.forEach((tb,i)=>{const e=document.createElement('div');e.className='tab'+(i===0?' active':'');e.textContent=tb.label;e.dataset.view=tb.view;
    if(i===0)$('#v-table').classList.add('active');
    e.onclick=()=>{document.querySelectorAll('.tab').forEach(x=>x.classList.remove('active'));document.querySelectorAll('.view').forEach(x=>x.classList.remove('active'));
      e.classList.add('active');$('#'+tb.view).classList.add('active');if(tb.view==='v-grid')drawGrid();};t.appendChild(e);});
  // first tab may not be v-table; activate the configured first view
  document.querySelectorAll('.view').forEach(x=>x.classList.remove('active'));
  $('#'+CONFIG.tabs[0].view).classList.add('active');
}
function drawGrid(){
  if(!CONFIG.grid)return;const el=$('#grid');if(!el||el.dataset.done)return;el.dataset.done=1;
  el.insertAdjacentHTML('beforeend','<div class="q" style="left:0;top:0">'+CONFIG.grid.q.tl+'</div><div class="q" style="right:0;top:0;text-align:right;justify-content:flex-end">'+CONFIG.grid.q.tr+'</div><div class="q" style="left:0;bottom:0;align-items:flex-end">'+CONFIG.grid.q.bl+'</div><div class="q" style="right:0;bottom:0;align-items:flex-end;justify-content:flex-end">'+CONFIG.grid.q.br+'</div>');
  el.insertAdjacentHTML('beforeend','<div class="axis" style="left:50%;bottom:-18px;transform:translateX(-50%)">Interest →</div><div class="axis" style="left:-6px;top:50%;transform:rotate(-90deg);transform-origin:left">Influence →</div>');
  // group co-located dots per (x,y) cell and apply a small deterministic jitter so every group stays discoverable
  const cells={};
  DATA.forEach(d=>{const x=d[CONFIG.grid.x],y=d[CONFIG.grid.y];if(!x||!y)return;const k=x+'_'+y;(cells[k]=cells[k]||[]).push(d);});
  Object.values(cells).forEach(list=>list.forEach((d,i)=>{
    const x=d[CONFIG.grid.x],y=d[CONFIG.grid.y];
    const ang=i*2.399963, r=7*Math.sqrt(i); // golden-angle spiral, px offsets, deterministic
    const dx=r*Math.cos(ang), dy=r*Math.sin(ang);
    const dot=document.createElement('div');dot.className='dot';
    dot.style.left='calc('+((x-1)/4*100)+'% + '+dx.toFixed(1)+'px)';
    dot.style.bottom='calc('+((y-1)/4*100)+'% + '+dy.toFixed(1)+'px)';
    dot.title=d[CONFIG.grid.label]+' (Infl '+y+', Int '+x+')'+(list.length>1?' — '+list.length+' groups at this point':'');
    el.appendChild(dot);}));
  // groups identified but not yet assessed render as a list, not as dots
  const pend=DATA.filter(d=>!d[CONFIG.grid.x]||!d[CONFIG.grid.y]);
  if(pend.length&&!document.getElementById('grid-pending'))
    el.insertAdjacentHTML('afterend','<div id="grid-pending" class="pv-note" style="max-width:620px;margin:6px 0 14px"><b>Pending assessment ('+pend.length+'):</b> '+pend.map(d=>esc(d[CONFIG.grid.label])).join('; ')+'</div>');
}
buildBar();buildTabs();renderPivots();render();
</script></body></html>"""

def render_html(out_path, *, title, subtitle, data, config, brand, extra_views=""):
    config = {**config, "title": title}
    htmlout = (TEMPLATE
        .replace("%%TITLE%%", html.escape(title))
        .replace("%%SUBTITLE%%", html.escape(subtitle))
        .replace("%%NAVY%%", brand["navy"]).replace("%%MAG%%", brand["magenta"])
        .replace("%%FONT%%", brand["font"]).replace("%%FOOTER%%", html.escape(brand["footer"]))
        .replace("%%EXTRA_VIEWS%%", extra_views)
        .replace("%%DATA%%", json.dumps(data, ensure_ascii=False))
        .replace("%%CONFIG%%", json.dumps(config, ensure_ascii=False))
        .replace("%%STATUSCOLORS%%", json.dumps(STATUS_COLORS, ensure_ascii=False))
        .replace("%%BRAND%%", json.dumps(brand, ensure_ascii=False)))
    open(out_path, "w", encoding="utf-8").write(htmlout)

# ----------------------------------------------------------------------------- builders
def build_sha(records, project, outdir):
    data = normalize_sha(records)
    brand = {**DEFAULT_BRAND, **project.get("brand", {})}
    pname = project.get("project_name", "Project")
    cols = [("id", "ID", 6), ("group", "Stakeholder Group / Role", 30), ("business_unit", "Business Unit", 16),
            ("location", "Location", 14), ("category", "Category", 22), ("description", "Group Description", 46),
            ("influence", "Influence", 10), ("interest", "Interest", 10),
            ("assessment_status", "Assessment Status", 18),
            ("influence_rationale", "Influence Rationale", 36), ("interest_rationale", "Interest Rationale", 36),
            ("impact_from_change", "Impact from Change", 36), ("decision_authority", "Decision Authority", 30),
            ("sentiment", "Current Sentiment", 16), ("pain_points", "Key Pain Points", 36),
            ("comms", "Communication Preferences", 30), ("engagement", "Engagement Strategy", 36),
            ("poc", "Main Point of Contact", 24), ("champion", "Change Champion", 20),
            ("notes", "Additional Notes", 36), ("follow_up", "Follow-Up Needed", 30)]
    write_xlsx(os.path.join(outdir, f"{pname} SHA.xlsx"), "Stakeholder Assessment", cols, data)
    extra = '<div class="view" id="v-grid"><div class="pv-head"><h3>Influence × Interest</h3></div><div class="grid2" id="grid"></div></div>'
    config = {
        "idKey": "id", "dimColors": DIM_COLORS,
        "tabs": [{"label": "Stakeholder Assessment", "view": "v-table"}, {"label": "Influence / Interest", "view": "v-grid"}],
        "columns": [{"key": "id", "label": "ID"}, {"key": "business_unit", "label": "BU"},
                    {"key": "location", "label": "Location"}, {"key": "group", "label": "Group / Role", "clip": 40},
                    {"key": "influence", "label": "Infl."}, {"key": "interest", "label": "Int."},
                    {"key": "description", "label": "Description", "clip": 80}, {"key": "engagement", "label": "Engagement", "clip": 30}],
        "filters": [{"key": "business_unit", "label": "All BUs"}, {"key": "location", "label": "All Locations"},
                    {"key": "category", "label": "All Categories"}, {"key": "influence", "label": "All Influence"}],
        "stats": [{"label": "Stakeholders", "type": "count"}, {"label": "Business Units", "type": "distinct", "key": "business_unit"},
                  {"label": "High Influence (5)", "countifKey": "influence", "countifVal": "5"},
                  {"label": "Manage Closely", "type": "countif", "countifKey": "quadrant", "countifVal": "Manage Closely"}],
        "detail": [{"label": "Group Description", "key": "description"}, {"label": "Assessment Status", "key": "assessment_status"},
                   {"label": "Influence Rationale", "key": "influence_rationale"},
                   {"label": "Interest Rationale", "key": "interest_rationale"}, {"label": "Impact from Change", "key": "impact_from_change"},
                   {"label": "Decision Authority", "key": "decision_authority"}, {"label": "Key Pain Points", "key": "pain_points"},
                   {"label": "Communication Preferences", "key": "comms"}, {"label": "Engagement Strategy", "key": "engagement"},
                   {"label": "Main Point of Contact", "key": "poc"}, {"label": "Change Champion", "key": "champion"},
                   {"label": "Follow-Up Needed", "key": "follow_up"}],
        "pivots": [],
        "grid": {"x": "interest", "y": "influence", "label": "group",
                 "q": {"tl": "Keep Satisfied", "tr": "Manage Closely", "bl": "Monitor", "br": "Keep Informed"}},
    }
    render_html(os.path.join(outdir, f"{pname} SHA Dashboard.html"),
                title=f"{pname} — Stakeholder Assessment",
                subtitle=f"{len(data)} stakeholder groups",
                data=data, config=config, brand=brand, extra_views=extra)
    return len(data)

def write_readme(outdir, project):
    pname = project.get("project_name", "Project")
    files = sorted(f for f in os.listdir(outdir)
                   if f.startswith(pname) and "SHA" in f and (f.endswith(".xlsx") or f.endswith(".html")))
    lines = [f"# {pname} \u2014 Stakeholder Assessment", "",
             "Generated by the `sha-builder` skill from transcripts/interviews.", "",
             "## Deliverables", ""]
    for f in files:
        what = ("interactive dashboard \u2014 open in a browser; filter, sort, click a row for full detail"
                if f.endswith(".html") else "data workbook \u2014 one row per stakeholder group, all fields")
        lines.append(f"- **{f}** \u2014 {what}")
    lines += ["", "## SHA \u2014 Stakeholder Assessment",
              "Each group is rated on Influence and Interest (1-5). The Engagement Strategy quadrant "
              "(Manage Closely / Keep Satisfied / Keep Informed / Monitor) is derived from the "
              "Influence x Interest grid \u2014 see the *Influence / Interest* tab.", "",
              "## Methodology", "See the `sha-builder` skill references (SHA_SCHEMA, METHODOLOGY, UPDATE_LOOP)."]
    open(os.path.join(outdir, "README.md"), "w", encoding="utf-8").write("\n".join(lines) + "\n")

# ----------------------------------------------------------------------------- cli
def _guard(outdir, names, force):
    """Refuse to overwrite existing deliverables unless --force is set."""
    import os, sys
    existing = [n for n in names if os.path.exists(os.path.join(outdir, n))]
    if existing and not force:
        sys.exit("Refusing to overwrite existing output(s) in %r:\n  %s\n"
                 "Re-run with --force to overwrite, or choose a different --outdir."
                 % (outdir, "\n  ".join(existing)))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("kind", choices=["sha", "readme"])
    ap.add_argument("--records"); ap.add_argument("--config", required=True)
    ap.add_argument("--outdir", default=".")
    ap.add_argument("--force", action="store_true", help="overwrite existing outputs in --outdir")
    a = ap.parse_args()
    os.makedirs(a.outdir, exist_ok=True)
    project = json.load(open(a.config, encoding="utf-8"))
    if a.kind == "sha":
        _pn = project.get("project_name", "Project")
        _guard(a.outdir, [f"{_pn} SHA.xlsx", f"{_pn} SHA Dashboard.html"], a.force)
    if a.kind == "readme":
        write_readme(a.outdir, project); print(f"README.md -> {a.outdir}"); return
    if not a.records:
        sys.exit("--records is required for sha")
    records = json.load(open(a.records, encoding="utf-8"))
    n = build_sha(records, project, a.outdir); print(f"SHA: {n} stakeholders -> {a.outdir}")

if __name__ == "__main__":
    main()
