#!/usr/bin/env python3
"""
tom_slide.py — render the TOM Agent "accelerator" capability slide (PPTX) for ANY
project. Builds on a blank branded 16:9 deck (no client template required), reading
brand from project.json and the value statement + 6 capability cards from
tom_blueprint.json["accelerator"]. Icons are embedded from ../assets/icons/.

Usage:
  python3 tom_slide.py --config project.json --blueprint tom_blueprint.json --outdir OUT

Requires python-pptx. Cross-platform.
"""
import argparse, json
from io import BytesIO
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ICON_DIR = Path(__file__).resolve().parent.parent / "assets" / "icons"
DEFAULT_ICONS = ["funnel", "map", "shapes", "ranking", "discover", "output"]

DEF_BRAND = {"navy": "#0E2841", "corp": "#156082", "accent": "#2773FB", "cyan": "#12ABDB",
             "card_bg": "#F4F5F7", "card_br": "#D9DDE3", "text_dk": "#231F20",
             "text_md": "#333333", "text_lt": "#5A6B8A",
             "font": "Calibri", "footer": "Confidential"}


def rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def soft_shadow(shape):
    spPr = shape._element.spPr
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    eff = spPr.makeelement(qn("a:effectLst"), {})
    sh = eff.makeelement(qn("a:outerShdw"),
                         {"blurRad": "50800", "dist": "25400", "dir": "5400000", "rotWithShape": "0"})
    clr = sh.makeelement(qn("a:srgbClr"), {"val": "8A94A6"})
    clr.append(clr.makeelement(qn("a:alpha"), {"val": "38000"}))
    sh.append(clr); eff.append(sh); spPr.append(eff)


def _guard(outdir, names, force):
    """Refuse to overwrite existing deliverables unless --force is set."""
    import os, sys
    existing = [n for n in names if os.path.exists(os.path.join(outdir, n))]
    if existing and not force:
        sys.exit("Refusing to overwrite existing output(s) in %r:\n  %s\n"
                 "Re-run with --force to overwrite, or choose a different --outdir."
                 % (outdir, "\n  ".join(existing)))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", required=True)
    ap.add_argument("--blueprint", required=True)
    ap.add_argument("--outdir", default=".")
    ap.add_argument("--force", action="store_true", help="overwrite existing outputs in --outdir")
    a = ap.parse_args()
    Path(a.outdir).mkdir(parents=True, exist_ok=True)

    project = json.load(open(a.config, encoding="utf-8"))
    blueprint = json.load(open(a.blueprint, encoding="utf-8"))
    acc = blueprint.get("accelerator", {})
    pname = project.get("project_name", "Project")
    _guard(a.outdir, [f"{pname} TOM Accelerator.pptx"], a.force)

    b = dict(DEF_BRAND)
    cfg_brand = project.get("brand", {})
    b["navy"] = cfg_brand.get("navy", b["navy"])
    for k in ("corp", "accent", "cyan", "card_bg", "card_br", "text_dk", "text_md", "text_lt", "font", "footer"):
        if k in cfg_brand:
            b[k] = cfg_brand[k]
    # "magenta" is an alias for the accent color (HTML renderer uses it); explicit "accent" wins.
    if "magenta" in cfg_brand and "accent" not in cfg_brand:
        b["accent"] = cfg_brand["magenta"]
    # python-pptx needs a single font name, not a CSS stack like "Calibri, system-ui, sans-serif".
    b["font"] = b["font"].split(",")[0].strip().strip("'\"")
    NAVY, CORP, ACCENT, CYAN = rgb(b["navy"]), rgb(b["corp"]), rgb(b["accent"]), rgb(b["cyan"])
    CARDBG, CARDBR = rgb(b["card_bg"]), rgb(b["card_br"])
    TXTDK, TXTMD, WHITE = rgb(b["text_dk"]), rgb(b["text_md"]), RGBColor(0xFF, 0xFF, 0xFF)
    FONT = b["font"]

    title = acc.get("title", "Target Operating Model (TOM) Agent")
    subtitle = acc.get("subtitle", "AI-accelerated business-process operating model design")
    value = acc.get("value_statement", "")
    caps = acc.get("capabilities", [])[:6]

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    shapes = slide.shapes

    def box(x, y, w, h):
        return shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame

    def run(tf, text, size, color, bold=False, italic=False, align=None):
        p = tf.paragraphs[0]
        if align:
            p.alignment = align
        r = p.add_run(); r.text = text
        r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color
        return p, r

    # left panel
    panel = shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(4.30), Inches(7.5))
    panel.fill.solid(); panel.fill.fore_color.rgb = NAVY; panel.line.fill.background(); panel.shadow.inherit = False
    rule = shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(2.18), Inches(0.62), Inches(0.06))
    rule.fill.solid(); rule.fill.fore_color.rgb = CYAN; rule.line.fill.background(); rule.shadow.inherit = False

    run(box(0.55, 1.65, 3.3, 0.4), "ACCELERATOR", 13, CYAN, bold=True)
    tb = box(0.55, 2.35, 3.35, 1.7); tb.word_wrap = True; run(tb, title, 27, WHITE, bold=True)
    st = box(0.55, 4.05, 3.35, 0.6); st.word_wrap = True
    run(st, subtitle, 11, rgb("#B8C6DB"), italic=True)

    # product mock window
    win = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(4.85), Inches(3.30), Inches(0.78))
    win.adjustments[0] = 0.12
    win.fill.solid(); win.fill.fore_color.rgb = WHITE; win.line.color.rgb = rgb("#E3E6EB"); win.line.width = Pt(0.75)
    win.shadow.inherit = False; soft_shadow(win)
    for i, c in enumerate(["#EE2C81", "#F08301", "#04A577"]):
        dot = shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.72 + i * 0.16), Inches(5.00), Inches(0.09), Inches(0.09))
        dot.fill.solid(); dot.fill.fore_color.rgb = rgb(c); dot.line.fill.background(); dot.shadow.inherit = False
    wl = box(0.62, 5.20, 3.16, 0.4)
    run(wl, title.split("(")[0].strip() if "(" in title else "TOM Agent", 13, NAVY, bold=True, align=PP_ALIGN.CENTER)

    # value statement (narrow so it clears a top-right logo)
    vs = box(4.75, 0.55, 6.55, 1.65); vs.word_wrap = True
    p, _ = run(vs, value, 13, TXTDK); p.line_spacing = 1.08
    run(box(4.75, 2.18, 8.15, 0.35), acc.get("capabilities_header", "Capabilities:"), 14, NAVY, bold=True)

    # 6 cards
    col_x = [4.75, 7.57, 10.39]; row_y = [2.70, 4.92]
    cw, chh = 2.50, 1.92
    for idx, cap in enumerate(caps):
        col, row = idx % 3, idx // 3
        x, y = col_x[col], row_y[row]
        card = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(cw), Inches(chh))
        card.adjustments[0] = 0.06
        card.fill.solid(); card.fill.fore_color.rgb = CARDBG
        card.line.color.rgb = CARDBR; card.line.width = Pt(0.75); card.shadow.inherit = False; soft_shadow(card)
        # number badge
        badge = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x - 0.02), Inches(y - 0.14), Inches(0.52), Inches(0.52))
        badge.adjustments[0] = 0.22
        badge.fill.solid(); badge.fill.fore_color.rgb = [CORP, ACCENT, CYAN][idx % 3]
        badge.line.fill.background(); badge.shadow.inherit = False; soft_shadow(badge)
        btf = badge.text_frame; btf.word_wrap = False
        btf.margin_left = btf.margin_right = btf.margin_top = btf.margin_bottom = 0
        run(btf, str(idx + 1), 18, WHITE, bold=True, align=PP_ALIGN.CENTER)
        # icon
        iname = cap.get("icon", DEFAULT_ICONS[idx] if idx < len(DEFAULT_ICONS) else "shapes")
        ipath = ICON_DIR / f"ic_{iname}.png"
        if ipath.exists():
            isz = 0.46
            shapes.add_picture(str(ipath), Inches(x + cw - 0.16 - isz), Inches(y + 0.10), Inches(isz), Inches(isz))
        # text
        ctf = card.text_frame; ctf.word_wrap = True; ctf.vertical_anchor = MSO_ANCHOR.TOP
        ctf.margin_left = Inches(0.16); ctf.margin_right = Inches(0.16)
        ctf.margin_top = Inches(0.58); ctf.margin_bottom = Inches(0.10)
        cp = ctf.paragraphs[0]; cp.alignment = PP_ALIGN.LEFT; cp.line_spacing = 1.04
        r1 = cp.add_run(); r1.text = cap.get("lead", ""); r1.font.name = FONT; r1.font.size = Pt(10.5)
        r1.font.bold = True; r1.font.color.rgb = NAVY
        r2 = cp.add_run(); r2.text = cap.get("rest", ""); r2.font.name = FONT; r2.font.size = Pt(10.5)
        r2.font.color.rgb = TXTDK

    # optional logo top-right
    logo = cfg_brand.get("logo_path")
    if logo and Path(logo).exists():
        shapes.add_picture(logo, Inches(11.39), Inches(0.38), Inches(1.59), Inches(0.82))

    # footer
    run(box(4.75, 7.12, 7.0, 0.22), b["footer"], 9, TXTMD)

    out = Path(a.outdir) / f"{pname} TOM Accelerator.pptx"
    prs.save(str(out))
    print(f"TOM slide -> {out}")


if __name__ == "__main__":
    main()
