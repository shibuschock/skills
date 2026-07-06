#!/usr/bin/env python3
"""
curriculum-architect — render a curriculum blueprint as an xlsx workbook + a
self-contained HTML blueprint, for ANY project, from curriculum.json
(see references/CURRICULUM_SCHEMA.md: learning_paths, modules, assessment_plan).

Usage:
  python3 curriculum_render.py --records curriculum.json --config project.json --outdir OUT [--force]

Only stdlib + openpyxl. Output HTML has no external dependencies.
"""
import argparse, html, json, os, sys

DEFAULT_BRAND = {"navy": "#162B75", "magenta": "#EE2C81",
                 "font": "Calibri, system-ui, sans-serif", "footer": "Confidential"}
BLOOM_COLORS = {"Remember": "#6b7280", "Understand": "#1d6fb8", "Apply": "#04A577",
                "Analyze": "#F08301", "Evaluate": "#EE2C81", "Create": "#7A4FBE"}
BAND_COLORS = {"experiential": "#04A577", "social": "#F08301", "formal": "#162B75"}
BAND_PCT = {"experiential": 70, "social": 20, "formal": 10}

def fmt(v):
    if isinstance(v, bool): return "Yes" if v else ""
    if isinstance(v, (list, tuple)): return ", ".join(str(x) for x in v)
    if isinstance(v, dict): return "; ".join(f"{k}: {x}" for k, x in v.items())
    return v

def style_sheet(ws, columns, navy):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    thin = Side(style="thin", color="D9D9D9"); border = Border(thin, thin, thin, thin)
    for ci, _ in enumerate(columns, 1):
        c = ws.cell(1, ci); c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=navy.lstrip("#"))
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = border
    ws.freeze_panes = "B2"; ws.row_dimensions[1].height = 26
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
    for ci, col in enumerate(columns, 1):
        L = get_column_letter(ci); ws.column_dimensions[L].width = col[2] if len(col) > 2 else 22
        for cell in ws[L][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top"); cell.border = border
            cell.font = Font(size=10)

def path_hours(path, mod_by_id):
    mins = sum((mod_by_id.get(mid, {}).get("duration_minutes") or 0) for mid in path.get("sequence", []))
    return round(mins / 60, 1)

def write_xlsx(path, cur, navy):
    from openpyxl import Workbook
    mods = cur.get("modules", []); paths = cur.get("learning_paths", [])
    mod_by_id = {m.get("id"): m for m in mods}
    wb = Workbook()
    mcols = [("id", "Module ID", 18), ("title", "Title", 34), ("audience", "Audience", 28),
             ("objectives", "Learning Objectives (Bloom)", 60), ("duration_minutes", "Duration (min)", 12),
             ("blend", "70-20-10 Blend", 60), ("prerequisites", "Prerequisites", 18),
             ("content_outline", "Content Outline", 44),
             ("assessment", "Assessment (type / mastery / Kirkpatrick)", 44),
             ("maps_to_needs", "Maps to Needs", 18), ("maps_to_impacts", "Maps to Impacts (CIA)", 44),
             ("super_user_track", "Super-User Track", 12),
             ("non_training_gates", "Non-Training Gates", 36), ("status", "Status", 18)]
    ws = wb.active; ws.title = "Modules Register"
    ws.append([c[1] for c in mcols])
    for m in mods:
        objs = "\n".join(f"[{o.get('bloom_level','')}] {o.get('text','')} (→ {o.get('capability_ref','')})"
                         for o in m.get("learning_objectives", []))
        blend = "\n".join(f"{k.title()} ({BAND_PCT.get(k,'')}%): {v}" for k, v in (m.get("modality_blend") or {}).items())
        a = m.get("assessment") or {}
        assess = f"{a.get('type','')} / {a.get('mastery_threshold','')} / {a.get('kirkpatrick_level','')}"
        ws.append([m.get("id"), m.get("title"), m.get("audience"), objs, m.get("duration_minutes"),
                   blend, fmt(m.get("prerequisites", [])), fmt(m.get("content_outline", [])),
                   assess, fmt(m.get("maps_to_needs", [])), fmt(m.get("maps_to_impacts", [])),
                   fmt(bool(m.get("super_user_track"))),
                   fmt(m.get("non_training_gates", [])),
                   (m.get("status") or "planned") + (f" — {m.get('gate')}" if m.get("gate") else "")])
    style_sheet(ws, mcols, navy)

    pcols = [("path_name", "Learning Path", 34), ("audience", "Audience", 30), ("business_unit", "Business Unit", 22),
             ("sequence", "Module Sequence", 44), ("total_hours", "Total Hours", 12),
             ("target_proficiency", "Target Proficiency", 16)]
    ws2 = wb.create_sheet("Learning Paths")
    ws2.append([c[1] for c in pcols])
    for p in paths:
        ws2.append([p.get("path_name"), p.get("audience"), p.get("business_unit"),
                    " → ".join(p.get("sequence", [])), path_hours(p, mod_by_id), p.get("target_proficiency", "")])
    style_sheet(ws2, pcols, navy)
    wb.save(path)

# ----------------------------------------------------------------------------- html
CSS = r"""
:root{--navy:%%NAVY%%;--mag:%%MAG%%;--ink:#1d2433;--mut:#6b7280;--line:#e5e7eb;--bg:#f5f6f8}
*{box-sizing:border-box}body{margin:0;font-family:%%FONT%%;color:var(--ink);background:var(--bg);font-size:14px}
header{background:linear-gradient(135deg,var(--navy),#0E2841);color:#fff;padding:18px 28px}
header h1{margin:0;font-size:1.15rem;font-weight:700}header .sub{opacity:.85;font-size:.85rem;margin-top:3px}
.wrap{max-width:1300px;margin:0 auto;padding:18px 28px 80px}
.stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:14px;margin:16px 0}
.card{background:#fff;border:1px solid var(--line);border-radius:10px;padding:14px 16px}
.card .v{font-size:1.6rem;font-weight:800;color:var(--navy)}.card .l{font-size:.72rem;color:var(--mut);text-transform:uppercase;letter-spacing:.04em;margin-top:2px}
h2.path{color:var(--navy);margin:26px 0 4px;font-size:1.05rem}
.pmeta{color:var(--mut);font-size:.8rem;margin-bottom:10px}
.seq{display:flex;flex-wrap:wrap;align-items:center;gap:6px;margin:8px 0 14px}
.seq .step{background:var(--navy);color:#fff;border-radius:16px;padding:4px 12px;font-size:.76rem;font-weight:700}
.seq .arr{color:var(--mag);font-weight:800}
.mod{background:#fff;border:1px solid var(--line);border-left:5px solid var(--navy);border-radius:10px;padding:14px 18px;margin:10px 0}
.mod.su{border-left-color:var(--mag)}
.mod h3{margin:0;font-size:.98rem;color:var(--navy);display:flex;gap:10px;align-items:baseline;flex-wrap:wrap}
.mod .dur{font-size:.72rem;color:var(--mut);font-weight:400}
.subadge{background:var(--mag);color:#fff;border-radius:9px;padding:1px 8px;font-size:.62rem;font-weight:800;text-transform:uppercase}
.prereq{font-size:.74rem;color:var(--mut);margin:4px 0}
.lbl{font-size:.68rem;text-transform:uppercase;letter-spacing:.04em;color:var(--mut);font-weight:700;margin:10px 0 4px}
ul.obj{margin:4px 0;padding-left:18px}ul.obj li{margin:4px 0}
.bloom{display:inline-block;padding:1px 8px;border-radius:9px;color:#fff;font-size:.64rem;font-weight:700;margin-left:6px}
.blend{display:flex;height:26px;border-radius:7px;overflow:hidden;max-width:640px;margin:4px 0}
.blend div{color:#fff;font-size:.66rem;font-weight:800;display:flex;align-items:center;justify-content:center}
.bandtxt{font-size:.78rem;margin:3px 0}.bandtxt b{text-transform:capitalize}
.chip{display:inline-block;padding:2px 9px;border-radius:11px;background:#5a6b86;color:#fff;font-size:.68rem;font-weight:600;margin:1px 3px 1px 0;white-space:nowrap}
.chip.need{background:#04A577}.chip.imp{background:#0E2841;opacity:.85}
.chip.gate{background:#B45309}
.mod.gated{border-left-color:#B45309;background:#fffaf1}
.gbadge{background:#B45309;color:#fff;border-radius:9px;padding:1px 8px;font-size:.62rem;font-weight:800;text-transform:uppercase}
.blendnote{font-size:.68rem;color:var(--mut);margin:2px 0 4px}
.assess{background:#f6f8fc;border:1px solid var(--line);border-radius:8px;padding:8px 12px;font-size:.8rem;margin-top:6px}
.panel{background:#fff;border:1px solid var(--line);border-radius:10px;padding:16px 20px;margin:26px 0}
.panel h2{color:var(--navy);font-size:1.02rem;margin:0 0 8px}
.panel p{margin:6px 0;font-size:.86rem}
footer{text-align:center;color:var(--mut);font-size:.72rem;padding:18px}
"""

def esc(s): return html.escape(str(s if s is not None else ""))

def chips(items, cls=""):
    return "".join(f'<span class="chip {cls}">{esc(x)}</span>' for x in (items or [])) or '<span style="color:#bbb">—</span>'

def module_card(m):
    su = bool(m.get("super_user_track"))
    objs = "".join(
        f'<li>{esc(o.get("text",""))}<span class="bloom" style="background:{BLOOM_COLORS.get(o.get("bloom_level"),"#888")}">'
        f'{esc(o.get("bloom_level",""))}</span> <span style="color:#6b7280;font-size:.72rem">→ {esc(o.get("capability_ref",""))}</span></li>'
        for o in m.get("learning_objectives", []))
    blend = m.get("modality_blend") or {}
    weights = m.get("blend_weights") or {}
    try:
        wsum = sum(float(weights.get(b, 0)) for b in ("experiential", "social", "formal"))
    except (TypeError, ValueError):
        wsum = 0
    if wsum > 0:
        pct = {b: round(float(weights.get(b, 0)) / wsum * 100) for b in ("experiential", "social", "formal")}
        blend_note = ""
    else:
        pct = BAND_PCT
        blend_note = '<div class="blendnote">default blend (70-20-10 reference model — no per-module weights supplied)</div>'
    bar = "".join(f'<div style="width:{pct[b]}%;background:{BAND_COLORS[b]}">{pct[b]}% {b}</div>'
                  for b in ("experiential", "social", "formal") if pct[b] > 0)
    bands = "".join(f'<div class="bandtxt"><b style="color:{BAND_COLORS[b]}">{b}</b> — {esc(blend.get(b,""))}</div>'
                    for b in ("experiential", "social", "formal") if blend.get(b))
    a = m.get("assessment") or {}
    prereq = m.get("prerequisites") or []
    status = (m.get("status") or "planned").lower()
    gated = status.startswith("gated")
    gates = m.get("non_training_gates") or []
    gates_html = f'<div class="lbl">Non-training gates</div>{chips(gates, "gate")}' if gates else ''
    gbadge = f'<span class="gbadge">Gated — do not build{(" · " + esc(m.get("gate"))) if m.get("gate") else ""}</span>' if gated else ''
    return f'''<div class="mod{' su' if su else ''}{' gated' if gated else ''}" id="mod-{esc(m.get("id"))}">
<h3>{esc(m.get("id"))} — {esc(m.get("title"))}<span class="dur">{esc(m.get("duration_minutes"))} min</span>{'<span class="subadge">super-user track</span>' if su else ''}{gbadge}</h3>
{f'<div class="prereq">Prerequisites: {esc(", ".join(prereq))}</div>' if prereq else ''}
<div class="lbl">Learning objectives (backward design)</div><ul class="obj">{objs}</ul>
<div class="lbl">70-20-10 blend</div><div class="blend">{bar}</div>{blend_note}{bands}
{gates_html}
<div class="assess"><b>Assessment:</b> {esc(a.get("type",""))} · <b>Mastery:</b> {esc(a.get("mastery_threshold",""))} · <b>Kirkpatrick:</b> {esc(a.get("kirkpatrick_level",""))}</div>
<div class="lbl">Traceability</div>Needs: {chips(m.get("maps_to_needs"), "need")}<br>Impacts: {chips(m.get("maps_to_impacts"), "imp")}
</div>'''

def write_html(path, cur, project, brand):
    pname = project.get("project_name", "Project")
    mods = cur.get("modules", []); paths = cur.get("learning_paths", [])
    mod_by_id = {m.get("id"): m for m in mods}
    body = []
    total_h = sum(path_hours(p, mod_by_id) for p in paths)
    unique_h = round(sum((m.get("duration_minutes") or 0) for m in mods) / 60, 1)
    stats = [("Learning Paths", len(paths)), ("Modules", len(mods)),
             ("Objectives", sum(len(m.get("learning_objectives", [])) for m in mods)),
             ("Total Path Hours (shared modules count per path)", round(total_h, 1)),
             ("Unique Module Hours (build/delivery volume)", unique_h),
             ("Super-User Modules", sum(1 for m in mods if m.get("super_user_track")))]
    body.append('<div class="stats">' + "".join(
        f'<div class="card"><div class="v">{v}</div><div class="l">{esc(l)}</div></div>' for l, v in stats) + "</div>")
    for p in paths:
        seq = p.get("sequence", [])
        seq_html = '<span class="arr">→</span>'.join(f'<span class="step">{esc(mid)}</span>' for mid in seq)
        body.append(f'<h2 class="path">{esc(p.get("path_name"))}</h2>'
                    f'<div class="pmeta">{esc(p.get("audience"))} · {esc(p.get("business_unit"))} · '
                    f'{path_hours(p, mod_by_id)} hrs · target: {esc(p.get("target_proficiency",""))}</div>'
                    f'<div class="seq">{seq_html}</div>')
        for mid in seq:
            m = mod_by_id.get(mid)
            body.append(module_card(m) if m else f'<div class="mod"><h3>{esc(mid)} — MISSING MODULE</h3></div>')
    ap = cur.get("assessment_plan") or {}
    if ap:
        rows = "".join(f'<p><b>{esc(k.replace("_", " ").title())}:</b> {esc(v)}</p>' for k, v in ap.items())
        body.append(f'<div class="panel"><h2>Assessment plan (program level)</h2>{rows}</div>')
    doc = ('<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">'
           '<meta name="viewport" content="width=device-width, initial-scale=1.0">'
           f'<title>{esc(pname)} — Curriculum Blueprint</title><style>'
           + CSS.replace("%%NAVY%%", brand["navy"]).replace("%%MAG%%", brand["magenta"]).replace("%%FONT%%", brand["font"])
           + f'</style></head><body><header><h1>{esc(pname)} — Curriculum Blueprint</h1>'
           f'<div class="sub">{len(paths)} learning paths · {len(mods)} modules · backward design: objectives → assessment → content</div></header>'
           f'<div class="wrap">{"".join(body)}</div><footer>{esc(brand["footer"])}</footer></body></html>')
    open(path, "w", encoding="utf-8").write(doc)

# ----------------------------------------------------------------------------- cli
def _guard(outdir, names, force):
    existing = [n for n in names if os.path.exists(os.path.join(outdir, n))]
    if existing and not force:
        sys.exit("Refusing to overwrite existing output(s) in %r:\n  %s\n"
                 "Re-run with --force to overwrite, or choose a different --outdir."
                 % (outdir, "\n  ".join(existing)))

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--records", required=True); ap.add_argument("--config", required=True)
    ap.add_argument("--outdir", default=".")
    ap.add_argument("--force", action="store_true", help="overwrite existing outputs in --outdir")
    a = ap.parse_args()
    os.makedirs(a.outdir, exist_ok=True)
    project = json.load(open(a.config, encoding="utf-8"))
    brand = {**DEFAULT_BRAND, **project.get("brand", {})}
    pname = project.get("project_name", "Project")
    names = [f"{pname} Curriculum.xlsx", f"{pname} Curriculum Blueprint.html"]
    _guard(a.outdir, names, a.force)
    cur = json.load(open(a.records, encoding="utf-8"))
    write_xlsx(os.path.join(a.outdir, names[0]), cur, brand["navy"])
    write_html(os.path.join(a.outdir, names[1]), cur, project, brand)
    print(f"Curriculum: {len(cur.get('modules', []))} modules / {len(cur.get('learning_paths', []))} paths -> {a.outdir}")

if __name__ == "__main__":
    main()
