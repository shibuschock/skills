#!/usr/bin/env python3
"""
persona-journey-builder — render a tiered change-persona set (personas + user
journeys + day-in-the-life narratives) as an xlsx register + a self-contained
interactive HTML explorer, for ANY project, from structured records (built by
Claude from CIA/SHA data or interview notes).

Methodology baked in (portable OCM):
  - Tiered persona model: Tier 1 deep persona + journey + day-in-the-life,
    Tier 2 lightweight profile, Tier 3 reference-only role list.
  - Journey anatomy: stages with current pain -> future moment, emotion curve,
    and a change-support cue per stage.
  - Optional CIA cross-check: roles_covered must map to CIA roles_impacted;
    unmapped high-severity roles are flagged (console + HTML banner).

Usage:
  python3 persona_render.py --records personas.json --config project.json \
      [--cia cia_records.json] --outdir OUT [--force]

Records = a JSON array of persona objects (see references/SCHEMA.md).
Only stdlib + openpyxl. Cross-platform. Output HTML has no external dependencies.
Shared scaffolding (guard, xlsx styling) follows the suite canonical source: cia-builder.
"""
import argparse, html, json, os, re, sys

DEFAULT_BRAND = {"navy": "#162B75", "magenta": "#EE2C81", "orange": "#F08301",
                 "teal": "#04A577", "coral": "#FF533C",
                 "font": "Calibri, system-ui, sans-serif", "footer": "Confidential"}
TIER_LABEL = {1: "Tier 1 — Deep persona", 2: "Tier 2 — Lightweight profile",
              3: "Tier 3 — Reference only"}
TIER_COLOR = {1: "#162B75", 2: "#04A577", 3: "#9aa3b2"}
INTENSITY_COLOR = {"High": "#EE2C81", "Medium": "#F08301", "Low": "#9aa3b2"}


def g(rec, *keys, default=""):
    for k in keys:
        if k in rec and rec[k] not in (None, ""):
            return rec[k]
    return default


def as_list(v):
    if v is None or v == "":
        return []
    if isinstance(v, (list, tuple)):
        return [str(x).strip() for x in v if str(x).strip()]
    return [s.strip() for s in re.split(r"[;,]", str(v)) if s.strip()]


def split_roles(v):
    """Split a roles field on TOP-LEVEL commas/semicolons only (commas inside
    parentheses don't split), then strip parenthetical content from each token."""
    if v is None or v == "":
        return []
    if isinstance(v, (list, tuple)):
        items = [str(x) for x in v]
    else:
        items, cur, depth = [], "", 0
        for ch in str(v):
            if ch == "(":
                depth += 1
            elif ch == ")":
                depth = max(0, depth - 1)
            if ch in ",;" and depth == 0:
                items.append(cur)
                cur = ""
            else:
                cur += ch
        items.append(cur)
    out = []
    for it in items:
        it = re.sub(r"\([^)]*\)", "", it).strip()
        if it:
            out.append(it)
    return out


def as_ditl(blocks):
    """Normalize day-in-the-life blocks to {time, activity} objects."""
    out = []
    for b in blocks or []:
        if isinstance(b, dict):
            out.append({"time": str(b.get("time", "")).strip(),
                        "activity": str(b.get("activity", "")).strip()})
        else:
            s = str(b)
            m = re.match(r"\s*(\d{1,2}:\d{2})\s*[—\-–]\s*(.*)", s)
            out.append({"time": m.group(1), "activity": m.group(2)} if m
                       else {"time": "", "activity": s.strip()})
    return out


def normalize(records):
    out = []
    for i, r in enumerate(records, 1):
        tier = int(g(r, "tier", default=3) or 3)
        journey = []
        for j, st in enumerate(r.get("journey") or [], 1):
            journey.append({"n": j, "stage": g(st, "stage"),
                            "current": g(st, "current"), "future": g(st, "future"),
                            "emotion_current": g(st, "emotion_current"),
                            "emotion_future": g(st, "emotion_future"),
                            "support": g(st, "support")})
        ditl = r.get("day_in_life") or {}
        out.append({
            "id": r.get("id", i),
            "name": g(r, "name"),
            "tier": tier,
            "tier_label": TIER_LABEL.get(tier, TIER_LABEL[3]),
            "archetype": g(r, "archetype"),
            "roles_covered": as_list(r.get("roles_covered")),
            "population_estimate": g(r, "population_estimate"),
            "change_intensity": g(r, "change_intensity"),
            "top_impacts": as_list(r.get("top_impacts")),
            "goals": as_list(r.get("goals")),
            "pain_points": as_list(r.get("pain_points")),
            "quote": g(r, "quote"),
            "sentiment": g(r, "sentiment"),
            "confidence": g(r, "confidence"),
            "journey": journey,
            "ditl_before": as_ditl(ditl.get("before")),
            "ditl_after": as_ditl(ditl.get("after")),
            "day_in_life_status": g(r, "day_in_life_status"),
        })
    return out


# ----------------------------------------------------------------------- CIA cross-check
def cross_check(personas, cia_records, aliases=None):
    """Compare persona roles_covered against CIA roles_impacted.
    Returns (warnings, unscored_msg) — warnings is a list of strings (empty =
    clean); unscored_msg is a banner string when the CIA carries no usable
    numeric severity, else None."""
    aliases = aliases or {}
    canon = lambda r: aliases.get(r, r)
    cia_roles = {}  # canonical role -> max severity seen
    scored = 0
    for rec in cia_records:
        sev = rec.get("severity")
        try:
            sev = int(float(sev))
            scored += 1
        except Exception:
            sev = 0
        for role in split_roles(rec.get("roles_impacted")):
            role = canon(role)
            cia_roles[role] = max(cia_roles.get(role, 0), sev)
    covered, dupes = set(), set()
    for p in personas:
        for role in p["roles_covered"]:
            role = canon(role)
            if role in covered:
                dupes.add(role)
            covered.add(role)
    warnings = []
    unscored_msg = None
    if cia_records and scored < max(1, len(cia_records) * 0.1):
        unscored_msg = ("high-severity coverage not evaluated — CIA unscored; "
                        "score the CIA first")
    unmapped = [r for r in sorted(cia_roles) if r not in covered]
    high = [r for r in unmapped if cia_roles[r] >= 4]
    if high:
        warnings.append("HIGH-SEVERITY CIA roles not covered by any persona: " + "; ".join(high))
    low = [r for r in unmapped if cia_roles[r] < 4]
    if low:
        warnings.append("CIA roles not covered by any persona: " + "; ".join(low))
    extra = sorted({canon(r) for p in personas for r in p["roles_covered"]} - set(cia_roles))
    if extra:
        warnings.append("Persona roles_covered not found in CIA roles_impacted "
                        "(check name alignment): " + "; ".join(extra))
    if dupes:
        warnings.append("Roles covered by more than one persona (each role should land "
                        "in exactly one): " + "; ".join(sorted(dupes)))
    return warnings, unscored_msg


# ----------------------------------------------------------------------------- xlsx
def fmt_cell(v):
    if isinstance(v, (list, tuple)):
        return "\n".join(str(x) for x in v)
    return v


def style_sheet(ws, columns, navy):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    thin = Side(style="thin", color="D9D9D9")
    border = Border(thin, thin, thin, thin)
    for ci, _ in enumerate(columns, 1):
        c = ws.cell(1, ci)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=navy.lstrip("#"))
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = border
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 26
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
    for ci, col in enumerate(columns, 1):
        L = get_column_letter(ci)
        ws.column_dimensions[L].width = col[2] if len(col) > 2 else 22
        for cell in ws[L][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            cell.font = Font(size=10)


def write_xlsx(path, personas, navy):
    from openpyxl import Workbook
    wb = Workbook()

    reg_cols = [("id", "ID", 6), ("name", "Persona", 30), ("tier_label", "Tier", 24),
                ("archetype", "Archetype", 50), ("roles_covered", "Roles Covered", 34),
                ("population_estimate", "Population (est.)", 14),
                ("change_intensity", "Change Intensity", 14),
                ("top_impacts", "Top Impacts (CIA titles)", 44),
                ("goals", "Goals", 40), ("pain_points", "Pain Points", 44),
                ("quote", "Quote (sourced)", 36), ("sentiment", "Sentiment", 12),
                ("confidence", "Confidence", 12)]
    ws = wb.active
    ws.title = "Persona Register"
    ws.append([c[1] for c in reg_cols])
    for p in personas:
        ws.append([fmt_cell(p.get(c[0], "")) for c in reg_cols])
    style_sheet(ws, reg_cols, navy)

    jm_cols = [("persona", "Persona", 28), ("n", "#", 5), ("stage", "Stage", 20),
               ("current", "Current (pain)", 46), ("future", "Future (moment)", 46),
               ("emotion_current", "Emotion Now", 14), ("emotion_future", "Emotion Future", 14),
               ("support", "Change Support Cue", 40)]
    ws2 = wb.create_sheet("Journey Moments")
    ws2.append([c[1] for c in jm_cols])
    for p in personas:
        for st in p["journey"]:
            row = dict(st)
            row["persona"] = p["name"]
            ws2.append([fmt_cell(row.get(c[0], "")) for c in jm_cols])
    style_sheet(ws2, jm_cols, navy)
    wb.save(path)


# ----------------------------------------------------------------------------- HTML
TEMPLATE = r"""<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0"><title>%%TITLE%%</title>
<style>
:root{--navy:%%NAVY%%;--mag:%%MAG%%;--ink:#1d2433;--mut:#6b7280;--line:#e5e7eb;--bg:#f5f6f8}
*{box-sizing:border-box}body{margin:0;font-family:%%FONT%%;color:var(--ink);background:var(--bg);font-size:14px}
header{background:linear-gradient(135deg,var(--navy),#0E2841);color:#fff;padding:18px 28px}
header h1{margin:0;font-size:1.15rem;font-weight:700}header .sub{opacity:.85;font-size:.85rem;margin-top:3px}
.wrap{max-width:1280px;margin:0 auto;padding:18px 28px 80px}
.warn{background:#fdecea;border:1px solid #f5b7b1;color:#b40020;border-radius:10px;padding:12px 16px;margin:14px 0;font-size:.85rem}
.warn b{display:block;margin-bottom:4px}
.amber{background:#fef6e7;border:1px solid #f0c36d;color:#8a5a00;border-radius:10px;padding:12px 16px;margin:14px 0;font-size:.85rem}
.amber b{display:block;margin-bottom:4px}
.ditl-tag{display:inline-block;padding:2px 8px;border-radius:6px;font-size:.66rem;font-weight:700;margin-left:8px;vertical-align:middle}
.ditl-tag.observed{background:#e2f5ee;color:#04785a;border:1px solid #9adfc8}
.ditl-tag.pending{background:#fef6e7;color:#8a5a00;border:1px solid #f0c36d}
.note{color:var(--mut);font-size:.8rem;margin:8px 0 16px}
h2.tier{color:var(--navy);font-size:1rem;margin:26px 0 10px;border-bottom:2px solid var(--line);padding-bottom:6px}
.tiles{display:grid;grid-template-columns:repeat(auto-fill,minmax(280px,1fr));gap:14px}
.pcard{background:#fff;border:1px solid var(--line);border-radius:12px;padding:14px 16px;cursor:pointer;transition:box-shadow .12s,transform .12s;border-top:4px solid var(--navy)}
.pcard:hover{box-shadow:0 6px 18px rgba(20,40,90,.10);transform:translateY(-1px)}
.pc-h{font-weight:700;color:var(--navy);font-size:1rem}
.pc-meta{color:var(--mut);font-size:.8rem;margin:6px 0;line-height:1.45}
.chip{display:inline-block;padding:2px 9px;border-radius:11px;color:#fff;font-size:.68rem;font-weight:600;margin:1px 3px 1px 0;white-space:nowrap}
.chip.role{background:#5a6b86}.chip.imp{background:#7A4FBE}
.pill{display:inline-block;padding:2px 8px;border-radius:6px;color:#fff;font-size:.66rem;font-weight:700}
.t3row{background:#fff;border:1px solid var(--line);border-radius:10px;padding:10px 14px;margin:6px 0;display:flex;gap:12px;align-items:baseline;flex-wrap:wrap}
.t3row b{color:var(--navy)}
#detail{display:none;position:fixed;inset:0;background:rgba(15,22,40,.55);z-index:50;overflow:auto;padding:30px 16px}
#detail .panel{max-width:980px;margin:0 auto;background:#fff;border-radius:14px;padding:24px 28px;position:relative}
#detail .close{position:absolute;top:14px;right:16px;border:none;background:var(--navy);color:#fff;border-radius:7px;padding:6px 12px;cursor:pointer;font:inherit}
.d-h{color:var(--navy);margin:0 44px 4px 0;font-size:1.2rem}
.sec{margin:14px 0}.lbl{font-size:.68rem;text-transform:uppercase;letter-spacing:.04em;color:var(--mut);font-weight:700;margin-bottom:4px}
.quote{border-left:3px solid var(--mag);padding:6px 12px;font-style:italic;color:#444;background:#fbfcfe;border-radius:0 8px 8px 0}
ul.pl{margin:4px 0;padding-left:18px}ul.pl li{margin:3px 0}
table.jt{width:100%;border-collapse:collapse;background:#fff;border:1px solid var(--line);border-radius:10px;overflow:hidden}
.jt th{background:var(--navy);color:#fff;text-align:left;padding:8px 10px;font-size:.72rem;white-space:nowrap}
.jt td{padding:8px 10px;border-top:1px solid var(--line);font-size:.8rem;vertical-align:top}
.jt td.cur{border-left:3px solid #c0392b;background:#fdf6f5}
.jt td.fut{border-left:3px solid #04A577;background:#f4fbf8}
.emo{white-space:nowrap;font-weight:600}.emo .a{color:var(--mut);padding:0 4px}
.emo .now{color:#c0392b}.emo .later{color:#04A577}
.ditl{display:grid;grid-template-columns:1fr 1fr;gap:14px}
.ditl .col{background:#fff;border:1px solid var(--line);border-radius:10px;padding:12px}
.ditl .col.before{border-top:4px solid #c0392b}.ditl .col.after{border-top:4px solid #04A577}
.ditl h4{margin:0 0 8px;font-size:.8rem;text-transform:uppercase;letter-spacing:.04em;color:var(--mut)}
.blk{display:flex;gap:10px;padding:6px 0;border-top:1px dashed var(--line);font-size:.82rem}
.blk:first-of-type{border-top:none}
.blk .t{min-width:48px;font-weight:700;color:var(--navy);font-variant-numeric:tabular-nums}
footer{text-align:center;color:var(--mut);font-size:.72rem;padding:18px}
@media(max-width:760px){.ditl{grid-template-columns:1fr}}
</style></head><body>
<header><h1>%%TITLE%%</h1><div class="sub">%%SUBTITLE%%</div></header>
<div class="wrap">
%%WARNINGS%%
<div class="note">Each persona is an archetype representing many people, not one literal individual. Click a Tier 1/2 card for the full profile, journey, and day-in-the-life.</div>
<div id="cards"></div>
</div>
<div id="detail"><div class="panel"><button class="close" onclick="closeDetail()">Close ✕</button><div id="dbody"></div></div></div>
<footer>%%FOOTER%%</footer>
<script>
const DATA=%%DATA%%, TIER_COLOR=%%TIERCOLORS%%, INT_COLOR=%%INTCOLORS%%;
const $=s=>document.querySelector(s), esc=s=>String(s==null?'':s).replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
function chipList(arr,cls){return (arr||[]).map(x=>'<span class="chip '+cls+'">'+esc(x)+'</span>').join('');}
function intensity(p){return p.change_intensity?'<span class="pill" style="background:'+(INT_COLOR[p.change_intensity]||'#888')+'">'+esc(p.change_intensity)+' intensity</span>':'';}
function card(p){
  return '<div class="pcard" style="border-top-color:'+(TIER_COLOR[p.tier]||'#888')+'" onclick="openDetail('+p.id+')">'
    +'<div class="pc-h">'+esc(p.name)+'</div>'
    +'<div class="pc-meta">'+esc(p.archetype)+'</div>'
    +'<div>'+intensity(p)+(p.population_estimate?' <span class="pill" style="background:#5a6b86">'+esc(p.population_estimate)+' people</span>':'')
    +(p.confidence?' <span class="pill" style="background:#9aa3b2">Confidence: '+esc(p.confidence)+'</span>':'')+'</div>'
    +'<div class="pc-meta" style="margin-top:8px">'+chipList(p.roles_covered,'role')+'</div></div>';
}
function build(){
  const host=$('#cards');let out='';
  [1,2,3].forEach(t=>{
    const ps=DATA.filter(p=>p.tier===t);if(!ps.length)return;
    out+='<h2 class="tier">'+esc(ps[0].tier_label)+'</h2>';
    if(t===3){out+=ps.map(p=>'<div class="t3row"><b>'+esc(p.name)+'</b><span>'+esc(p.archetype)+'</span><span>'+chipList(p.roles_covered,'role')+'</span></div>').join('');}
    else{out+='<div class="tiles">'+ps.map(card).join('')+'</div>';}
  });
  host.innerHTML=out;
}
function journeyTable(p){
  if(!p.journey.length)return '';
  return '<div class="sec"><div class="lbl">User journey — current → future</div><div style="overflow-x:auto"><table class="jt"><thead><tr><th>#</th><th>Stage</th><th>Current (pain)</th><th>Future (moment)</th><th>Emotion</th><th>Change support</th></tr></thead><tbody>'
    +p.journey.map(s=>'<tr><td>'+s.n+'</td><td><b>'+esc(s.stage)+'</b></td><td class="cur">'+esc(s.current)+'</td><td class="fut">'+esc(s.future)+'</td><td><span class="emo"><span class="now">'+esc(s.emotion_current||'—')+'</span><span class="a">→</span><span class="later">'+esc(s.emotion_future||'—')+'</span></span></td><td>'+esc(s.support)+'</td></tr>').join('')
    +'</tbody></table></div></div>';
}
function ditlCol(blocks,cls,title){
  return '<div class="col '+cls+'"><h4>'+title+'</h4>'+blocks.map(b=>'<div class="blk"><span class="t">'+esc(b.time)+'</span><span>'+esc(b.activity)+'</span></div>').join('')+'</div>';
}
function ditlTag(p){
  const s=p.day_in_life_status;if(!s)return '';
  const cls=s==='observed'?'observed':'pending';
  return '<span class="ditl-tag '+cls+'">'+esc(s)+'</span>';
}
function ditl(p){
  if(!p.ditl_before.length&&!p.ditl_after.length)return '';
  return '<div class="sec"><div class="lbl">Day in the life — before / after'+ditlTag(p)+'</div><div class="ditl">'
    +ditlCol(p.ditl_before,'before','Today')+ditlCol(p.ditl_after,'after','Future')+'</div></div>';
}
function ul(arr){return arr.length?'<ul class="pl">'+arr.map(x=>'<li>'+esc(x)+'</li>').join('')+'</ul>':'<span style="color:#bbb">—</span>';}
function openDetail(id){
  const p=DATA.find(x=>x.id===id);if(!p)return;
  $('#dbody').innerHTML='<h2 class="d-h">'+esc(p.name)+'</h2>'
    +'<div>'+'<span class="pill" style="background:'+(TIER_COLOR[p.tier]||'#888')+'">'+esc(p.tier_label)+'</span> '+intensity(p)
    +(p.population_estimate?' <span class="pill" style="background:#5a6b86">'+esc(p.population_estimate)+' people</span>':'')
    +(p.sentiment?' <span class="pill" style="background:#F08301">'+esc(p.sentiment)+'</span>':'')
    +(p.confidence?' <span class="pill" style="background:#9aa3b2">Confidence: '+esc(p.confidence)+'</span>':'')+'</div>'
    +'<div class="sec">'+esc(p.archetype)+'</div>'
    +(p.quote?'<div class="sec quote">“'+esc(p.quote)+'”</div>':'')
    +'<div class="sec"><div class="lbl">Roles covered</div>'+chipList(p.roles_covered,'role')+'</div>'
    +(p.top_impacts.length?'<div class="sec"><div class="lbl">Top change impacts (CIA)</div>'+chipList(p.top_impacts,'imp')+'</div>':'')
    +'<div class="sec"><div class="lbl">Goals</div>'+ul(p.goals)+'</div>'
    +'<div class="sec"><div class="lbl">Pain points</div>'+ul(p.pain_points)+'</div>'
    +journeyTable(p)+ditl(p);
  $('#detail').style.display='block';window.scrollTo(0,0);
}
function closeDetail(){$('#detail').style.display='none';}
document.addEventListener('keydown',e=>{if(e.key==='Escape')closeDetail();});
$('#detail').addEventListener('click',e=>{if(e.target.id==='detail')closeDetail();});
build();
</script></body></html>"""


def render_html(out_path, *, title, subtitle, personas, brand, warnings, unscored=None):
    warn_html = ""
    if unscored:
        warn_html += ('<div class="amber"><b>UNSCORED-CIA WARNING</b>'
                      + html.escape(unscored) + "</div>")
    if warnings:
        warn_html += ('<div class="warn"><b>CIA cross-check flags — resolve before delivering</b>'
                      + "".join(f"<div>• {html.escape(w)}</div>" for w in warnings) + "</div>")
    out = (TEMPLATE
           .replace("%%TITLE%%", html.escape(title))
           .replace("%%SUBTITLE%%", html.escape(subtitle))
           .replace("%%NAVY%%", brand["navy"]).replace("%%MAG%%", brand["magenta"])
           .replace("%%FONT%%", brand["font"]).replace("%%FOOTER%%", html.escape(brand["footer"]))
           .replace("%%WARNINGS%%", warn_html)
           .replace("%%DATA%%", json.dumps(personas, ensure_ascii=False))
           .replace("%%TIERCOLORS%%", json.dumps(TIER_COLOR))
           .replace("%%INTCOLORS%%", json.dumps(INTENSITY_COLOR)))
    open(out_path, "w", encoding="utf-8").write(out)


# ----------------------------------------------------------------------------- main
def _guard(outdir, names, force):
    """Refuse to overwrite existing deliverables unless --force is set."""
    existing = [n for n in names if os.path.exists(os.path.join(outdir, n))]
    if existing and not force:
        sys.exit("Refusing to overwrite existing output(s) in %r:\n  %s\n"
                 "Re-run with --force to overwrite, or choose a different --outdir."
                 % (outdir, "\n  ".join(existing)))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--records", required=True, help="personas.json")
    ap.add_argument("--config", required=True, help="project.json")
    ap.add_argument("--cia", help="optional cia_records.json for the role cross-check")
    ap.add_argument("--outdir", default=".")
    ap.add_argument("--force", action="store_true", help="overwrite existing outputs in --outdir")
    a = ap.parse_args()
    os.makedirs(a.outdir, exist_ok=True)
    project = json.load(open(a.config, encoding="utf-8"))
    pname = project.get("project_name", "Project")
    xlsx_name = f"{pname} Personas.xlsx"
    html_name = f"{pname} Personas & Journeys.html"
    _guard(a.outdir, [xlsx_name, html_name], a.force)

    raw = json.load(open(a.records, encoding="utf-8"))
    aliases = {}
    if isinstance(raw, dict):  # {"role_aliases": {...}, "personas": [...]}
        aliases = raw.get("role_aliases") or {}
        raw = raw.get("personas") or []
    personas = normalize(raw)
    brand = {**DEFAULT_BRAND, **project.get("brand", {})}

    warnings, unscored = [], None
    if a.cia:
        warnings, unscored = cross_check(personas, json.load(open(a.cia, encoding="utf-8")), aliases)
        if unscored:
            print("WARNING:", unscored)
        for w in warnings:
            print("WARN:", w)

    write_xlsx(os.path.join(a.outdir, xlsx_name), personas, brand["navy"])
    t1 = sum(1 for p in personas if p["tier"] == 1)
    roles = sum(len(p["roles_covered"]) for p in personas)
    render_html(os.path.join(a.outdir, html_name),
                title=f"{pname} — Personas, Journeys & Day in the Life",
                subtitle=f"{len(personas)} personas ({t1} Tier 1) covering {roles} roles",
                personas=personas, brand=brand, warnings=warnings, unscored=unscored)
    print(f"Personas: {len(personas)} ({t1} Tier 1) / {roles} roles -> {a.outdir}")


if __name__ == "__main__":
    main()
