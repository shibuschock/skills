#!/usr/bin/env python3
"""
change-network-builder — render a Change Champion Network package as an xlsx
workbook (roster + coverage) + a self-contained interactive HTML dashboard,
for ANY project, from a structured network plan (designed by Claude).

Methodology baked in (portable OCM):
  - Tiered network (sponsor / change lead / champion) with selection criteria,
    time commitment by phase, cadence, recognition, nomination workflow.
  - Coverage discipline: Open seats are tracked, never invented; coverage by
    BU/location is derived, not hand-computed.

Usage:
  python3 network_render.py --plan network_plan.json --config project.json --outdir OUT

Plan = a JSON object with "network", "cadence", "roster" (see references/SCHEMA.md).
Only stdlib + openpyxl. Cross-platform. Output HTML has no external dependencies.

Render scaffolding pattern shared with the OCM suite (canonical source: cia-builder).
"""
import argparse, html, json, os, sys

STATUSES = ["Open", "Candidate", "Nominated", "Confirmed", "Onboarded", "Departed"]
STATUS_COLORS = {"Open": "#b40020", "Candidate": "#7E57C2", "Nominated": "#F08301",
                 "Confirmed": "#1d6fb8", "Onboarded": "#04A577", "Departed": "#9aa3b2"}
DEFAULT_BRAND = {"navy": "#162B75", "magenta": "#EE2C81", "orange": "#F08301",
                 "teal": "#04A577", "coral": "#FF533C", "font": "Calibri, system-ui, sans-serif",
                 "footer": "Confidential"}
FILLED = ("Nominated", "Confirmed", "Onboarded")

def g(rec, *keys, default=""):
    for k in keys:
        if k in rec and rec[k] not in (None, ""):
            return rec[k]
    return default

def as_int(v):
    try: return int(float(v))
    except Exception: return None

# ----------------------------------------------------------------------------- normalize
def normalize_roster(roster):
    out = []
    for i, r in enumerate(roster, 1):
        status = g(r, "status", default="Open")
        if status not in STATUSES:
            status = "Open"
        out.append({
            "id": r.get("id", i),
            "name": g(r, "name"),
            "role_title": g(r, "role_title"),
            "tier": g(r, "tier", default="Champion"),
            "business_unit": g(r, "business_unit", "bu"),
            "location": g(r, "location"),
            "team": g(r, "team"),
            "email": g(r, "email"),
            "status": status,
            "headcount_covered": as_int(g(r, "headcount_covered")) or "",
            "nominated_by": g(r, "nominated_by"),
            "notes": g(r, "notes"),
            "filled": status in FILLED,
        })
    return out

def coverage_rows(data, key):
    """Aggregate seats by a roster key (business_unit or location)."""
    groups = {}
    for d in data:
        k = d[key] or "(unassigned)"
        grp = groups.setdefault(k, {"seats": 0, "filled": 0, "open": 0, "candidate": 0, "onboarded": 0, "headcount": 0})
        grp["seats"] += 1
        if d["filled"]: grp["filled"] += 1
        if d["status"] == "Open": grp["open"] += 1
        if d["status"] == "Candidate": grp["candidate"] += 1
        if d["status"] == "Onboarded": grp["onboarded"] += 1
        if isinstance(d["headcount_covered"], int): grp["headcount"] += d["headcount_covered"]
    rows = []
    for k in sorted(groups):
        v = groups[k]
        champs = sum(1 for d in data if (d[key] or "(unassigned)") == k and d["filled"] and d["tier"].lower() not in ("sponsor",))
        ratio = f"1:{round(v['headcount']/champs)}" if champs and v["headcount"] else ""
        rows.append({key: k, "seats": v["seats"], "filled": v["filled"], "open": v["open"],
                     "candidate": v["candidate"],
                     "onboarded": v["onboarded"], "pct_filled": round(100*v["filled"]/v["seats"]) if v["seats"] else 0,
                     "headcount": v["headcount"] or "", "ratio": ratio,
                     "gap": "GAP" if v["open"] or v["candidate"] else ""})
    return rows

# ----------------------------------------------------------------------------- xlsx
def style_sheet(ws, columns, navy):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    thin = Side(style="thin", color="D9D9D9"); border = Border(thin, thin, thin, thin)
    for ci, _ in enumerate(columns, 1):
        c = ws.cell(1, ci); c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=navy.lstrip("#"))
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = border
    ws.freeze_panes = "A2"; ws.row_dimensions[1].height = 26
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
    for ci, col in enumerate(columns, 1):
        L = get_column_letter(ci); ws.column_dimensions[L].width = col[2] if len(col) > 2 else 20
        for cell in ws[L][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top"); cell.border = border
            cell.font = Font(size=10)

def write_xlsx(path, roster, cov_bu, cov_loc, navy):
    from openpyxl import Workbook
    wb = Workbook()
    r_cols = [("id", "ID", 6), ("name", "Name", 22), ("role_title", "Role / Title", 24),
              ("tier", "Tier", 14), ("business_unit", "Business Unit", 18),
              ("location", "Location", 18), ("team", "Team / Group", 22),
              ("email", "Email", 26), ("status", "Status", 12),
              ("headcount_covered", "Headcount Covered", 12),
              ("nominated_by", "Nominated By", 20), ("notes", "Notes", 34)]
    ws = wb.active; ws.title = "Roster"
    ws.append([c[1] for c in r_cols])
    for r in roster:
        ws.append([r.get(c[0], "") for c in r_cols])
    style_sheet(ws, r_cols, navy)

    for title, key, rows in (("Coverage by BU", "business_unit", cov_bu),
                             ("Coverage by Location", "location", cov_loc)):
        c_cols = [(key, title.replace("Coverage by ", ""), 24), ("seats", "Seats", 9),
                  ("filled", "Filled", 9), ("open", "Open", 9), ("candidate", "Candidate", 11),
                  ("onboarded", "Onboarded", 11),
                  ("pct_filled", "% Filled", 10), ("headcount", "Headcount Covered", 12),
                  ("ratio", "Champion Ratio", 14), ("gap", "Gap Flag", 10)]
        ws2 = wb.create_sheet(title[:31])
        ws2.append([c[1] for c in c_cols])
        for r in rows:
            ws2.append([r.get(c[0], "") for c in c_cols])
        style_sheet(ws2, c_cols, navy)
    wb.save(path)

# ----------------------------------------------------------------------------- HTML
TEMPLATE = r"""<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0"><title>%%TITLE%%</title>
<style>
:root{--navy:%%NAVY%%;--mag:%%MAG%%;--ink:#1d2433;--mut:#6b7280;--line:#e5e7eb;--bg:#f5f6f8}
*{box-sizing:border-box}body{margin:0;font-family:%%FONT%%;color:var(--ink);background:var(--bg);font-size:14px}
header{background:linear-gradient(135deg,var(--navy),#0E2841);color:#fff;padding:18px 28px}
header h1{margin:0;font-size:1.15rem;font-weight:700}header .sub{opacity:.85;font-size:.85rem;margin-top:3px}
.wrap{max-width:1400px;margin:0 auto;padding:18px 28px 80px}
.stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:14px;margin:16px 0}
.card{background:#fff;border:1px solid var(--line);border-radius:10px;padding:14px 16px}
.card .v{font-size:1.6rem;font-weight:800;color:var(--navy)}.card .v.warn{color:#b40020}
.card .l{font-size:.72rem;color:var(--mut);text-transform:uppercase;letter-spacing:.04em;margin-top:2px}
.bar{display:flex;flex-wrap:wrap;gap:8px;align-items:center;margin:10px 0}
.bar select,.bar input{padding:7px 10px;border:1px solid var(--line);border-radius:7px;font:inherit;background:#fff}
.bar input{flex:1;min-width:180px}
.tabs{display:flex;gap:4px;border-bottom:2px solid var(--line);margin-top:8px;flex-wrap:wrap}
.tab{padding:8px 14px;cursor:pointer;border-radius:7px 7px 0 0;font-weight:600;color:var(--mut)}
.tab.active{color:var(--navy);background:#fff;border:1px solid var(--line);border-bottom:2px solid #fff;margin-bottom:-2px}
.view{display:none}.view.active{display:block}
table{width:100%;border-collapse:collapse;background:#fff;border:1px solid var(--line);border-radius:10px;overflow:hidden}
thead th{position:sticky;top:0;background:var(--navy);color:#fff;text-align:left;padding:9px 11px;font-size:.74rem;white-space:nowrap;cursor:pointer}
tbody td{padding:9px 11px;border-top:1px solid var(--line);font-size:.82rem;vertical-align:top}
tbody tr:hover{background:#f0f3fb}
.statc{display:inline-block;padding:2px 8px;border-radius:6px;color:#fff;font-size:.66rem;font-weight:700;white-space:nowrap}
.gapflag{display:inline-block;padding:1px 7px;border-radius:9px;background:#fdecea;color:#b40020;font-size:.62rem;font-weight:700}
.open-name{color:#b40020;font-style:italic;font-weight:600}
.covbar{background:#eef0f4;border-radius:6px;height:12px;overflow:hidden;min-width:110px}
.covbar>div{height:100%;background:var(--navy)}
h3.sec{color:var(--navy);margin:20px 0 8px}
.panel{background:#fff;border:1px solid var(--line);border-radius:10px;padding:14px 18px;margin:10px 0}
.panel h4{margin:0 0 6px;color:var(--navy)}.panel .mut{color:var(--mut);font-size:.82rem}
.tiergrid{display:grid;grid-template-columns:repeat(auto-fit,minmax(280px,1fr));gap:14px}
.tier{background:#fff;border:1px solid var(--line);border-left:5px solid var(--navy);border-radius:10px;padding:14px 16px}
.tier h4{margin:0;color:var(--navy)}.tier .who{color:var(--mut);font-size:.8rem;margin:4px 0}
.tier .tc{display:inline-block;margin-top:8px;padding:2px 9px;border-radius:11px;background:var(--mag);color:#fff;font-size:.68rem;font-weight:700}
ul.plain{margin:6px 0;padding-left:20px}ul.plain li{margin:4px 0}
.calgrid{display:grid;grid-template-columns:repeat(auto-fit,minmax(250px,1fr));gap:14px}
.calcol{background:#fff;border:1px solid var(--line);border-radius:10px;padding:12px 14px}
.calcol h4{margin:0 0 8px;color:#fff;background:var(--navy);border-radius:6px;padding:5px 10px;font-size:.8rem;display:inline-block}
.mtg{border-left:3px solid var(--mag);padding:6px 10px;margin:8px 0;background:#fbfcfe;border-radius:0 6px 6px 0}
.mtg b{display:block}.mtg .m{color:var(--mut);font-size:.76rem;margin-top:2px}
footer{text-align:center;color:var(--mut);font-size:.72rem;padding:18px}
</style></head><body>
<header><h1>%%TITLE%%</h1><div class="sub">%%SUBTITLE%%</div></header>
<div class="wrap">
  <div class="stats" id="stats"></div>
  <div class="tabs" id="tabs"></div>
  <div class="view active" id="v-roster">
    <div class="bar" id="bar"></div>
    <div style="overflow:auto"><table><thead><tr id="head"></tr></thead><tbody id="body"></tbody></table></div>
  </div>
  <div class="view" id="v-coverage">
    <h3 class="sec">Coverage by business unit</h3>
    <div style="overflow:auto"><table><thead><tr id="cov-bu-head"></tr></thead><tbody id="cov-bu"></tbody></table></div>
    <h3 class="sec">Coverage by location</h3>
    <div style="overflow:auto"><table><thead><tr id="cov-loc-head"></tr></thead><tbody id="cov-loc"></tbody></table></div>
  </div>
  <div class="view" id="v-cadence"><h3 class="sec">Operating rhythm</h3><div class="calgrid" id="cal"></div></div>
  <div class="view" id="v-design" ></div>
</div>
<footer>%%FOOTER%%</footer>
<script>
const ROSTER=%%ROSTER%%, COVBU=%%COVBU%%, COVLOC=%%COVLOC%%, NET=%%NET%%, CAD=%%CADENCE%%, SC=%%STATUSCOLORS%%;
const $=s=>document.querySelector(s), esc=s=>String(s==null?'':s).replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
let sortKey=null, sortDir=1;
const COLS=[["id","ID"],["name","Name"],["role_title","Role / Title"],["tier","Tier"],["business_unit","Business Unit"],["location","Location"],["team","Team"],["status","Status"],["headcount_covered","Headcount"],["nominated_by","Nominated By"],["notes","Notes"]];
const FILTERS=[["tier","All tiers"],["business_unit","All BUs"],["location","All locations"],["status","All statuses"]];
function distinct(k){return [...new Set(ROSTER.map(d=>d[k]).filter(v=>v!==''&&v!=null))].sort();}
function buildBar(){
  const bar=$('#bar');
  FILTERS.forEach(([k,label])=>{const s=document.createElement('select');s.id='f-'+k;
    s.innerHTML='<option value="">'+esc(label)+'</option>'+distinct(k).map(v=>'<option>'+esc(v)+'</option>').join('');
    s.onchange=render;bar.appendChild(s);});
  const q=document.createElement('input');q.id='q';q.placeholder='Search roster…';q.oninput=render;bar.appendChild(q);
}
function filtered(){
  const q=($('#q').value||'').toLowerCase();
  return ROSTER.filter(d=>FILTERS.every(([k])=>{const v=$('#f-'+k).value;return !v||d[k]===v;})
    && (!q||COLS.some(([k])=>String(d[k]??'').toLowerCase().includes(q))));
}
function render(){
  let rows=filtered();
  if(sortKey)rows=[...rows].sort((a,b)=>((a[sortKey]??'')>(b[sortKey]??'')?1:-1)*sortDir);
  $('#head').innerHTML=COLS.map(([k,l])=>'<th data-k="'+k+'">'+esc(l)+(sortKey===k?(sortDir>0?' ▲':' ▼'):'')+'</th>').join('');
  document.querySelectorAll('#head th').forEach(th=>th.onclick=()=>{const k=th.dataset.k;sortDir=sortKey===k?-sortDir:1;sortKey=k;render();});
  $('#body').innerHTML=rows.map(d=>'<tr>'+COLS.map(([k])=>{
    if(k==='status')return '<td><span class="statc" style="background:'+(SC[d.status]||'#6b7280')+'">'+esc(d.status)+'</span></td>';
    if(k==='name'&&!d.name)return '<td><span class="open-name">OPEN — needs nominee</span></td>';
    return '<td>'+esc(d[k])+'</td>';}).join('')+'</tr>').join('')
    ||'<tr><td colspan="'+COLS.length+'">No seats match the filters.</td></tr>';
}
function tiles(){
  const seats=ROSTER.length, filled=ROSTER.filter(d=>d.filled).length, open=ROSTER.filter(d=>d.status==='Open').length;
  const cand=ROSTER.filter(d=>d.status==='Candidate').length;
  const onb=ROSTER.filter(d=>d.status==='Onboarded').length;
  const hc=ROSTER.reduce((s,d)=>s+(typeof d.headcount_covered==='number'?d.headcount_covered:0),0);
  const t=[["Seats",seats],["Filled",filled],["Open",open,open>0],["Candidate",cand],["Onboarded",onb],["% Filled",seats?Math.round(100*filled/seats)+'%':'—'],["Headcount covered",hc||'—']];
  $('#stats').innerHTML=t.map(([l,v,w])=>'<div class="card"><div class="v'+(w?' warn':'')+'">'+v+'</div><div class="l">'+esc(l)+'</div></div>').join('');
}
function covTable(rows,key,headId,bodyId,label){
  const cols=[[key,label],["seats","Seats"],["filled","Filled"],["open","Open"],["candidate","Candidate"],["onboarded","Onboarded"],["pct_filled","% Filled"],["headcount","Headcount"],["ratio","Ratio"],["gap",""]];
  $(headId).innerHTML=cols.map(c=>'<th>'+esc(c[1])+'</th>').join('');
  $(bodyId).innerHTML=rows.map(r=>'<tr>'+cols.map(([k])=>{
    if(k==='pct_filled')return '<td><div class="covbar"><div style="width:'+r.pct_filled+'%"></div></div>'+r.pct_filled+'%</td>';
    if(k==='gap')return '<td>'+(r.gap?'<span class="gapflag">GAP</span>':'')+'</td>';
    return '<td>'+esc(r[k])+'</td>';}).join('')+'</tr>').join('');
}
function cadence(){
  const order=["Weekly","Bi-weekly","Monthly","Quarterly","After each briefing","Always on"];
  const freqs=[...new Set(CAD.map(c=>c.frequency))].sort((a,b)=>{const i=order.indexOf(a),j=order.indexOf(b);return (i<0?99:i)-(j<0?99:j);});
  $('#cal').innerHTML=freqs.map(f=>'<div class="calcol"><h4>'+esc(f)+'</h4>'+CAD.filter(c=>c.frequency===f).map(c=>
    '<div class="mtg"><b>'+esc(c.name)+(c.duration?' · '+esc(c.duration):'')+'</b><div class="m">'+esc(c.purpose||'')+
    (c.audience?'<br>Audience: '+esc(c.audience):'')+(c.owner?' · Owner: '+esc(c.owner):'')+'</div></div>').join('')+'</div>').join('');
}
function design(){
  let h='<h3 class="sec">'+esc(NET.name||'Change Network')+'</h3>';
  if(NET.purpose)h+='<div class="panel"><h4>Purpose</h4><div>'+esc(NET.purpose)+'</div></div>';
  if(NET.tiers&&NET.tiers.length)h+='<h3 class="sec">Tiers</h3><div class="tiergrid">'+NET.tiers.map(t=>
    '<div class="tier"><h4>'+esc(t.tier)+'</h4><div class="who">'+esc(t.who||'')+'</div><div>'+esc(t.role||'')+'</div>'+
    (t.time_commitment?'<span class="tc">'+esc(t.time_commitment)+'</span>':'')+'</div>').join('')+'</div>';
  if(NET.sizing_basis)h+='<div class="panel"><h4>Sizing basis</h4><div>'+esc(NET.sizing_basis)+'</div></div>';
  if(NET.selection_criteria&&NET.selection_criteria.length)h+='<div class="panel"><h4>Selection criteria</h4><ul class="plain">'+NET.selection_criteria.map(c=>'<li>'+esc(c)+'</li>').join('')+'</ul></div>';
  if(NET.time_by_phase&&NET.time_by_phase.length)h+='<div class="panel"><h4>Time commitment by phase</h4><ul class="plain">'+NET.time_by_phase.map(p=>'<li><b>'+esc(p.phase)+'</b> — '+esc(p.time)+' · '+esc(p.activities||'')+'</li>').join('')+'</ul></div>';
  if(NET.recognition&&NET.recognition.length)h+='<div class="panel"><h4>Recognition &amp; sustainment</h4><ul class="plain">'+NET.recognition.map(c=>'<li>'+esc(c)+'</li>').join('')+'</ul></div>';
  if(NET.nomination){const n=NET.nomination;h+='<div class="panel"><h4>Nomination workflow</h4><div class="mut">Nominated by: '+esc(n.nominated_by||'—')+(n.deadline?' · Deadline: '+esc(n.deadline):'')+'</div>'+
    (n.workflow&&n.workflow.length?'<ol class="plain">'+n.workflow.map(w=>'<li>'+esc(w)+'</li>').join('')+'</ol>':'')+'</div>';}
  $('#v-design').innerHTML=h;
}
function buildTabs(){
  const tabs=[["v-roster","Roster"],["v-coverage","Coverage"],["v-cadence","Cadence"],["v-design","Network Design"]];
  $('#tabs').innerHTML=tabs.map(([id,l],i)=>'<div class="tab'+(i===0?' active':'')+'" data-v="'+id+'">'+esc(l)+'</div>').join('');
  document.querySelectorAll('.tab').forEach(t=>t.onclick=()=>{
    document.querySelectorAll('.tab').forEach(x=>x.classList.remove('active'));
    document.querySelectorAll('.view').forEach(x=>x.classList.remove('active'));
    t.classList.add('active');document.getElementById(t.dataset.v).classList.add('active');});
}
tiles();buildTabs();buildBar();render();
covTable(COVBU,'business_unit','#cov-bu-head','#cov-bu','Business Unit');
covTable(COVLOC,'location','#cov-loc-head','#cov-loc','Location');
cadence();design();
</script></body></html>"""

def render_html(out_path, *, title, subtitle, roster, cov_bu, cov_loc, network, cadence, brand):
    out = (TEMPLATE
        .replace("%%TITLE%%", html.escape(title))
        .replace("%%SUBTITLE%%", html.escape(subtitle))
        .replace("%%NAVY%%", brand["navy"]).replace("%%MAG%%", brand["magenta"])
        .replace("%%FONT%%", brand["font"]).replace("%%FOOTER%%", html.escape(brand["footer"]))
        .replace("%%ROSTER%%", json.dumps(roster, ensure_ascii=False))
        .replace("%%COVBU%%", json.dumps(cov_bu, ensure_ascii=False))
        .replace("%%COVLOC%%", json.dumps(cov_loc, ensure_ascii=False))
        .replace("%%NET%%", json.dumps(network, ensure_ascii=False))
        .replace("%%CADENCE%%", json.dumps(cadence, ensure_ascii=False))
        .replace("%%STATUSCOLORS%%", json.dumps(STATUS_COLORS, ensure_ascii=False)))
    open(out_path, "w", encoding="utf-8").write(out)

# ----------------------------------------------------------------------------- main
def _guard(outdir, names, force):
    """Refuse to overwrite existing deliverables unless --force is set."""
    if force: return
    clashes = [n for n in names if os.path.exists(os.path.join(outdir, n))]
    if clashes:
        sys.exit("Refusing to overwrite existing output(s): " + ", ".join(clashes) +
                 "\nRe-run with --force to overwrite, or choose a different --outdir.")

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--plan", required=True, help="network_plan.json")
    ap.add_argument("--config", required=True, help="project.json")
    ap.add_argument("--outdir", required=True)
    ap.add_argument("--force", action="store_true", help="overwrite existing outputs in --outdir")
    a = ap.parse_args()

    plan = json.load(open(a.plan, encoding="utf-8"))
    project = json.load(open(a.config, encoding="utf-8"))
    brand = {**DEFAULT_BRAND, **project.get("brand", {})}
    pname = project.get("project_name", "Project")
    os.makedirs(a.outdir, exist_ok=True)
    xlsx_name = f"{pname} Change Network Roster.xlsx"
    html_name = f"{pname} Change Network Dashboard.html"
    _guard(a.outdir, [xlsx_name, html_name], a.force)

    roster = normalize_roster(plan.get("roster") or [])
    network = plan.get("network") or {}
    cadence = plan.get("cadence") or []
    cov_bu = coverage_rows(roster, "business_unit")
    cov_loc = coverage_rows(roster, "location")

    # sanity: BUs in project.json with no roster coverage
    missing = [b for b in project.get("business_units", [])
               if not any(d["business_unit"] == b for d in roster)]
    if missing:
        print("WARN: business unit(s) with no roster seats: " + ", ".join(missing))

    write_xlsx(os.path.join(a.outdir, xlsx_name), roster, cov_bu, cov_loc, brand["navy"])
    filled = sum(1 for d in roster if d["filled"])
    subtitle = (f"{len(roster)} seats · {filled} filled · "
                f"{sum(1 for d in roster if d['status']=='Open')} open · "
                f"{network.get('name', 'Change Network')}")
    render_html(os.path.join(a.outdir, html_name),
                title=f"{pname} — Change Network Dashboard", subtitle=subtitle,
                roster=roster, cov_bu=cov_bu, cov_loc=cov_loc,
                network=network, cadence=cadence, brand=brand)
    print(f"Wrote: {os.path.join(a.outdir, xlsx_name)}")
    print(f"Wrote: {os.path.join(a.outdir, html_name)}")

if __name__ == "__main__":
    main()
