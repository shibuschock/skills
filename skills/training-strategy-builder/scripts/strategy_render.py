#!/usr/bin/env python3
"""
training-strategy-builder — render a Training Strategy as an xlsx register +
a self-contained HTML strategy document, for ANY change program, from a
structured plan (training_strategy.json drafted by Claude).

Methodology baked in (portable OCM/L&D): guiding principles with implications,
modality strategy by workforce type, governance/ownership, resourcing
(core team + multipliers + vendor stance), build-vs-buy, Kirkpatrick
measurement framing, phasing relative to go-live, and a strict no-fabrication
discipline (open questions instead of invented answers).

Usage:
  python3 strategy_render.py --plan training_strategy.json --config project.json --outdir OUT

Only stdlib + openpyxl. Output HTML has no external dependencies.
Render scaffolding conventions follow cia-builder/scripts/cia_render.py (canonical source).
"""
import argparse, datetime, html, json, os, sys

DEFAULT_BRAND = {"navy": "#162B75", "magenta": "#EE2C81", "orange": "#F08301",
                 "teal": "#04A577", "coral": "#FF533C", "font": "Calibri, system-ui, sans-serif",
                 "footer": "Confidential"}

def e(s):
    return html.escape(str(s if s is not None else ""))

def lst(v):
    if not v: return []
    if isinstance(v, (list, tuple)): return [str(x) for x in v if str(x).strip()]
    return [str(v)]

# ----------------------------------------------------------------------------- xlsx
def write_xlsx(path, brand, plan):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    navy = brand["navy"].lstrip("#")
    thin = Side(style="thin", color="D9D9D9"); border = Border(thin, thin, thin, thin)
    wb = Workbook(); wb.remove(wb.active)

    def sheet(name, columns, rows):
        ws = wb.create_sheet(name[:31])
        ws.append([c[0] for c in columns])
        for ci, _ in enumerate(columns, 1):
            c = ws.cell(1, ci); c.font = Font(bold=True, color="FFFFFF", size=11)
            c.fill = PatternFill("solid", fgColor=navy)
            c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = border
        ws.freeze_panes = "A2"; ws.row_dimensions[1].height = 24
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
        for r in rows:
            ws.append([", ".join(v) if isinstance(v, list) else (v or "") for v in r])
        for ci, col in enumerate(columns, 1):
            L = get_column_letter(ci); ws.column_dimensions[L].width = col[1]
            for cell in ws[L][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = border; cell.font = Font(size=10)

    meta = plan.get("meta", {})
    sheet("Strategy on a Page",
          [("Section", 22), ("Content", 90)],
          [("Program", meta.get("program", "")),
           ("Workforce types", ", ".join(lst(meta.get("workforce_types")))),
           ("Go-live reference", meta.get("go_live_reference", "")),
           ("Principles", "; ".join(p.get("principle", "") for p in plan.get("principles", []))),
           ("Resourcing — core team", plan.get("resourcing", {}).get("core_team", "")),
           ("Resourcing — multipliers", plan.get("resourcing", {}).get("multipliers", "")),
           ("Resourcing — vendor stance", plan.get("resourcing", {}).get("vendor_stance", "")),
           ("Measurement approach", plan.get("measurement", {}).get("approach", "")),
           ("Open questions", str(len(plan.get("open_questions", []))))])
    sheet("Principles",
          [("Principle", 22), ("Statement", 50), ("Implication (so we will...)", 50)],
          [(p.get("principle", ""), p.get("statement", ""), p.get("implication", ""))
           for p in plan.get("principles", [])])
    sheet("Modality Matrix",
          [("Workforce type", 24), ("Primary modalities", 40), ("Secondary modalities", 30), ("Rationale", 55)],
          [(m.get("workforce_type", ""), lst(m.get("primary_modalities")),
            lst(m.get("secondary_modalities")), m.get("rationale", ""))
           for m in plan.get("modality_strategy", [])])
    sheet("Governance",
          [("Role", 28), ("Owns", 60), ("Named person", 24)],
          [(gv.get("role", ""), gv.get("owns", ""), gv.get("named_person", ""))
           for gv in plan.get("governance", [])])
    sheet("Build vs Buy",
          [("Category", 24), ("Stance", 20), ("Rationale", 60)],
          [(b.get("category", ""), b.get("stance", ""), b.get("rationale", ""))
           for b in plan.get("build_vs_buy", [])])
    sheet("Phasing",
          [("Phase", 26), ("Timing (vs go-live)", 22), ("Description", 60)],
          [(p.get("phase", ""), p.get("timing", ""), p.get("description", ""))
           for p in plan.get("phasing", [])])
    sheet("Risks",
          [("Risk", 50), ("Mitigation", 50), ("Owner", 24)],
          [(r.get("risk", ""), r.get("mitigation", ""), r.get("owner", ""))
           for r in plan.get("risks", [])])
    sheet("Assumptions & Open Qs",
          [("Type", 16), ("Item", 60), ("Needed from", 24), ("Needed by", 24)],
          [("Assumption", a, "", "") for a in plan.get("assumptions", [])] +
          [("Open question", q.get("question", ""), q.get("needed_from", ""), q.get("needed_by", ""))
           for q in plan.get("open_questions", [])])
    wb.save(path)

# ----------------------------------------------------------------------------- html
def build_html(plan, project, brand):
    meta = plan.get("meta", {})
    program = meta.get("program") or project.get("project_name", "Project")
    title = f"{program} Training Strategy"

    def tiles(items):
        out = []
        for p in items:
            out.append(f'<div class="tile"><h3>{e(p.get("principle"))}</h3>'
                       f'<p>{e(p.get("statement"))}</p>'
                       f'<p class="imp"><b>So we will:</b> {e(p.get("implication"))}</p></div>')
        return "".join(out)

    def modality_rows():
        out = []
        for m in plan.get("modality_strategy", []):
            prim = "".join(f'<span class="chip prim">{e(x)}</span>' for x in lst(m.get("primary_modalities")))
            sec = "".join(f'<span class="chip">{e(x)}</span>' for x in lst(m.get("secondary_modalities")))
            out.append(f'<tr><td class="wt">{e(m.get("workforce_type"))}</td>'
                       f'<td>{prim}</td><td>{sec or "&mdash;"}</td><td>{e(m.get("rationale"))}</td></tr>')
        return "".join(out)

    def gov_rows():
        out = []
        for gv in plan.get("governance", []):
            person = f' <span class="mut">({e(gv["named_person"])})</span>' if gv.get("named_person") else ""
            out.append(f'<tr><td class="wt">{e(gv.get("role"))}{person}</td><td>{e(gv.get("owns"))}</td></tr>')
        return "".join(out)

    def bvb_rows():
        return "".join(f'<tr><td class="wt">{e(b.get("category"))}</td>'
                       f'<td><span class="chip prim">{e(b.get("stance"))}</span></td>'
                       f'<td>{e(b.get("rationale"))}</td></tr>'
                       for b in plan.get("build_vs_buy", []))

    def phase_steps():
        out = []
        for i, p in enumerate(plan.get("phasing", []), 1):
            out.append(f'<div class="step"><div class="dot">{i}</div>'
                       f'<div class="stx"><b>{e(p.get("phase"))}</b>'
                       f'<span class="tim">{e(p.get("timing"))}</span>'
                       f'<p>{e(p.get("description"))}</p></div></div>')
        return "".join(out)

    def risk_rows():
        return "".join(f'<tr><td>{e(r.get("risk"))}</td><td>{e(r.get("mitigation"))}</td>'
                       f'<td class="wt">{e(r.get("owner"))}</td></tr>' for r in plan.get("risks", []))

    def li(items):
        return "".join(f"<li>{e(x)}</li>" for x in items)

    def oq_items():
        out = []
        for q in plan.get("open_questions", []):
            who = f' <span class="mut">— needed from {e(q["needed_from"])}' if q.get("needed_from") else ""
            if who and q.get("needed_by"): who += f', by {e(q["needed_by"])}'
            if who: who += "</span>"
            out.append(f'<li><b>?</b> {e(q.get("question"))}{who}</li>')
        return "".join(out)

    res = plan.get("resourcing", {})
    meas = plan.get("measurement", {})
    wtypes = ", ".join(lst(meta.get("workforce_types")))
    today = datetime.date.today().isoformat()

    return f"""<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0"><title>{e(title)}</title>
<style>
:root{{--navy:{brand["navy"]};--mag:{brand["magenta"]};--teal:{brand["teal"]};--orange:{brand["orange"]};
--ink:#1d2433;--mut:#6b7280;--line:#e5e7eb;--bg:#f5f6f8}}
*{{box-sizing:border-box;margin:0}}
body{{font-family:{brand["font"]};background:var(--bg);color:var(--ink);font-size:14px}}
header{{background:var(--navy);color:#fff;padding:26px 34px}}
header h1{{font-size:24px;font-weight:700}} header p{{opacity:.85;margin-top:6px}}
header .k{{display:inline-block;margin-right:24px;margin-top:10px;font-size:12.5px;opacity:.9}}
header .k b{{color:#fff;opacity:1}}
main{{max-width:1100px;margin:0 auto;padding:26px 34px 60px}}
section{{margin-top:34px}}
h2{{font-size:17px;color:var(--navy);border-left:4px solid var(--mag);padding-left:10px;margin-bottom:14px}}
.grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(300px,1fr));gap:14px}}
.tile{{background:#fff;border:1px solid var(--line);border-radius:10px;padding:16px;border-top:3px solid var(--navy)}}
.tile h3{{font-size:14.5px;color:var(--navy);margin-bottom:6px}}
.tile p{{font-size:13px;color:var(--ink);line-height:1.45}}
.tile .imp{{margin-top:8px;color:var(--mut)}} .tile .imp b{{color:var(--mag)}}
table{{width:100%;border-collapse:collapse;background:#fff;border:1px solid var(--line);border-radius:10px;overflow:hidden}}
th{{background:var(--navy);color:#fff;text-align:left;padding:9px 12px;font-size:12.5px}}
td{{padding:10px 12px;border-top:1px solid var(--line);vertical-align:top;line-height:1.45}}
td.wt{{font-weight:600;color:var(--navy);white-space:nowrap}}
.chip{{display:inline-block;background:#eef0f4;border-radius:999px;padding:2px 10px;margin:2px 4px 2px 0;font-size:12px}}
.chip.prim{{background:var(--navy);color:#fff}}
.mut{{color:var(--mut);font-weight:400;font-size:12.5px}}
.panel{{background:#fff;border:1px solid var(--line);border-radius:10px;padding:16px 18px}}
.panel h3{{font-size:13.5px;color:var(--navy);margin:10px 0 4px}} .panel h3:first-child{{margin-top:0}}
.steps{{display:flex;flex-wrap:wrap;gap:14px}}
.step{{flex:1 1 220px;background:#fff;border:1px solid var(--line);border-radius:10px;padding:14px;display:flex;gap:12px}}
.dot{{flex:0 0 30px;height:30px;border-radius:50%;background:var(--mag);color:#fff;display:flex;align-items:center;justify-content:center;font-weight:700}}
.stx b{{color:var(--navy)}} .stx .tim{{display:block;font-size:12px;color:var(--orange);font-weight:600;margin:2px 0 4px}}
.stx p{{font-size:12.5px;color:var(--mut);line-height:1.4}}
ul.plain{{list-style:none;padding:0}} ul.plain li{{padding:7px 0;border-bottom:1px solid var(--line);line-height:1.45}}
ul.plain li b{{color:var(--mag);margin-right:6px}}
.two{{display:grid;grid-template-columns:1fr 1fr;gap:14px}} @media(max-width:820px){{.two{{grid-template-columns:1fr}}}}
footer{{text-align:center;color:var(--mut);font-size:12px;padding:20px}}
</style></head><body>
<header><h1>{e(title)}</h1><p>The strategic layer: principles, modalities, governance, and phasing — upstream of the training needs analysis.</p>
<span class="k"><b>Workforce types:</b> {e(wtypes)}</span>
<span class="k"><b>Go-live reference:</b> {e(meta.get("go_live_reference",""))}</span>
<span class="k"><b>Generated:</b> {today}</span></header>
<main>
<section><h2>Guiding principles</h2><div class="grid">{tiles(plan.get("principles", []))}</div></section>
<section><h2>Modality strategy by workforce type</h2>
<table><tr><th>Workforce type</th><th>Primary modalities</th><th>Secondary</th><th>Rationale</th></tr>{modality_rows()}</table></section>
<section><h2>Governance — who owns what</h2>
<table><tr><th>Role</th><th>Owns</th></tr>{gov_rows()}</table></section>
<section class="two"><div><h2>Resourcing model</h2><div class="panel">
<h3>Core team</h3><p>{e(res.get("core_team",""))}</p>
<h3>Multipliers</h3><p>{e(res.get("multipliers",""))}</p>
<h3>Vendor stance</h3><p>{e(res.get("vendor_stance",""))}</p></div></div>
<div><h2>Measurement</h2><div class="panel">
<h3>Approach</h3><p>{e(meas.get("approach",""))}</p>
<h3>Links</h3><ul class="plain">{li(lst(meas.get("links")))}</ul></div></div></section>
<section><h2>Build vs buy</h2>
<table><tr><th>Category</th><th>Stance</th><th>Rationale</th></tr>{bvb_rows()}</table></section>
<section><h2>Phasing relative to go-live</h2><div class="steps">{phase_steps()}</div></section>
<section><h2>Risks</h2>
<table><tr><th>Risk</th><th>Mitigation</th><th>Owner</th></tr>{risk_rows()}</table></section>
<section class="two"><div><h2>Assumptions</h2><div class="panel"><ul class="plain">{li(plan.get("assumptions", []))}</ul></div></div>
<div><h2>Open questions</h2><div class="panel"><ul class="plain">{oq_items()}</ul></div></div></section>
</main><footer>{e(brand["footer"])} &middot; Strategy decisions constrain the downstream TNA and curriculum &mdash; no module lists or hours here by design.</footer>
</body></html>"""

# ----------------------------------------------------------------------------- cli
def _guard(outdir, names, force):
    """Refuse to overwrite existing deliverables unless --force is set."""
    existing = [n for n in names if os.path.exists(os.path.join(outdir, n))]
    if existing and not force:
        sys.exit("Refusing to overwrite existing output(s) in %r:\n  %s\n"
                 "Re-run with --force to overwrite, or choose a different --outdir."
                 % (outdir, "\n  ".join(existing)))

def main():
    ap = argparse.ArgumentParser(description="Render a Training Strategy (xlsx + HTML).")
    ap.add_argument("--plan", required=True, help="training_strategy.json")
    ap.add_argument("--config", required=True, help="project.json")
    ap.add_argument("--outdir", required=True)
    ap.add_argument("--force", action="store_true", help="overwrite existing outputs in --outdir")
    a = ap.parse_args()
    os.makedirs(a.outdir, exist_ok=True)
    project = json.load(open(a.config, encoding="utf-8"))
    plan = json.load(open(a.plan, encoding="utf-8"))
    brand = dict(DEFAULT_BRAND); brand.update(project.get("brand") or {})
    pn = plan.get("meta", {}).get("program") or project.get("project_name", "Project")
    xlsx_name = f"{pn} Training Strategy.xlsx"; html_name = f"{pn} Training Strategy.html"
    _guard(a.outdir, [xlsx_name, html_name], a.force)
    write_xlsx(os.path.join(a.outdir, xlsx_name), brand, plan)
    open(os.path.join(a.outdir, html_name), "w", encoding="utf-8").write(build_html(plan, project, brand))
    print(f"Training Strategy: {len(plan.get('principles', []))} principles / "
          f"{len(plan.get('modality_strategy', []))} workforce types -> {a.outdir}")

if __name__ == "__main__":
    main()
