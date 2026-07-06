#!/usr/bin/env python3
"""
training-needs-builder — render a Training Needs Analysis (TNA) as an xlsx
workbook + a self-contained interactive HTML dashboard, for ANY project, from
training_needs.json (see references/NEEDS_SCHEMA.md).

HARD RULE: a TNA never surfaces training hours / effort sizing — that is a
curriculum-stage output. Any hours/effort field in the input is dropped.

Usage:
  python3 needs_render.py --records training_needs.json --config project.json --outdir OUT [--force]

Only stdlib + openpyxl. Output HTML has no external dependencies.
"""
import argparse, html, json, os, re, sys

DEFAULT_BRAND = {"navy": "#162B75", "magenta": "#EE2C81",
                 "font": "Calibri, system-ui, sans-serif", "footer": "Confidential"}
PRI_COLORS = {"Critical": "#b40020", "High": "#d35400", "Medium": "#1d6fb8", "Low": "#6b7280"}
PRI_FILL   = {"Critical": "FDECEA", "High": "FDF3E7", "Medium": "E8F1FA", "Low": "F0F1F3"}
PROF = ["Novice", "Advanced Beginner", "Competent", "Proficient", "Expert"]
HOURS_RE = re.compile(r"hour|effort", re.I)

def as_list(v):
    if v is None or v == "": return []
    if isinstance(v, (list, tuple)): return [str(x).strip() for x in v if str(x).strip()]
    return [s.strip() for s in re.split(r"[;,]", str(v)) if s.strip()]

def normalize(records):
    out = []
    for i, r in enumerate(records, 1):
        r = {k: v for k, v in r.items() if not HOURS_RE.search(k)}  # TNA carries no hours/effort
        out.append({
            "id": r.get("id") or f"TN-{i:02d}",
            "capability": r.get("capability", ""),
            "audience": r.get("audience", ""),
            "business_unit": r.get("business_unit", ""),
            "source_impacts": as_list(r.get("source_impacts")),
            "driver_dimensions": as_list(r.get("driver_dimensions")),
            "current_proficiency": r.get("current_proficiency", ""),
            "target_proficiency": r.get("target_proficiency", ""),
            "gap": r.get("gap", ""),
            "bloom_level": r.get("bloom_level", ""),
            "priority": r.get("priority", ""),
            "recommended_modalities": as_list(r.get("recommended_modalities")),
            "population_size": r.get("population_size", ""),
            "access_constraints": r.get("access_constraints", ""),
            "training_alone_insufficient": bool(r.get("training_alone_insufficient")),
            "non_training_response": r.get("non_training_response", ""),
            "notes": r.get("notes", ""),
        })
    return out

# ----------------------------------------------------------------------------- xlsx
def style_sheet(ws, columns, navy):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    thin = Side(style="thin", color="D9D9D9"); border = Border(thin, thin, thin, thin)
    for ci, _ in enumerate(columns, 1):
        c = ws.cell(1, ci); c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=navy.lstrip("#"))
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = border
    ws.freeze_panes = "B2"; ws.row_dimensions[1].height = 26
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
    for ci, col in enumerate(columns, 1):
        L = get_column_letter(ci); ws.column_dimensions[L].width = col[2] if len(col) > 2 else 22
        for cell in ws[L][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top"); cell.border = border
            cell.font = Font(size=10)

def fmt(v):
    if isinstance(v, bool): return "Yes" if v else ""
    if isinstance(v, (list, tuple)): return ", ".join(str(x) for x in v)
    return v

def write_xlsx(path, data, navy):
    from openpyxl import Workbook
    wb = Workbook()
    cols = [("id", "ID", 8), ("capability", "Capability", 46), ("audience", "Audience", 28),
            ("business_unit", "Business Unit", 14), ("source_impacts", "Source Impacts (CIA)", 40),
            ("driver_dimensions", "Driver Dimensions", 28),
            ("current_proficiency", "Current Proficiency", 16), ("target_proficiency", "Target Proficiency", 16),
            ("gap", "Gap", 10), ("bloom_level", "Bloom Level", 12), ("priority", "Priority", 10),
            ("recommended_modalities", "Recommended Modalities", 30),
            ("population_size", "Population Size", 22), ("access_constraints", "Access Constraints", 34),
            ("training_alone_insufficient", "Training Alone Insufficient", 14),
            ("non_training_response", "Non-Training Response", 44), ("notes", "Notes", 40)]
    ws = wb.active; ws.title = "Needs Matrix"
    ws.append([c[1] for c in cols])
    for r in data:
        ws.append([fmt(r.get(c[0], "")) for c in cols])
    style_sheet(ws, cols, navy)

    # by-audience summary
    auds = {}
    for r in data:
        a = auds.setdefault((r["audience"], r["business_unit"]),
                            {"needs": 0, "Critical": 0, "High": 0, "Medium": 0, "Low": 0, "flagged": 0, "mods": set()})
        a["needs"] += 1
        if r["priority"] in a: a[r["priority"]] += 1
        if r["training_alone_insufficient"]: a["flagged"] += 1
        a["mods"].update(r["recommended_modalities"])
    scols = [("audience", "Audience", 30), ("business_unit", "Business Unit", 14), ("needs", "Needs", 8),
             ("Critical", "Critical", 9), ("High", "High", 8), ("Medium", "Medium", 9), ("Low", "Low", 8),
             ("flagged", "Training-Alone-Insufficient", 14), ("mods", "Modalities in Play", 40)]
    ws2 = wb.create_sheet("By Audience")
    ws2.append([c[1] for c in scols])
    for (aud, bu), a in sorted(auds.items()):
        ws2.append([aud, bu, a["needs"], a["Critical"], a["High"], a["Medium"], a["Low"],
                    a["flagged"], ", ".join(sorted(a["mods"]))])
    style_sheet(ws2, scols, navy)
    wb.save(path)

# ----------------------------------------------------------------------------- html
TEMPLATE = r"""<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0"><title>%%TITLE%%</title>
<style>
:root{--navy:%%NAVY%%;--mag:%%MAG%%;--ink:#1d2433;--mut:#6b7280;--line:#e5e7eb;--bg:#f5f6f8}
*{box-sizing:border-box}body{margin:0;font-family:%%FONT%%;color:var(--ink);background:var(--bg);font-size:14px}
header{background:linear-gradient(135deg,var(--navy),#0E2841);color:#fff;padding:18px 28px}
header h1{margin:0;font-size:1.15rem;font-weight:700}header .sub{opacity:.85;font-size:.85rem;margin-top:3px}
.wrap{max-width:1500px;margin:0 auto;padding:18px 28px 80px}
.stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:14px;margin:16px 0}
.card{background:#fff;border:1px solid var(--line);border-radius:10px;padding:14px 16px}
.card .v{font-size:1.6rem;font-weight:800;color:var(--navy)}.card .l{font-size:.72rem;color:var(--mut);text-transform:uppercase;letter-spacing:.04em;margin-top:2px}
.bar{display:flex;flex-wrap:wrap;gap:8px;align-items:center;margin:10px 0}
.bar select,.bar input{padding:7px 10px;border:1px solid var(--line);border-radius:7px;font:inherit;background:#fff}
.bar input{flex:1;min-width:180px}.btn{background:var(--navy);color:#fff;border:none;border-radius:7px;padding:8px 14px;cursor:pointer;font:inherit}
table{width:100%;border-collapse:collapse;background:#fff;border:1px solid var(--line);border-radius:10px;overflow:hidden}
thead th{position:sticky;top:0;background:var(--navy);color:#fff;text-align:left;padding:9px 11px;font-size:.74rem;white-space:nowrap}
tbody td{padding:9px 11px;border-top:1px solid var(--line);font-size:.82rem;vertical-align:top}
tbody tr.row{cursor:pointer}tbody tr.row:hover{background:#f0f3fb}
tr.p-Critical td:first-child{border-left:4px solid #b40020}tr.p-High td:first-child{border-left:4px solid #d35400}
tr.p-Medium td:first-child{border-left:4px solid #1d6fb8}tr.p-Low td:first-child{border-left:4px solid #6b7280}
.pri{display:inline-block;padding:2px 9px;border-radius:6px;color:#fff;font-size:.68rem;font-weight:700}
.chip{display:inline-block;padding:2px 9px;border-radius:11px;background:#5a6b86;color:#fff;font-size:.68rem;font-weight:600;margin:1px 2px 1px 0;white-space:nowrap}
.chip.dim{background:#7A4FBE}.chip.src{background:#0E2841;opacity:.85}
.prof{white-space:nowrap;font-size:.78rem}.prof b{color:var(--navy)}
.prof .arrow{color:var(--mag);font-weight:800;padding:0 4px}
.gapdots{display:inline-flex;gap:2px;margin-left:6px;vertical-align:middle}
.gapdots i{width:8px;height:8px;border-radius:50%;background:#dfe3ea}
.gapdots i.on{background:var(--mag)}
.flag{display:inline-block;padding:2px 8px;border-radius:9px;background:#fdecea;color:#b40020;font-size:.64rem;font-weight:800;text-transform:uppercase;white-space:nowrap}
.detail{background:#fbfcfe}.detail .sec{margin:9px 0}
.detail .lbl{font-size:.68rem;text-transform:uppercase;letter-spacing:.04em;color:var(--mut);font-weight:700;margin-bottom:3px}
.warnbox{border:1px solid #f3c1cf;background:#fdf0f5;border-radius:8px;padding:9px 12px;margin:9px 0}
.warnbox .lbl{color:#b40020}
footer{text-align:center;color:var(--mut);font-size:.72rem;padding:18px}
</style></head><body>
<header><h1>%%TITLE%%</h1><div class="sub">%%SUBTITLE%%</div></header>
<div class="wrap">
  <div class="stats" id="stats"></div>
  <div class="bar" id="bar"></div>
  <div style="overflow:auto"><table><thead><tr id="head"></tr></thead><tbody id="body"></tbody></table></div>
</div>
<footer>%%FOOTER%%</footer>
<script>
const DATA=%%DATA%%, PRI_COLORS=%%PRICOLORS%%, PROF=%%PROF%%;
const $=s=>document.querySelector(s), esc=s=>String(s==null?'':s).replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
let expanded=null;
const FILTERS=[["audience","All Audiences"],["business_unit","All BUs"],["priority","All Priorities"],["bloom_level","All Bloom Levels"]];
function distinct(k){return [...new Set(DATA.map(d=>d[k]).filter(Boolean))].sort();}
function buildBar(){
  const bar=$('#bar');
  FILTERS.forEach(([k,l])=>{const s=document.createElement('select');s.id='f-'+k;
    s.innerHTML='<option value="">'+esc(l)+'</option>'+distinct(k).map(v=>'<option>'+esc(v)+'</option>').join('');
    s.onchange=render;bar.appendChild(s);});
  const q=document.createElement('input');q.id='q';q.placeholder='Search…';q.oninput=render;bar.appendChild(q);
  const b=document.createElement('button');b.className='btn';b.textContent='Reset';
  b.onclick=()=>{FILTERS.forEach(([k])=>$('#f-'+k).value='');$('#q').value='';render();};bar.appendChild(b);
}
function filtered(){
  let d=DATA.slice();
  FILTERS.forEach(([k])=>{const v=$('#f-'+k).value;if(v)d=d.filter(x=>String(x[k])===v);});
  const q=($('#q').value||'').toLowerCase();
  if(q)d=d.filter(x=>JSON.stringify(x).toLowerCase().includes(q));
  return d;
}
function profCell(d){
  const ci=PROF.indexOf(d.current_proficiency), ti=PROF.indexOf(d.target_proficiency);
  const dots=(ci>=0&&ti>=0)?('<span class="gapdots">'+PROF.map((_,i)=>'<i class="'+(i>ci&&i<=ti?'on':'')+'"></i>').join('')+'</span>'):'';
  return '<span class="prof">'+esc(d.current_proficiency||'—')+'<span class="arrow">→</span><b>'+esc(d.target_proficiency||'—')+'</b></span>'+dots;
}
function chips(a,cls){return (a||[]).map(x=>'<span class="chip '+(cls||'')+'">'+esc(x)+'</span>').join('')||'<span style="color:#bbb">—</span>';}
function render(){
  const rows=filtered();
  $('#head').innerHTML=['ID','Capability','Audience','BU','Priority','Proficiency (current → target)','Modalities','Flags'].map(h=>'<th>'+h+'</th>').join('');
  $('#body').innerHTML=rows.map(d=>{
    const pri='<span class="pri" style="background:'+(PRI_COLORS[d.priority]||'#6b7280')+'">'+esc(d.priority||'—')+'</span>';
    const flag=d.training_alone_insufficient?'<span class="flag">training alone insufficient</span>':'';
    const tr='<tr class="row p-'+esc(d.priority)+'" data-id="'+esc(d.id)+'"><td>'+esc(d.id)+'</td><td>'+esc(d.capability)+'</td><td>'+esc(d.audience)+'</td><td>'+esc(d.business_unit)+'</td><td>'+pri+'</td><td>'+profCell(d)+'</td><td>'+chips(d.recommended_modalities)+'</td><td>'+flag+'</td></tr>';
    const det=(expanded===d.id)?'<tr class="detail"><td colspan="8">'+detail(d)+'</td></tr>':'';
    return tr+det;}).join('');
  $('#body').querySelectorAll('tr.row').forEach(tr=>tr.onclick=()=>{expanded=(expanded===tr.dataset.id)?null:tr.dataset.id;render();});
  buildStats(rows);
}
function detail(d){
  let h='';
  h+='<div class="sec"><div class="lbl">Traceability — source change impacts (CIA)</div>'+chips(d.source_impacts,'src')+'</div>';
  h+='<div class="sec"><div class="lbl">Driver dimensions</div>'+chips(d.driver_dimensions,'dim')+'</div>';
  if(d.gap)h+='<div class="sec"><div class="lbl">Gap</div>'+esc(d.gap)+'</div>';
  if(d.bloom_level)h+='<div class="sec"><div class="lbl">Bloom level</div>'+esc(d.bloom_level)+'</div>';
  if(d.population_size)h+='<div class="sec"><div class="lbl">Population size</div>'+esc(d.population_size)+'</div>';
  if(d.access_constraints)h+='<div class="sec"><div class="lbl">Access constraints</div>'+esc(d.access_constraints)+'</div>';
  if(d.training_alone_insufficient)h+='<div class="warnbox"><div class="lbl">Training alone insufficient — non-training response</div>'+esc(d.non_training_response||'See notes.')+'</div>';
  if(d.notes)h+='<div class="sec"><div class="lbl">Notes</div>'+esc(d.notes)+'</div>';
  return h;
}
function buildStats(rows){
  const stats=[['Needs',rows.length],['Audiences',new Set(rows.map(r=>r.audience)).size],
    ['Business Units',new Set(rows.map(r=>r.business_unit)).size],
    ['Critical',rows.filter(r=>r.priority==='Critical').length],
    ['High',rows.filter(r=>r.priority==='High').length],
    ['Training-Alone-Insufficient',rows.filter(r=>r.training_alone_insufficient).length]];
  $('#stats').innerHTML=stats.map(([l,v])=>'<div class="card"><div class="v">'+v+'</div><div class="l">'+esc(l)+'</div></div>').join('');
}
buildBar();render();
</script></body></html>"""

def write_html(path, data, project, brand):
    pname = project.get("project_name", "Project")
    auds = len({d["audience"] for d in data})
    out = (TEMPLATE
           .replace("%%TITLE%%", html.escape(f"{pname} — Training Needs Analysis"))
           .replace("%%SUBTITLE%%", html.escape(
               f"{len(data)} needs across {auds} audiences — no hours/effort here; sizing is a curriculum-stage output"))
           .replace("%%NAVY%%", brand["navy"]).replace("%%MAG%%", brand["magenta"])
           .replace("%%FONT%%", brand["font"]).replace("%%FOOTER%%", html.escape(brand["footer"]))
           .replace("%%DATA%%", json.dumps(data, ensure_ascii=False))
           .replace("%%PRICOLORS%%", json.dumps(PRI_COLORS))
           .replace("%%PROF%%", json.dumps(PROF)))
    open(path, "w", encoding="utf-8").write(out)

# ----------------------------------------------------------------------------- cli
def _guard(outdir, names, force):
    existing = [n for n in names if os.path.exists(os.path.join(outdir, n))]
    if existing and not force:
        sys.exit("Refusing to overwrite existing output(s) in %r:\n  %s\n"
                 "Re-run with --force to overwrite, or choose a different --outdir."
                 % (outdir, "\n  ".join(existing)))

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--records", required=True); ap.add_argument("--config", required=True)
    ap.add_argument("--outdir", default=".")
    ap.add_argument("--force", action="store_true", help="overwrite existing outputs in --outdir")
    a = ap.parse_args()
    os.makedirs(a.outdir, exist_ok=True)
    project = json.load(open(a.config, encoding="utf-8"))
    brand = {**DEFAULT_BRAND, **project.get("brand", {})}
    pname = project.get("project_name", "Project")
    names = [f"{pname} TNA.xlsx", f"{pname} TNA Dashboard.html"]
    _guard(a.outdir, names, a.force)
    data = normalize(json.load(open(a.records, encoding="utf-8")))
    write_xlsx(os.path.join(a.outdir, names[0]), data, brand["navy"])
    write_html(os.path.join(a.outdir, names[1]), data, project, brand)
    print(f"TNA: {len(data)} needs -> {a.outdir}")

if __name__ == "__main__":
    main()
