#!/usr/bin/env python3
"""
ocm-playbook-builder — render an OCM Playbook + 90-day tactical plan as an xlsx
workbook + a self-contained interactive HTML playbook, for ANY project, from a
structured playbook.json (written by Claude from engagement context).

Methodology baked in (portable OCM):
  - Workstreams + 30/60/90 (or custom) phased activities, each with owner,
    output, done test, dependencies; cadence & escalation registers; handoff.
  - With --cia, high-severity CIA impacts (severity >= 4) with no activity
    carrying them in linked_impacts are flagged as GAPs.

Usage:
  python3 playbook_render.py --plan playbook.json --config project.json --outdir OUT
  python3 playbook_render.py --plan playbook.json --config project.json --cia cia_records.json --outdir OUT

Only stdlib + openpyxl. Cross-platform. Output HTML has no external dependencies.
Render conventions follow the suite canon (cia-builder/scripts/cia_render.py).
"""
import argparse, datetime, html, json, os, sys

DEFAULT_BRAND = {"navy": "#162B75", "magenta": "#EE2C81", "orange": "#F08301",
                 "teal": "#04A577", "coral": "#FF533C", "font": "Calibri, system-ui, sans-serif",
                 "footer": "Confidential"}
PHASE_LABELS = {"30": "Days 1-30 · Mobilize & Assess", "60": "Days 31-60 · Build & Engage",
                "90": "Days 61-90 · Execute & Measure"}
STATUS_COLORS = {"Done": "#04A577", "In Progress": "#F08301", "Blocked": "#FF533C",
                 "Not Started": "#9aa3b2"}

def esc(s): return html.escape(str(s if s is not None else ""))

def g(rec, k, default=""):
    v = rec.get(k)
    return default if v in (None, "") else v

# ----------------------------------------------------------------------------- normalize
def normalize(plan):
    meta = plan.get("meta") or {}
    workstreams = plan.get("workstreams") or []
    ws_names = [w.get("name", "") for w in workstreams]
    acts = []
    for i, a in enumerate(plan.get("activities") or [], 1):
        ws = g(a, "workstream")
        if ws not in ws_names:
            ws = "(Unassigned)"
        acts.append({
            "id": g(a, "id", f"A{i:02d}"), "workstream": ws,
            "phase": str(g(a, "phase", "30")), "week": a.get("week", ""),
            "title": g(a, "title"), "description": g(a, "description"),
            "owner": g(a, "owner", "TBD"), "output": g(a, "output"),
            "done_test": g(a, "done_test"),
            "depends_on": a.get("depends_on") or [],
            "status": g(a, "status", "Not Started"),
            "linked_impacts": a.get("linked_impacts") or [],
        })
    if any(a["workstream"] == "(Unassigned)" for a in acts) and "(Unassigned)" not in ws_names:
        workstreams = workstreams + [{"name": "(Unassigned)"}]
    # phase order = 30/60/90 first (if used), then customs by first appearance
    phases, seen = [], set()
    for p in ["30", "60", "90"]:
        if any(a["phase"] == p for a in acts):
            phases.append(p); seen.add(p)
    for a in acts:
        if a["phase"] not in seen:
            phases.append(a["phase"]); seen.add(a["phase"])
    return meta, workstreams, acts, phases

def cia_is_unscored(cia_records):
    """True when all (or nearly all) CIA records lack a numeric severity,
    so the high-severity GAP check has nothing to test."""
    if not cia_records:
        return False
    scored = 0
    for r in cia_records:
        try:
            float(r.get("severity"))
            scored += 1
        except (TypeError, ValueError):
            pass
    return scored < max(1, len(cia_records) * 0.1)

def find_gaps(cia_records, acts):
    """High-severity CIA impacts (severity >= 4) with no activity linking them."""
    linked = set()
    for a in acts:
        linked.update(str(t).strip() for t in a["linked_impacts"])
    gaps = []
    for r in cia_records:
        try: sev = int(float(r.get("severity") or 0))
        except Exception: sev = 0
        title = str(r.get("title", "")).strip()
        if sev >= 4 and title and title not in linked:
            gaps.append({"title": title, "severity": sev,
                         "business_unit": g(r, "business_unit"),
                         "roles_impacted": g(r, "roles_impacted")})
    return gaps

# ----------------------------------------------------------------------------- xlsx
def write_xlsx(path, brand, meta, workstreams, acts, cadences):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    navy = brand["navy"].lstrip("#")
    thin = Side(style="thin", color="D9D9D9"); border = Border(thin, thin, thin, thin)
    wb = Workbook()

    def sheet(ws, columns, rows):
        ws.append([c[1] for c in columns])
        for ci in range(1, len(columns) + 1):
            c = ws.cell(1, ci); c.font = Font(bold=True, color="FFFFFF", size=11)
            c.fill = PatternFill("solid", fgColor=navy)
            c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = border
        ws.freeze_panes = "A2"; ws.row_dimensions[1].height = 26
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
        for r in rows:
            ws.append([", ".join(map(str, r.get(c[0]))) if isinstance(r.get(c[0]), list)
                       else r.get(c[0], "") for c in columns])
        for ci, col in enumerate(columns, 1):
            L = get_column_letter(ci)
            ws.column_dimensions[L].width = col[2] if len(col) > 2 else 22
            for cell in ws[L][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = border; cell.font = Font(size=10)

    act_cols = [("id", "ID", 7), ("workstream", "Workstream", 22), ("phase", "Phase", 8),
                ("week", "Week", 7), ("title", "Activity", 34), ("description", "Description", 40),
                ("owner", "Owner", 16), ("output", "Output", 30), ("done_test", "Done Test", 30),
                ("depends_on", "Depends On", 12), ("status", "Status", 12),
                ("linked_impacts", "Linked Impacts (CIA)", 34)]
    ws1 = wb.active; ws1.title = "Activities"
    sheet(ws1, act_cols, acts)

    ws2 = wb.create_sheet("By Workstream")
    by_ws_rows = []
    for w in workstreams:
        rows = sorted([a for a in acts if a["workstream"] == w.get("name")],
                      key=lambda a: (a["phase"], a["week"] if isinstance(a["week"], int) else 999))
        for a in rows:
            by_ws_rows.append({"workstream": w.get("name", ""), "lead": w.get("lead", ""),
                               "phase": a["phase"], "week": a["week"], "id": a["id"],
                               "title": a["title"], "owner": a["owner"], "status": a["status"]})
    sheet(ws2, [("workstream", "Workstream", 24), ("lead", "Workstream Lead", 16),
                ("phase", "Phase", 8), ("week", "Week", 7), ("id", "ID", 7),
                ("title", "Activity", 40), ("owner", "Owner", 16), ("status", "Status", 12)],
          by_ws_rows)

    ws3 = wb.create_sheet("Cadences")
    sheet(ws3, [("name", "Cadence", 26), ("frequency", "Frequency", 18),
                ("audience", "Audience", 26), ("owner", "Owner", 16), ("purpose", "Purpose", 44)],
          cadences)
    wb.save(path)

# ----------------------------------------------------------------------------- html
TEMPLATE = r"""<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0"><title>%%TITLE%%</title>
<style>
:root{--navy:%%NAVY%%;--mag:%%MAG%%;--ink:#1d2433;--mut:#6b7280;--line:#e5e7eb;--bg:#f5f6f8}
*{box-sizing:border-box;margin:0}body{font-family:%%FONT%%;background:var(--bg);color:var(--ink);font-size:14px}
header{background:var(--navy);color:#fff;padding:18px 28px}
header h1{font-size:20px}header .sub{opacity:.8;font-size:12px;margin-top:4px}
.wrap{max-width:1240px;margin:0 auto;padding:20px 28px 60px}
h2{font-size:15px;margin:26px 0 10px;color:var(--navy)}
.tiles{display:flex;gap:12px;flex-wrap:wrap}
.tile{background:#fff;border:1px solid var(--line);border-radius:10px;padding:14px 18px;min-width:150px;flex:1}
.tile .n{font-size:26px;font-weight:700;color:var(--navy)}.tile .l{font-size:11px;color:var(--mut);text-transform:uppercase;letter-spacing:.04em}
.tile.warn .n{color:%%CORAL%%}
.swim{overflow-x:auto}.swim table{border-collapse:collapse;width:100%;min-width:760px}
.swim th,.swim td{border:1px solid var(--line);padding:8px;vertical-align:top;background:#fff}
.swim th{background:var(--navy);color:#fff;font-size:12px;text-align:left}
.swim td.ws{font-weight:700;font-size:12px;width:170px;background:#fafbfc}
.swim td.ws .lead{font-weight:400;color:var(--mut);font-size:11px}
.card{border-left:3px solid var(--mag);background:#f8f9fc;border-radius:4px;padding:5px 8px;margin:4px 0;font-size:12px}
.card .wk{color:var(--mut);font-size:10px}
.dot{display:inline-block;width:8px;height:8px;border-radius:50%;margin-right:5px}
.controls{display:flex;gap:10px;flex-wrap:wrap;margin:10px 0}
select,input[type=search]{padding:6px 10px;border:1px solid var(--line);border-radius:6px;font-family:inherit;font-size:13px;background:#fff}
table.reg{border-collapse:collapse;width:100%;background:#fff}
table.reg th{background:var(--navy);color:#fff;font-size:11px;padding:7px 8px;text-align:left;position:sticky;top:0}
table.reg td{border-bottom:1px solid var(--line);padding:7px 8px;font-size:12px;vertical-align:top}
table.reg tr:hover td{background:#f4f6fb}
.regwrap{overflow-x:auto;max-height:560px;overflow-y:auto;border:1px solid var(--line);border-radius:8px}
.pill{display:inline-block;padding:1px 8px;border-radius:10px;font-size:11px;color:#fff}
.tbd{color:%%CORAL%%;font-weight:700}
.panel{background:#fff;border:1px solid var(--line);border-radius:10px;padding:16px 18px;margin-bottom:12px}
.panel h3{font-size:13px;color:var(--navy);margin-bottom:8px}
.panel ul{margin-left:18px}.panel li{margin:4px 0;font-size:13px}
.gap{border-left:4px solid %%CORAL%%;background:#fff4f2;border-radius:6px;padding:10px 12px;margin:6px 0;font-size:13px}
.gap b{color:%%CORAL%%}
.two{display:grid;grid-template-columns:1fr 1fr;gap:14px}@media(max-width:860px){.two{grid-template-columns:1fr}}
.chk{list-style:none;margin-left:0!important}
.chk li{padding-left:24px;position:relative}
.chk li::before{content:"";position:absolute;left:0;top:4px;width:13px;height:13px;border:2px solid var(--mut);border-radius:3px}
.chk li.done::before{border-color:#04A577;background:#04A577}
footer{color:var(--mut);font-size:11px;text-align:center;padding:16px}
</style></head><body>
<header><h1>%%TITLE%%</h1><div class="sub">%%SUBTITLE%%</div></header>
<div class="wrap">
<h2>Summary</h2><div class="tiles" id="tiles"></div>
%%GAPS%%
<h2>Plan swimlane (by workstream × phase)</h2><div class="swim" id="swim"></div>
<h2>Activity register</h2>
<div class="controls">
<select id="fWs"><option value="">All workstreams</option></select>
<select id="fPhase"><option value="">All phases</option></select>
<select id="fStatus"><option value="">All statuses</option></select>
<input type="search" id="fQ" placeholder="Search activities…">
</div>
<div class="regwrap"><table class="reg"><thead><tr>
<th>ID</th><th>Workstream</th><th>Phase</th><th>Wk</th><th>Activity</th><th>Owner</th><th>Output</th><th>Done test</th><th>Depends</th><th>Status</th><th>Linked impacts</th>
</tr></thead><tbody id="rows"></tbody></table></div>
<h2>Cadence &amp; governance</h2>
<div class="two">
<div class="panel"><h3>Standing cadences</h3><ul id="cad"></ul></div>
<div class="panel"><h3>Decision &amp; escalation path</h3><ul id="esc"></ul></div>
</div>
%%HANDOFF%%
</div>
<footer>%%FOOTER%% · generated %%DATE%%</footer>
<script>
const D=%%DATA%%;
const SC={"Done":"#04A577","In Progress":"#F08301","Blocked":"#FF533C","Not Started":"#9aa3b2"};
const el=id=>document.getElementById(id);
const e=s=>{const d=document.createElement('div');d.textContent=s==null?"":String(s);return d.innerHTML};
// tiles
const isTBD=o=>String(o||"").includes("TBD");
const tbd=D.acts.filter(a=>isTBD(a.owner)).length;
let tiles=[["Activities",D.acts.length],["Workstreams",D.workstreams.length],
 ["Done",D.acts.filter(a=>a.status==="Done").length],
 ["Owners TBD",tbd,tbd>0],["Cadences",D.cadences.length]];
if(D.gaps!==null) tiles.push(["CIA gaps",D.gaps.length,D.gaps.length>0]);
el('tiles').innerHTML=tiles.map(t=>`<div class="tile${t[2]?' warn':''}"><div class="n">${t[1]}</div><div class="l">${t[0]}</div></div>`).join('');
// swimlane
let sw='<table><tr><th>Workstream</th>'+D.phases.map(p=>`<th>${e(D.phase_labels[p]||("Phase "+p))}</th>`).join('')+'</tr>';
for(const w of D.workstreams){
 sw+=`<tr><td class="ws">${e(w.name)}${w.lead?`<div class="lead">${e(w.lead)}</div>`:''}</td>`;
 for(const p of D.phases){
  const cards=D.acts.filter(a=>a.workstream===w.name&&a.phase===p)
   .sort((a,b)=>(a.week||99)-(b.week||99))
   .map(a=>`<div class="card"><span class="dot" style="background:${SC[a.status]||'#9aa3b2'}"></span>${e(a.title)}<div class="wk">${a.week?('Wk '+a.week+' · '):''}${isTBD(a.owner)?`<span class="tbd">${e(a.owner)}</span>`:e(a.owner)}</div></div>`).join('');
  sw+=`<td>${cards}</td>`;}
 sw+='</tr>';}
el('swim').innerHTML=sw+'</table>';
// filters
const uniq=k=>[...new Set(D.acts.map(a=>a[k]))];
uniq('workstream').forEach(v=>el('fWs').insertAdjacentHTML('beforeend',`<option>${e(v)}</option>`));
D.phases.forEach(v=>el('fPhase').insertAdjacentHTML('beforeend',`<option value="${e(v)}">${e(D.phase_labels[v]||v)}</option>`));
uniq('status').forEach(v=>el('fStatus').insertAdjacentHTML('beforeend',`<option>${e(v)}</option>`));
function draw(){
 const ws=el('fWs').value,ph=el('fPhase').value,st=el('fStatus').value,q=el('fQ').value.toLowerCase();
 el('rows').innerHTML=D.acts.filter(a=>(!ws||a.workstream===ws)&&(!ph||a.phase===ph)&&(!st||a.status===st)
  &&(!q||(a.title+' '+a.description+' '+a.owner+' '+a.output).toLowerCase().includes(q)))
  .map(a=>`<tr><td>${e(a.id)}</td><td>${e(a.workstream)}</td><td>${e(a.phase)}</td><td>${e(a.week)}</td>
   <td><b>${e(a.title)}</b>${a.description?`<br><span style="color:var(--mut)">${e(a.description)}</span>`:''}</td>
   <td>${isTBD(a.owner)?`<span class="tbd">${e(a.owner)}</span>`:e(a.owner)}</td>
   <td>${e(a.output)}</td><td>${e(a.done_test)}</td><td>${e(a.depends_on.join(', '))}</td>
   <td><span class="pill" style="background:${SC[a.status]||'#9aa3b2'}">${e(a.status)}</span></td>
   <td>${e(a.linked_impacts.join('; '))}</td></tr>`).join('');}
['fWs','fPhase','fStatus'].forEach(id=>el(id).onchange=draw);el('fQ').oninput=draw;draw();
// cadences + escalation
el('cad').innerHTML=D.cadences.map(c=>`<li><b>${e(c.name)}</b> — ${e(c.frequency)}${c.owner?` · ${e(c.owner)}`:''}${c.audience?` · ${e(c.audience)}`:''}${c.purpose?`<br><span style="color:var(--mut)">${e(c.purpose)}</span>`:''}</li>`).join('')||'<li>None defined.</li>';
el('esc').innerHTML=D.escalation.map(x=>`<li><b>L${e(x.level)} ${e(x.forum)}</b> — ${e(x.trigger)}<br><span style="color:var(--mut)">Decides: ${e(x.decision_authority)}</span></li>`).join('')||'<li>None defined.</li>';
</script></body></html>"""

def build_html(path, brand, project, meta, workstreams, acts, phases, cadences,
               escalation, handoff, gaps, cia_unscored=False):
    sub = " · ".join(x for x in [
        g(meta, "program_phase"), f"Start: {g(meta,'start_reference')}" if g(meta, "start_reference") else "",
        f"Go-live horizon: {g(meta,'go_live_horizon')}" if g(meta, "go_live_horizon") else "",
        f"v{g(meta,'version')}" if g(meta, "version") else ""] if x)
    gaps_html = ""
    if cia_unscored:
        gaps_html = ('<h2>CIA coverage gaps</h2>'
                     '<div class="panel" style="border-left:4px solid #F08301;background:#fff8ee">'
                     '<b style="color:#B45309">WARNING</b> — high-severity coverage not evaluated '
                     '&mdash; CIA unscored; score the CIA first.</div>')
    elif gaps is not None:
        if gaps:
            items = "".join(f'<div class="gap"><b>GAP</b> — high-severity impact with no linked activity: '
                            f'<b>{esc(x["title"])}</b> (severity {x["severity"]}'
                            f'{", " + esc(x["business_unit"]) if x["business_unit"] else ""})'
                            f'{"<br>Roles: " + esc(x["roles_impacted"]) if x["roles_impacted"] else ""}</div>'
                            for x in gaps)
            gaps_html = f"<h2>CIA coverage gaps</h2>{items}"
        else:
            gaps_html = ('<h2>CIA coverage gaps</h2><div class="panel">All high-severity CIA impacts '
                         'have at least one linked activity.</div>')
    handoff_html = ""
    if handoff:
        chk = "".join(f'<li class="{ "done" if str(i.get("status","")).lower()=="ready" else "" }">'
                      f'{esc(i.get("item",""))} <span style="color:var(--mut)">— {esc(i.get("status",""))}</span></li>'
                      for i in handoff.get("items") or [])
        oq = "".join(f"<li>{esc(q)}</li>" for q in handoff.get("open_questions") or [])
        handoff_html = (f'<h2>Handoff</h2><div class="two"><div class="panel">'
                        f'<h3>Handoff checklist — to {esc(handoff.get("recipient",""))}'
                        f'{" (backup: " + esc(handoff["backup"]) + ")" if handoff.get("backup") else ""}</h3>'
                        f'<ul class="chk">{chk}</ul></div>'
                        f'<div class="panel"><h3>Open questions for the recipient</h3><ul>{oq or "<li>None.</li>"}</ul></div></div>')
    data = {"workstreams": workstreams, "acts": acts, "phases": phases,
            "phase_labels": PHASE_LABELS, "cadences": cadences,
            "escalation": escalation, "gaps": gaps}
    out = (TEMPLATE
           .replace("%%TITLE%%", esc(f"{project} — OCM Playbook"))
           .replace("%%SUBTITLE%%", esc(sub))
           .replace("%%NAVY%%", brand["navy"]).replace("%%MAG%%", brand["magenta"])
           .replace("%%CORAL%%", brand.get("coral", DEFAULT_BRAND["coral"]))
           .replace("%%FONT%%", brand["font"]).replace("%%FOOTER%%", esc(brand["footer"]))
           .replace("%%DATE%%", datetime.date.today().isoformat())
           .replace("%%GAPS%%", gaps_html).replace("%%HANDOFF%%", handoff_html)
           .replace("%%DATA%%", json.dumps(data, ensure_ascii=False)))
    with open(path, "w", encoding="utf-8") as f:
        f.write(out)

# ----------------------------------------------------------------------------- main
def _guard(outdir, names, force):
    """Refuse to overwrite existing deliverables unless --force is set."""
    existing = [n for n in names if os.path.exists(os.path.join(outdir, n))]
    if existing and not force:
        sys.exit("Refusing to overwrite existing outputs: " + ", ".join(existing) +
                 "\nRe-run with --force to overwrite, or choose a different --outdir.")

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--plan", required=True)
    ap.add_argument("--config", required=True)
    ap.add_argument("--cia", help="optional cia_records.json for GAP flagging")
    ap.add_argument("--outdir", default=".")
    ap.add_argument("--force", action="store_true", help="overwrite existing outputs in --outdir")
    a = ap.parse_args()
    with open(a.plan, encoding="utf-8") as f: plan = json.load(f)
    with open(a.config, encoding="utf-8") as f: cfg = json.load(f)
    brand = dict(DEFAULT_BRAND); brand.update(cfg.get("brand") or {})
    project = cfg.get("project_name", "Project")
    os.makedirs(a.outdir, exist_ok=True)
    xlsx_name = f"{project} OCM Tactical Plan.xlsx"
    html_name = f"{project} OCM Playbook.html"
    _guard(a.outdir, [xlsx_name, html_name], a.force)

    meta, workstreams, acts, phases = normalize(plan)
    cadences = plan.get("cadences") or []
    escalation = plan.get("escalation") or []
    handoff = plan.get("handoff") or {}
    gaps = None
    cia_unscored = False
    if a.cia:
        with open(a.cia, encoding="utf-8") as f:
            cia_records = json.load(f)
        cia_unscored = cia_is_unscored(cia_records)
        if cia_unscored:
            print("WARNING: high-severity coverage not evaluated — CIA unscored; "
                  "score the CIA first.")
            gaps = None
        else:
            gaps = find_gaps(cia_records, acts)

    write_xlsx(os.path.join(a.outdir, xlsx_name), brand, meta, workstreams, acts, cadences)
    build_html(os.path.join(a.outdir, html_name), brand, project, meta, workstreams,
               acts, phases, cadences, escalation, handoff, gaps, cia_unscored)
    print(f"Wrote: {os.path.join(a.outdir, xlsx_name)}")
    print(f"Wrote: {os.path.join(a.outdir, html_name)}")
    if gaps:
        print(f"GAPs: {len(gaps)} high-severity CIA impact(s) with no linked activity.")

if __name__ == "__main__":
    main()
